"""Translator mode (Decision 26, Phase V2-10): the setting and its override, the history and feedback fields, the
provenance and licence of the preset values, the presets on a fresh install, the miss gate and the shipped template.

V10.5  every command in translator mode records the user's wording, whether the strict core missed it, the rule,
       synonym or preset that fired, the phrase handed to the core, any proposal, and whether the user confirmed,
       changed or declined it, with the translator's name, version and rules SHA-256; Export feedback carries them
V10.6  (owner ruling, 21 October 2026) every shipped preset value has a source and a licence tag; standard values
       (IEEE 802.3 standards and drafts, the COM 3.70 workbooks, OIF-CEI-05.3) carry the COM workbook terms; nothing
       from an OIF member working text is shipped; NOTICE carries the License Notice; every preset works on a fresh
       install, with no overlay; a value a preset takes from the standard is not asked about by rule Q14, while the
       same value typed by hand still is

The use-case rows of groups PS, PP, CF and SY score the outcomes themselves (tests\\test_use_cases.py, gate V10.2).

    python tests\\test_translate.py
"""
# Copyright 2026 Todd Bermensolo
# SPDX-License-Identifier: BSD-3-Clause
import csv
import hashlib
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, 'cases'))
sys.dont_write_bytecode = True

from check import check, finish  # noqa: E402
import fixtures  # noqa: E402

PKG = os.path.join(ROOT, 'sicopr_assist')
PRESETS = os.path.join(PKG, 'data', 'presets.csv')
TEMPLATE = os.path.join(PKG, 'data', 'templates', 'config_com-4p13p0_802p3dj_d2p3_200G_C2M_TP0_TP2_Egress_26_01_27.xlsx')
TEMPLATE_SHA256 = '57e9da3f17dee495f386622f2bab30895a55b2ed40f2df997636d1a73af62e6e'   # 802-COM com_code 4.17.0
ENV = 'SICOPR_ASSIST_TRANSLATOR'
TRANSLATION_FIELDS = ('translator', 'version', 'rules_sha256', 'wording', 'strict', 'fired', 'preset', 'phrase',
                      'proposal', 'used', 'response')


# ----------------------------------------------------------------------------- V10.6 provenance and licence
# Owner ruling, 21 October 2026: COM parameter values from the IEEE 802.3 standards and drafts, the COM 3.70 workbooks
# and the public OIF-CEI-05.3 ship under the terms of the COM configuration workbooks, each tagged with its source;
# nothing whose only source is an OIF member working text (the CEI-224G MR, VSR and LR drafts) ships.
L_COM = 'BSD-3-Clause, 802-COM workbook (Copyright 2025 802-COM Authors)'
L_STD = 'BSD-3-Clause, standard value under the COM workbook terms (Copyright 2025 802-COM Authors)'
L_OWN = 'BSD-3-Clause, SiCoPR-Assist (Copyright 2026 Todd Bermensolo)'
STANDARD_SOURCES = re.compile(r'^(?:IEEE Std 802\.3ck-2022|IEEE P802\.3dj/D2\.4|IEEE COM 3\.70 workbook|OIF-CEI-05\.3 )')
MEMBER_TEXTS = re.compile(r'oif20\d\d\.\d|working text|CEI-224G|CEI-06\.0|\bVSR\b|224G-LR', re.I)


def provenance(fx=None):
    with open(PRESETS, newline='', encoding='utf-8') as f:
        rows = list(csv.DictReader(f))
    check('V10.6: the preset file has the columns preset, keyword, value, when, source, license',
          rows and list(rows[0]) == ['preset', 'keyword', 'value', 'when', 'source', 'license'],
          list(rows[0]) if rows else None)
    check('V10.6: every shipped preset value names its source', all(r['source'].strip() for r in rows),
          [r['preset'] + ' ' + r['keyword'] for r in rows if not r['source'].strip()])
    check('V10.6: every shipped preset value carries a licence tag (the 802-COM workbook terms, a standard value under '
          'them, or SiCoPR-Assist\'s own)', all(r.get('license') in (L_COM, L_STD, L_OWN) for r in rows),
          sorted({r.get('license') for r in rows} - {L_COM, L_STD, L_OWN}))
    std = [r for r in rows if STANDARD_SOURCES.search(r['source'])]
    check('V10.6: every value taken from an IEEE standard or draft, a COM 3.70 workbook or OIF-CEI-05.3 is tagged as a '
          'standard value under the COM workbook terms (%d values)' % len(std),
          std and all(r.get('license') == L_STD for r in std),
          [(r['preset'], r['keyword']) for r in std if r.get('license') != L_STD][:5])
    check('V10.6: a value tagged as a standard value names a standard, a draft, a COM 3.70 workbook or OIF-CEI-05.3',
          all(STANDARD_SOURCES.search(r['source']) for r in rows if r.get('license') == L_STD),
          [(r['preset'], r['keyword'], r['source'][:60]) for r in rows if r.get('license') == L_STD
           and not STANDARD_SOURCES.search(r['source'])][:5])
    member = [(r['preset'], r['keyword'], r['source'][:80]) for r in rows if MEMBER_TEXTS.search(r['source'])]
    check('V10.6: nothing from an OIF member working text is shipped (no source names a CEI-224G draft or an oif20xx '
          'working text)', not member, member[:5])
    check('V10.6: no preset value is left to a local overlay (no marker overlay, no whole-preset marker *)',
          not [r for r in rows if r['value'] == 'overlay' or r['keyword'] == '*'],
          [(r['preset'], r['keyword']) for r in rows if r['value'] == 'overlay' or r['keyword'] == '*'])
    for name in ('802.3ck KR', '802.3ck C2C', '802.3ck CR', '802.3ck C2M eye', '802.3dj C2M', 'low power CTLE only',
                 'low power long DFE'):
        mine = [r for r in rows if r['preset'] == name]
        check('V10.6: the %s preset ships its values (%d)' % (name, len(mine)), mine, name)
    tdr = [r['value'] for r in rows if r['preset'] == '802.3ck C2M eye' and r['keyword'] == 'TDR_W_TXPKG']
    check('ruling 4: the 802.3ck C2M eye preset takes TDR_W_TXPKG = 1, the value of its COM 3.70 workbook', tdr == ['1'],
          tdr)


def license_notice():
    """NOTICE carries the License Notice of the COM workbooks, line for line as the shipped template's sheet holds it,
    and says that the preset values ship under those terms; the README says so too."""
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True)
    lines = [str(c.value).rstrip() for r in wb['License Notice'].iter_rows() for c in r if c.value is not None]
    wb.close()
    with open(os.path.join(ROOT, 'NOTICE'), encoding='utf-8') as f:
        notice = f.read()
    got = [ln.rstrip() for ln in notice.splitlines()]
    at = next((i for i in range(len(got)) if got[i:i + len(lines)] == lines), None)
    check('V10.6: NOTICE carries the COM workbooks\' License Notice line for line (%d lines, from the shipped '
          'template\'s License Notice sheet)' % len(lines), lines and at is not None, lines[:3])
    check('V10.6: NOTICE says the preset values ship under those terms, each tagged with its source',
          'presets.csv' in notice and 'license' in notice and 'P802.3dj' in notice and 'OIF-CEI-05.3' in notice
          and 'working text' in notice and 'preset_overlay' not in notice)
    with open(os.path.join(ROOT, 'README.md'), encoding='utf-8') as f:
        readme = f.read()
    check('V10.6: the README states the terms of the preset values and no longer describes an overlay',
          'Copyright 2025 802-COM Authors' in readme and 'preset_overlay' not in readme)


def template_ships():
    with open(TEMPLATE, 'rb') as f:
        sha = hashlib.sha256(f.read()).hexdigest()
    check('V10.6: the 802-COM C2M egress template ships byte for byte (SHA-256 of com_code 4.17.0)', sha == TEMPLATE_SHA256,
          sha)
    import openpyxl
    wb = openpyxl.load_workbook(TEMPLATE, read_only=True)
    notice = wb['License Notice'] if 'License Notice' in wb.sheetnames else None
    text = ' '.join(str(c.value) for r in notice.iter_rows(max_row=40) for c in r if c.value) if notice else ''
    check('V10.6: the template keeps its License Notice sheet (BSD-3-Clause, 802-COM Authors)',
          'Copyright 2025 802-COM Authors' in text and 'Redistribution and use' in text, text[:120])
    with open(os.path.join(ROOT, 'NOTICE'), encoding='utf-8') as f:
        notice_file = f.read()
    check('V10.6: NOTICE names the shipped template', os.path.basename(TEMPLATE) in notice_file)


# ----------------------------------------------------------------------------- the setting, the presets on a fresh install
def _after(o, path):
    return next((c['after'] for c in o.get('changes') or [] if c['json_path'] == path), None)


def _fresh(fx, phrase, answers=()):
    """The outcome of a phrase (and its answers) in a fresh project: the starting config and a THRU, no settings file,
    so the defaults are in force (translator on, no overlay of any kind)."""
    import prior_states
    from sicopr_assist import core
    from sicopr_assist.project import Project
    ctx = prior_states.Ctx(fx, 'Case1')
    try:
        o = core.handle(phrase, Project(ctx.root))
        for a in answers:
            o = core.handle(a, Project(ctx.root))
        return o
    finally:
        fx.remove_project(ctx.root)


def presets_fresh_install(fx):
    """Owner rulings of 21 October 2026: every preset works on a fresh install (no overlay), the 802.3ck presets run in
    one step as a PLAN (a value a preset takes from the standard is not asked about by rule Q14), the 802.3ck C2M eye
    preset takes TDR_W_TXPKG = 1, and a b_min typed by hand is still asked about."""
    from sicopr_assist import validate
    old = os.environ.pop(ENV, None)
    try:
        for iface in ('KR', 'C2C', 'CR'):
            o = _fresh(fx, 'Setup a 802.3ck config for %s and run my channels.' % iface)
            tr = o.get('translation') or {}
            check('ruling 1, fresh install: the 802.3ck %s preset is one PLAN, its b_min(1) from the standard not asked '
                  'about by Q14' % iface, o['type'] == 'PLAN' and tr.get('preset') == '802.3ck %s' % iface and
                  _after(o, '/main/b_min(1)') not in (None, '', 0, '0'), _brief(o))
            check('ruling 1: the plan of 802.3ck %s says the preset\'s values were not asked about (rule Q14)' % iface,
                  any('Q14' in a and 'preset' in a for a in o.get('announce') or []), o.get('announce'))
        o = _fresh(fx, '100G per lane backplane please')
        check('ruling 1: "100G per lane backplane please" (blind B170) is the 802.3ck KR preset in one PLAN',
              o['type'] == 'PLAN' and (o.get('translation') or {}).get('preset') == '802.3ck KR', _brief(o))
        o = _fresh(fx, 'Setup a 802.3ck config for C2M and run my channels.', ['eye height and VEC'])
        check('rulings 1 and 4, fresh install: the 802.3ck C2M eye height and VEC run is one PLAN with TDR_W_TXPKG = 1',
              o['type'] == 'PLAN' and str(_after(o, '/main/TDR_W_TXPKG')) in ('1', '1.0'),
              (_brief(o), _after(o, '/main/TDR_W_TXPKG')))
        o = _fresh(fx, 'Setup a 802.3ck config for C2M and run my channels.', ['C2C-parameter COM'])
        check('ruling 1, fresh install: the 802.3ck C2M C2C-parameter COM is one PLAN', o['type'] == 'PLAN', _brief(o))
        o = _fresh(fx, 'Setup a 802.3dj config for C2M and run my channels.')
        check('ruling 5, fresh install: the 802.3dj C2M preset plans with the draft\'s eta_0 = 7.5e-9',
              o['type'] == 'PLAN' and str(_after(o, '/main/eta_0')) in ('7.5e-9', '7.5e-09'), (_brief(o), _after(o, '/main/eta_0')))
        q = _fresh(fx, 'Run COM on data\\thru.s4p with data\\config.xlsx.', ['Restrict the receiver equalization to low power.'])
        opts = [x.lower() for x in (q.get('question') or {}).get('options') or []]
        check('ruling 5, fresh install: the low-power question offers all five receivers', q['type'] == 'QUESTION' and
              len(opts) == 6 and 'long dfe (112g-mr-like)' in opts and 'overlay' not in q['question']['text'],
              (q['type'], opts))
        o = _fresh(fx, 'Run COM on data\\thru.s4p with data\\config.xlsx.',
                   ['Restrict the receiver equalization to low power.', 'CTLE only'])
        check('ruling 5, fresh install: CTLE only takes the OIF-CEI-05.3 XSR ranges (g_DC = [-8:1:0])',
              o['type'] == 'PLAN' and str(_after(o, '/main/g_DC')).replace(' ', '') == '[-8:1:0]',
              (_brief(o), _after(o, '/main/g_DC')))
        # a value typed by hand is still asked about (strict mode, and translator mode, where the strict core reads it)
        for mode in ('off', 'on'):
            os.environ[ENV] = mode
            o = _fresh(fx, 'Run COM on data\\thru.s4p with data\\config.xlsx and b_min(1) = 0.3.')
            check('ruling 1: b_min(1) = 0.3 typed by hand still asks rule Q14 (translator %s)' % mode,
                  o['type'] == 'QUESTION' and 'minimum above 0' in o['question']['text'], _brief(o))
        os.environ.pop(ENV, None)
        wc = fx.working_config('Case1')
        r = validate.validate(wc, [('/main/b_min(1)', '0.3')], allow=['Q14:b_min(2..N_b)'])
        check('ruling 1: a preset\'s pass on Q14 is for its own keywords only (one for b_min(2..N_b) leaves b_min(1) '
              'asked)', r['status'] == 'ask' and r['rule'] == 'Q14', (r['status'], r['rule']))
        r = validate.validate(wc, [('/main/b_min(1)', '0.3')], allow=['Q14:b_min(1)'])
        check('ruling 1: a preset\'s pass on Q14 for b_min(1) turns the question into an announcement',
              r['status'] == 'ok' and any('Q14' in t or 'minimum above 0' in t for t in r['announce']),
              (r['status'], r['announce']))
    finally:
        if old is None:
            os.environ.pop(ENV, None)
        else:
            os.environ[ENV] = old


def settings_and_mode(fx):
    import prior_states
    from sicopr_assist import core
    from sicopr_assist.project import Project, read_settings
    old = os.environ.pop(ENV, None)
    try:
        ctx = prior_states.Ctx(fx, 'Case1')
        prior_states.seed_run(ctx)
        s = read_settings(ctx.project)['values']
        check('V10.6: the setting translator (on) is shipped with its default, and there is no setting preset_overlay '
              '(ruling 5: nothing needs an overlay)', s['translator'] is True and 'preset_overlay' not in s,
              (s['translator'], s.get('preset_overlay')))
        o = core.handle('Re-run with baud rate 53.125.', Project(ctx.root))
        check('translator on (the default): baud rate is f_b', o['type'] == 'PLAN' and
              any(c['json_path'] == '/main/f_b' and c.get('kind') == 'set' for c in o['changes']), _brief(o))
        prior_states.project_settings(ctx, 'translator = off')
        o = core.handle('Re-run with baud rate 53.125.', Project(ctx.root))
        check('translator = off in the project settings: strict mode, the phrase asks', o['type'] == 'QUESTION'
              and not o.get('translation'), _brief(o))
        core.handle('cancel', Project(ctx.root))                # the strict question above is answered first
        os.environ[ENV] = 'on'
        o = core.handle('Re-run with baud rate 53.125.', Project(ctx.root))
        check('the environment override %s = on wins over the setting' % ENV, o['type'] == 'PLAN', _brief(o))
        fx.remove_project(ctx.root)
    finally:
        if old is None:
            os.environ.pop(ENV, None)
        else:
            os.environ[ENV] = old


def _brief(o):
    return (o['type'], o.get('reason') or (o.get('question') or {}).get('text') or [s['op'] for s in o['steps']])


# ----------------------------------------------------------------------------- V10.5 history and feedback
def history_fields(fx):
    import prior_states
    from sicopr_assist import core, feedback
    from sicopr_assist.project import Project
    old = os.environ.pop(ENV, None)
    os.environ[ENV] = 'on'
    try:
        ctx = prior_states.Ctx(fx, 'Case1')
        prior_states.seed_run(ctx)
        p = Project(ctx.root)
        said = ['Re-run with baud rate 53.125.', 'Give me more Rx FFE and re-run my setup.', 'yes',
                'Give me more Rx FFE and re-run my setup.', 'cancel',
                'Restrict the receiver equalization to low power.', 'short fixed FFE', 'Show settings.']
        outs = [core.handle(ph, p) for ph in said]
        hist = [r for r in p.read_history() if r.get('phrase') in said]
        recs = dict(zip(range(len(said)), hist[-len(said):]))
        t0 = recs[0].get('translation') or {}
        check('V10.5: a command in translator mode records every translation field (%s)' % ', '.join(TRANSLATION_FIELDS),
              all(k in t0 for k in TRANSLATION_FIELDS), sorted(set(TRANSLATION_FIELDS) - set(t0)))
        check("V10.5: the user's wording", t0.get('wording') == said[0], t0.get('wording'))
        check('V10.5: whether the strict core missed it', (t0.get('strict') or {}).get('missed') is True, t0.get('strict'))
        check('V10.5: the rule, synonym or preset that fired', 'synonym: baud rate is f_b' in (t0.get('fired') or []),
              t0.get('fired'))
        check('V10.5: the phrase handed to the core', t0.get('phrase') == 'Re-run with f_b 53.125.' and
              t0.get('used') is True, t0.get('phrase'))
        check("V10.5: the translator's name, version and rules SHA-256",
              t0.get('translator') == 'rules and presets' and t0.get('version') and
              re.fullmatch(r'[0-9a-f]{64}', t0.get('rules_sha256') or ''), (t0.get('translator'), t0.get('version')))
        t1 = recs[1].get('translation') or {}
        check('V10.5: a proposal is recorded with its values', 'ffe_pre_tap_len 6 -> 7' in (t1.get('proposal') or ''),
              t1.get('proposal'))
        check('V10.5: the proposal confirmed (its first option chosen)',
              (recs[2].get('translation') or {}).get('response') == 'confirmed', recs[2].get('translation'))
        check('V10.5: the proposal declined (cancel)',
              (recs[4].get('translation') or {}).get('response') == 'declined', recs[4].get('translation'))
        check('V10.5: the proposal changed (another option than the proposal)',
              (recs[6].get('translation') or {}).get('response') == 'changed', recs[6].get('translation'))
        t7 = recs[7].get('translation') or {}
        check('V10.5: a phrase the strict core reads is recorded too, with nothing fired and not used',
              t7.get('wording') == said[7] and t7.get('used') is False and (t7.get('strict') or {}).get('missed') is False,
              t7)
        lines = feedback.export_lines(p)
        text = '\n'.join(lines)
        sec = text[text.find('translator mode ('):]
        check('V10.5: Export feedback has a translator-mode section with every command made in translator mode',
              'translator mode (%d commands)' % len([r for r in p.read_history() if r.get('translation')]) in text,
              text[text.find('translator mode'):][:200])
        for what, needle in (('the wording', 'wording: Re-run with baud rate'), ('strict missed', 'strict missed: yes'),
                             ('what fired', 'fired: synonym: baud rate is f_b'),
                             ('the phrase handed to the core', 'handed to the core: Re-run with f_b 53.125.'),
                             ('the proposal', 'proposal: ffe_pre_tap_len 6 -> 7'), ('the response', 'response: confirmed'),
                             ('declined', 'response: declined'), ('changed', 'response: changed'),
                             ('the translator', 'translator: rules and presets')):
            check('V10.5: Export feedback carries %s' % what, needle in sec, sec[:300])
        fx.remove_project(ctx.root)
        # strict mode records no translation
        os.environ[ENV] = 'off'
        ctx = prior_states.Ctx(fx, 'Case1')
        prior_states.seed_run(ctx)
        p = Project(ctx.root)
        core.handle('Re-run with baud rate 53.125.', p)
        last = p.read_history()[-1]
        check('V10.5: in strict mode a command records no translation', last.get('translation') is None,
              last.get('translation'))
        fx.remove_project(ctx.root)
    finally:
        if old is None:
            os.environ.pop(ENV, None)
        else:
            os.environ[ENV] = old


# ----------------------------------------------------------------------------- the miss gate and the layer's reach
def miss_gate(fx):
    import prior_states
    from sicopr_assist import core
    from sicopr_assist.project import Project
    old = os.environ.get(ENV)
    os.environ[ENV] = 'on'
    try:
        ctx = prior_states.Ctx(fx, 'Case1')
        prior_states.seed_run(ctx)
        p = Project(ctx.root)
        o = core.handle('set DER_0 to 2 and rerun', p)
        check('miss gate: an illegal value on a request the core understood is a STOP, not reinterpreted',
              o['type'] == 'STOP' and not (o.get('translation') or {}).get('used'), _brief(o))
        o = core.handle('email the results to the team', p)
        check('miss gate: an out-of-scope request (email) is never handed to the layer',
              o['type'] == 'STOP' and not (o.get('translation') or {}).get('fired'), _brief(o))
        o = core.handle('explain the baud rate', p)
        check('miss gate: an out-of-scope request the strict core misses (explain, with a word the layer knows) is '
              'never handed to the layer', o['type'] == 'STOP' and not (o.get('translation') or {}).get('fired'),
              (_brief(o), (o.get('translation') or {}).get('fired')))
        q = core.handle('Sweep DER_0', p)
        o = core.handle('set the baud rate to 53.125 and rerun', p)
        check('miss gate: while a question is pending, a reply that is not one of its answers never goes to the layer',
              q['type'] == 'QUESTION' and o['type'] == 'QUESTION' and not (o.get('translation') or {}).get('fired'),
              (q['type'], _brief(o), (o.get('translation') or {}).get('fired')))
        fx.remove_project(ctx.root)
    finally:
        if old is None:
            os.environ.pop(ENV, None)
        else:
            os.environ[ENV] = old
    with open(os.path.join(PKG, 'translate.py'), encoding='utf-8') as f:
        src = f.read()
    imports = re.findall(r'^\s*(?:from\s+(\S+)\s+import\s+([^\n]+)|import\s+([^\n]+))', src, re.M)
    names = ' '.join(' '.join(x) for x in imports)
    check('the layer never runs the engine or writes the run CSV: translate.py imports neither run, run_case, '
          'subprocess nor engine, and opens no file for writing',
          not re.search(r'\brun\b|run_case|subprocess|\bengine\b', names) and not re.search(r"open\([^)]*'[wa]", src),
          names)


# ----------------------------------------------------------------------------- the core fixes, beyond their rows
def core_fixes(fx):
    """What rows CF01, CF10 and CF12 do not compare by machine: Q14's text for a b_min limit above 0 (core fix 1), the
    run CSV holding a MATLAB vector expression written out (core fix 3), and the verdict of a run in eye height and
    VEC mode (core fix 5)."""
    import prior_states
    from sicopr_assist import config_model, core, summary, validate
    from sicopr_assist.project import Project
    wc = fx.working_config('Case1')
    r = validate.validate(wc, [('/main/b_min(1)', '0.3')])
    check('core fix 1: Q14 says that b_min(1) = 0.3 is a minimum above 0 that forces the first DFE tap up, not a '
          'weight outside -1 to 1', r['status'] == 'ask' and 'minimum above 0' in r['text'] and 'first DFE tap' in
          r['text'] and 'outside -1 to 1' not in r['text'], r['text'])
    r = validate.validate(wc, [('/main/b_max(1)', '0.5')])
    check('core fix 1: a b_max(1) of 0.5 is not asked about (the sign rule is for b_min above 0 and b_max below 0)',
          r['status'] == 'ok', r['text'])
    old = os.environ.get(ENV)
    os.environ[ENV] = 'off'
    try:
        ctx = prior_states.Ctx(fx, 'Case1')
        prior_states.seed_run(ctx)
        o = core.handle('Re-run with N_b = 12 and b_max(2..N_b) = [0.25 0.15*ones(1,3) 0.05*ones(1,7)].', Project(ctx.root))
        after = next((c['after'] for c in o['changes'] if c['json_path'] == '/main/b_max(2..N_b)'), None)
        check('core fix 3: a MATLAB vector expression is written out as the values MATLAB reads',
              o['type'] == 'PLAN' and after == '[0.25 0.15 0.15 0.15 0.05 0.05 0.05 0.05 0.05 0.05 0.05]', after)
        new = config_model.apply_changes(ctx.prior_wc, [(c['json_path'], c['after']) for c in o['changes']
                                                         if c['kind'] in ('set', 'add')])['config']
        fields, err = config_model.loaded_fields(new)
        check("core fix 3: SiCoPR's loader reads the run CSV of that change", err is None and fields is not None, err)
        fx.remove_project(ctx.root)
        ctx = prior_states.Ctx(fx, 'Case1')
        prior_states.c2m_eye_run(ctx)
        p = Project(ctx.root)
        s = summary.run_summary(p, p.read_history()[-1]['run_ids'][0])
        text = '\n'.join(s['lines'])
        check('core fix 5: a run in eye height and VEC mode is judged by VEC Pass threshold and EH_min, not COM: fail, '
              'the eye height below EH_min', 'fail: VEC_dB 8.6 is within VEC Pass threshold (10); the eye height '
              '(VEO_mV 9.1) is below EH_min (12)' in text and 'COM Pass threshold' not in text and 'not judged' in text,
              text)
        fx.remove_project(ctx.root)
        ctx = prior_states.Ctx(fx, 'Case1')
        prior_states.strict_ooxml(ctx)
        o = core.handle('Run COM on data\\thru.s4p with data\\config.xlsx.', Project(ctx.root))
        check('core fix 4: a strict OOXML workbook is a STOP that names it and says how to save it',
              o['type'] == 'STOP' and 'strict OOXML workbook' in o.get('reason', '') and
              'ordinary Excel Workbook' in o.get('reason', ''), _brief(o))
        fx.remove_project(ctx.root)
    finally:
        if old is None:
            os.environ.pop(ENV, None)
        else:
            os.environ[ENV] = old


def main():
    if not fixtures.find_starting_configs():
        print('SKIP: the four starting workbooks are absent or hidden; the translator checks need a working config')
        provenance()
        license_notice()
        template_ships()
        finish()
    template_ships()
    license_notice()
    with fixtures.Fixtures(variants=False) as fx:
        provenance(fx)
        settings_and_mode(fx)
        presets_fresh_install(fx)
        history_fields(fx)
        miss_gate(fx)
        core_fixes(fx)
    finish()


if __name__ == '__main__':
    main()
