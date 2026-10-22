#!/usr/bin/env python3
"""gen_use_cases.py - regenerate the parameter, enumeration and sweep use cases for any COM config sheet.

Reference implementation of the scoring oracle for SiCoPR-Assist. It holds no expected results
for particular sheets. It derives them from each config it is given:

  1. discover   find parameters the way the loader does: by keyword, anywhere in the sheet,
                value in the cell to the right. No cell addresses or column layout are assumed.
  2. import     build the config JSON (main, packages, links). Stored values are kept exactly;
                a formula becomes a link by parameter name.
  3. probe      apply one change in the JSON model, load it with sicopr.read_ParamConfigFile,
                and record which loaded fields moved.
  4. classify   turn the probe into PLAN, QUESTION or STOP with the rules in classify().
                Sweeps (sweep_rows) do the same for every point and count which points are
                new runs and which leave the loaded parameters as they were.

Usage (same in PowerShell):
  python gen_use_cases.py --sicopr sicopr.py --out rows.csv Case1=a.xlsx Case2=b.xlsx
  python gen_use_cases.py --sicopr sicopr.py --check someone_elses.xlsx
  python gen_use_cases.py --sicopr sicopr.py --compare reference.xlsx someone_elses.xlsx
  python gen_use_cases.py --variants a.xlsx variants_dir

The first config listed fixes the phrase ids; configs listed after it reuse those ids and
append new ones. --variants writes twelve altered copies of a sheet (a parameter removed,
rows moved, blocks renamed and so on); it expects the 802.3dj sheet layout it was written
against. Needs openpyxl, plus whatever sicopr.py imports. Nothing here runs COM.
"""
# Copyright 2026 Todd Bermensolo
# SPDX-License-Identifier: BSD-3-Clause
import argparse, collections, contextlib, copy, csv, difflib, hashlib, importlib.util, io, json, os, re, sys, tempfile
from types import SimpleNamespace
import openpyxl

REF = re.compile(r'\$?([A-Z]{1,3})\$?(\d+)')
SHEET = 'COM_Settings'
COLS = ['id', 'group', 'priority', 'starting_config', 'prior_state', 'phrase', 'expected_type', 'expected_detail',
        'fallback', 'paraphrases', 'open_decision', 'applies_when', 'target_param', 'target_block', 'json_path',
        'value_before', 'value_after', 'linked_changes', 'loader_effect', 'verified', 'note']
PRIOR = 'run complete with the starting config and data\\thru.s4p'

# ----------------------------------------------------------------------------- sicopr access
def load_sicopr(path):
    spec = importlib.util.spec_from_file_location('sicopr_under_test', path)
    mod = importlib.util.module_from_spec(spec)
    with contextlib.redirect_stdout(io.StringIO()):
        spec.loader.exec_module(mod)
    return mod

def known_keywords(sicopr_path):
    """Keywords the loader looks up, read from its own source so the list cannot drift."""
    lines = open(sicopr_path, encoding='utf-8').read().split('\n')
    def body(name):
        a = next(i for i, l in enumerate(lines) if l.startswith('def %s(' % name))
        b = next(i for i in range(a + 1, len(lines)) if lines[i].startswith('def '))
        return '\n'.join(lines[a:b])
    main = [k for _, k in re.findall(r"_read_ParamConfigFile__xls_param(?:_txffe)?\(\s*parameter\s*,\s*(['\"])(.+?)\1",
                                      body('read_ParamConfigFile'))]
    pkg = [k for _, k in re.findall(r"xp\(\s*(['\"])(.+?)\1", body('_read_ParamConfigFile__read_pkg_params'))]
    zp = ['z_p (TX)', 'z_p (NEXT)', 'z_p (FEXT)', 'z_p (RX)']          # looked up through a variable in the loader
    # SiCoPR e1651c4 reads b_max(2..N_b) and b_min(2..N_b) through a helper, _dfe_limits(key, default)
    main += [k for _, k in re.findall(r"_dfe_limits\(\s*(['\"])(.+?)\1", body('read_ParamConfigFile'))]
    src = body('read_ParamConfigFile'); mandatory = set()
    for m in re.finditer(r"_read_ParamConfigFile__xls_param\(\s*parameter\s*,\s*(['\"])(.+?)\1", src):
        i = m.end(); depth = 1; commas = 0                         # count the arguments after the keyword
        while depth and i < len(src):
            ch = src[i]
            if ch in '([': depth += 1
            elif ch in ')]': depth -= 1
            elif ch == ',' and depth == 1: commas += 1
            i += 1
        if commas <= 1: mandatory.add(m.group(2).lower())          # no default given
    MANDATORY.clear(); MANDATORY.update(mandatory)
    return {k.lower() for k in main + zp}, {k.lower() for k in pkg + zp}

MANDATORY = set()

def is_main_keyword(label, KW):
    k = label.strip().lower()
    return k in KW[0] or re.fullmatch(r'c\(-?\d+\)', k) is not None      # c(-N), c(N) are probed in a loop

def bootstrap_OP():
    """Mirrors the OP bootstrap at the top of sicopr._run_com."""
    return SimpleNamespace(DEBUG=False, DISPLAY_WINDOW=False, TDMODE=False, GET_FD=True, RxFFE=False,
                           RxFFE_with_MMSE=False, FFE_OPT_METHOD='FOM', RX_CALIBRATION=False, PSDRXCAL=False,
                           ERL_ONLY=False, WC_PORTZ=False, SNDR_REF=False, RESULT_DIR='', CONFIG2MAT_ONLY=False,
                           TESTING=0)

# ----------------------------------------------------------------------------- discovery and import
def _col(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26); s = chr(65 + r) + s
    return s

def discover(path, KW, label=None):
    """xlsx -> config dict {main, packages, links} plus bookkeeping (cells, issues)."""
    wb = openpyxl.load_workbook(path); wv = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    wsv = wv[ws.title]
    raw = {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
    cached = {c.coordinate: c.value for row in wsv.iter_rows() for c in row if c.value is not None}
    starts = sorted(int(c[1:]) for c, v in raw.items() if re.fullmatch(r'A\d+', c) and v == '.START')
    ends = sorted(int(c[1:]) for c, v in raw.items() if re.fullmatch(r'A\d+', c) and v == '.END')
    cfg = dict(schema='sicopr-config/1', label=label or os.path.basename(path),
               source=dict(file=os.path.basename(path), sheet=ws.title,
                           sha256=hashlib.sha256(open(path, 'rb').read()).hexdigest()),
               main=collections.OrderedDict(), packages=collections.OrderedDict(), links={},
               cells={}, issues=[], sheets=wb.sheetnames)
    if len(starts) != len(ends):
        cfg['issues'].append('number of .START and .END rows differs'); return cfg
    first_start = starts[0] if starts else 10 ** 9
    lost = []                                                       # formula cells with no stored result
    def value_of(cell):
        v = raw.get(cell)
        if isinstance(v, str) and v.startswith('='):
            f = v; v = cached.get(cell)
            if v is None:
                if not REF.search(f) and re.fullmatch(r'[0-9eE\.\+\-\*/\(\) ]+', f[1:]): v = eval(f[1:])   # a constant such as =0.005
                else: lost.append(cell)
        if isinstance(v, float) and v.is_integer() and abs(v) < 2 ** 53: v = int(v)
        return v
    def cells_in(r0, r1):
        out = []
        for c, v in raw.items():
            m = re.fullmatch(r'([A-Z]+)(\d+)', c); r = int(m.group(2))
            if r0 <= r <= r1 and isinstance(v, str): out.append((r, openpyxl.utils.column_index_from_string(m.group(1)), c, v))
        return sorted(out)
    valcell = {}                                                    # value cell -> qualified key
    def register(scope_key, target, name, r, ci):
        vc = '%s%d' % (_col(ci + 1), r)
        if name in target:
            cfg['issues'].append('keyword "%s" occurs more than once (%s and %s): the loader raises on this'
                                 % (name, cfg['cells'][scope_key + name], vc)); return
        target[name] = value_of(vc); cfg['cells'][scope_key + name] = vc; valcell[vc] = scope_key + name
    for r, ci, c, v in cells_in(1, first_start - 1):
        if is_main_keyword(v, KW): register('', cfg['main'], v.strip(), r, ci)
    for s, e in zip(starts, ends):
        blk = str(raw.get('B%d' % s)); cfg['packages'][blk] = collections.OrderedDict()
        for r, ci, c, v in cells_in(s + 1, e - 1):
            if v.strip().lower() in KW[1]: register('packages.%s.' % blk, cfg['packages'][blk], v.strip(), r, ci)
    # formulas -> links by parameter name; a precedent that is not a parameter is taken in by its label
    for vc, key in list(valcell.items()):
        f = raw.get(vc)
        if not (isinstance(f, str) and f.startswith('=') and REF.search(f)): continue
        scope_key = key.rsplit('.', 1)[0] + '.' if key.startswith('packages.') else ''
        ok = True
        def name_of(m):
            nonlocal ok
            pc = m.group(1) + m.group(2)
            if pc not in valcell:
                lab = raw.get('%s%s' % (_col(openpyxl.utils.column_index_from_string(m.group(1)) - 1), m.group(2)))
                if isinstance(lab, str) and lab.strip() and not lab.startswith('='):
                    r_ = int(m.group(2))
                    sk = next(('packages.%s.' % b for b, (s, e) in zip(cfg['packages'], zip(starts, ends)) if s < r_ < e), '')
                    tgt = cfg['main'] if not sk else cfg['packages'][sk.split('.')[1]]
                    if lab.strip() not in tgt:
                        tgt[lab.strip()] = value_of(pc); cfg['cells'][sk + lab.strip()] = pc
                    valcell[pc] = sk + lab.strip()
                else:
                    ok = False; return m.group(0)
            return '{' + valcell[pc] + '}'
        expr = REF.sub(name_of, f[1:])
        if ok: cfg['links'][key] = expr
        else: cfg['issues'].append('formula in %s (%s) depends on an unlabelled cell: imported as a value, not a link' % (vc, key))
    if lost:
        done = 0
        for _ in range(len(cfg['links']) + 1):
            moved = False
            for key, expr in cfg['links'].items():
                sc, nm = _scope(cfg, key)
                if sc.get(nm) is None:
                    try: sc[nm] = _eval_link(cfg, expr); moved = True; done += 1
                    except Exception: pass
            if not moved: break
        cfg['issues'].append('%d formula cells have no stored result (the file was last saved by a tool that does not calculate): '
                             '%d recomputed from their formulas on import; the loader itself cannot read this file' % (len(lost), done))
    def pos(key):
        m = re.fullmatch(r'([A-Z]+)(\d+)', cfg['cells'][key]); return (int(m.group(2)), openpyxl.utils.column_index_from_string(m.group(1)))
    cfg['main'] = collections.OrderedDict(sorted(cfg['main'].items(), key=lambda kv: pos(kv[0])))      # sheet order, row by row
    for b in cfg['packages']:
        cfg['packages'][b] = collections.OrderedDict(sorted(cfg['packages'][b].items(), key=lambda kv: pos('packages.%s.%s' % (b, kv[0]))))
    # labels that look like a keyword but are not one: the loader ignores them and uses its default
    allkw = sorted(KW[0])
    for r, ci, c, v in cells_in(1, first_start - 1):
        k = v.strip().lower()
        if is_main_keyword(v, KW) or len(k) < 4 or ' ' in k and len(k) > 30: continue
        if raw.get('%s%d' % (_col(ci + 1), r)) is None: continue
        near = difflib.get_close_matches(k, allkw, n=1, cutoff=0.8)
        if near: cfg['issues'].append('label "%s" in %s is not a keyword the loader reads; did you mean "%s"?' % (v.strip(), c, near[0]))
    return cfg

def selected_blocks(cfg):
    return str(cfg['main'].get('PKG_NAME') or '').split()

# ----------------------------------------------------------------------------- JSON model
def _scope(cfg, key):
    if key.startswith('packages.'):
        _, blk, name = key.split('.', 2); return cfg['packages'][blk], name
    return cfg['main'], key

def _eval_link(cfg, expr):
    m = re.fullmatch(r'\{([^}]+)\}', expr)
    def get(k):
        s, n = _scope(cfg, k); return s[n]
    if m: return get(m.group(1))
    def sub(mm):
        x = get(mm.group(1))
        if not isinstance(x, (int, float)) or isinstance(x, bool): raise ValueError('%s is not a number' % mm.group(1))
        return repr(float(x))
    py = re.sub(r'\{([^}]+)\}', sub, expr)
    if not re.fullmatch(r'[0-9eE\.\+\-\*/\(\) ]+', py): raise ValueError('unsupported link ' + expr)
    return eval(py)                                                                      # digits and operators only

def _deps(cfg, key):
    out, todo = set(), [key]
    while todo:
        for n in re.findall(r'\{([^}]+)\}', cfg['links'].get(todo.pop(), '')):
            if n not in out: out.add(n); todo.append(n)
    return out

def apply(cfg, overrides):
    """An explicit value removes that key's link; links downstream of a change are recomputed."""
    new = copy.deepcopy(cfg); changed = {}
    # a value equal to the current one is not a change: it must not drop a link or trigger a recompute
    overrides = {k: v for k, v in overrides.items() if not (_scope(cfg, k)[1] in _scope(cfg, k)[0] and norm(v) == norm(_scope(cfg, k)[0][_scope(cfg, k)[1]]))}
    for k, v in overrides.items():
        s, n = _scope(new, k); s[n] = v; new['links'].pop(k, None)
    for _ in range(len(new['links']) + 1):
        moved = False
        for k in new['links']:
            if _deps(new, k) & set(overrides):
                s, n = _scope(new, k); val = _eval_link(new, new['links'][k])
                if s[n] != val:
                    changed.setdefault(k, [_scope(cfg, k)[0].get(n), None])[1] = val; s[n] = val; moved = True
        if not moved: break
    return new, {k: tuple(v) for k, v in changed.items()}

def _csv_cell(v):
    if v is None or (isinstance(v, float) and v != v): return ''
    if isinstance(v, bool): return '1' if v else '0'
    if isinstance(v, float): return repr(v)
    return str(v)

def to_csv(cfg, path):
    """The run CSV SiCoPR-Assist hands the engine (prompt v3, Decision 2): keyword,value rows, then a
    .START,<block> ... .END block per package. The Reference Code reads the same file with xlsread."""
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        for k, v in cfg['main'].items(): w.writerow([k, _csv_cell(v)])
        for blk, params in cfg['packages'].items():
            w.writerow(['.START', blk])
            for k, v in params.items(): w.writerow([k, _csv_cell(v)])
            w.writerow(['.END'])

# ----------------------------------------------------------------------------- loading and diffing
def _num(x):
    try:
        x = float(x)
        return str(int(x)) if x == int(x) and abs(x) < 1e9 else '%.6g' % x
    except Exception:
        return str(x)

def fmt(v):
    import numpy as np
    if isinstance(v, np.ndarray):
        if v.ndim == 2 and v.size <= 24: return '[' + ' ; '.join(' '.join(_num(x) for x in row) for row in v) + ']'
        if v.size <= 8: return '[' + ' '.join(_num(x) for x in v.ravel()) + ']'
        f = v.ravel()
        try: return '%s values %s..%s' % ('x'.join(str(n) for n in v.shape), _num(f[0]), _num(f[-1]))
        except Exception: return 'array %s' % (v.shape,)
    if isinstance(v, (float, int)) and not isinstance(v, bool): return _num(v)
    if isinstance(v, (list, tuple)):
        if len(v) <= 8: return '[' + ' '.join(fmt(x) for x in v) + ']'
        return '%d values %s..%s' % (len(v), fmt(v[0]), fmt(v[-1]))
    return str(v)

def _flatten(ns, prefix, out):
    for k, v in vars(ns).items():
        if isinstance(v, SimpleNamespace): _flatten(v, prefix + k + '.', out)
        elif isinstance(v, dict):
            for kk, vv in v.items():
                if isinstance(vv, SimpleNamespace): _flatten(vv, prefix + k + '.' + str(kk) + '.', out)
                else: out[prefix + k + '.' + str(kk)] = vv
        else: out[prefix + k] = v
    return out

def load_fields(S, path=None, cfg=None):
    """Loaded parameters as {field: display text}. Give an xlsx path, or a config dict, which is written
    as the run CSV and read by the engine's own CSV reader (no reader is patched in)."""
    tmp = None
    if cfg is not None:
        fd, tmp = tempfile.mkstemp(suffix='.csv'); os.close(fd)
        to_csv(cfg, tmp)
        path = tmp
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            try: p, OP = S.read_ParamConfigFile(path, bootstrap_OP())
            except Exception as e: return None, '%s: %s' % (type(e).__name__, e)
    finally:
        if tmp and os.path.exists(tmp): os.remove(tmp)
        mat = (tmp or '')[:-4] + '.mat'
        if tmp and os.path.exists(mat): os.remove(mat)
    out = {}; _flatten(p, 'param.', out); _flatten(OP, 'OP.', out)
    return out, None

def _same(a, b):
    import numpy as np
    try:
        if isinstance(a, np.ndarray) or isinstance(b, np.ndarray):
            a = np.asarray(a); b = np.asarray(b)
            if a.shape != b.shape: return False
            num = a.dtype.kind in 'fiu' and b.dtype.kind in 'fiu'
            return bool(np.array_equal(a, b, equal_nan=True) if num else np.array_equal(a, b))
        if isinstance(a, float) and isinstance(b, float) and a != a and b != b: return True
        if isinstance(a, (int, float)) and isinstance(b, (int, float)): return a == b
        return type(a) == type(b) and a == b
    except Exception:
        return repr(a) == repr(b)

def diff(a, b):
    d = []
    for k in sorted(set(a) | set(b)):
        if k not in b: d.append('%s: %s -> (absent)' % (k, fmt(a[k])))
        elif k not in a: d.append('%s: (absent) -> %s' % (k, fmt(b[k])))
        elif not _same(a[k], b[k]): d.append('%s: %s -> %s' % (k, fmt(a[k]), fmt(b[k])))
    return d

# ----------------------------------------------------------------------------- test values
TEST_VALUES = {  # value to request per parameter; anything not listed gets an automatic numeric change
 'SAVE_CONFIG2MAT': 1, 'DIAGNOSTICS': 0, 'ERL Pass threshold': 10, 'f_b': 53.125, 'DISPLAY_WINDOW': 1, 'COM Pass threshold': 4,
 'RX_CALIBRATION': 1, 'f_min': 0.01, 'CSV_REPORT': 0, 'DER_0': 1e-4, 'Sigma BBN step': 0.002, 'Delta_f': 0.005,
 'RESULT_DIR': '.\\results\\trial_{date}\\', 'T_r': 0.005, 'C_d': '[0.5e-4 0.9e-4 1.1e-4 ; 0.5e-4 0.9e-4 1.1e-4]',
 'SAVE_FIGURES': 1, 'FORCE_TR': 0, 'T_t': 5, 'L_s': '[0.12 0.15 0.14 ; 0.12 0.15 0.14]', 'Port Order': '[1 2 3 4]',
 'PMD_type': 'C2M', 'f_v': 0.4, 'C_b': '[0.2e-4 0.2e-4]', 'RUNTAG': 'KR_pkgB_', 'EW': 0, 'T_ft': 5, 'R_0': 50,
 'COM_CONTRIBUTION': 0, 'MLSE': 0, 'T_nt': 5, 'PKG_NAME': 'PKG_HiR_CLASSB PKG_HiR_CLASSB', 'ts_anchor': 0, 'f_f': 0.6,
 'z_p select': '[ 2 ]', 'sample_adjustment': '[-16 16]', 'f_n': 0.6, 'L': 2, 'TDR': 0, 'Local Search': 0, 'f_1': 0.05,
 'M': 24, 'ERL': 0, 'flim': 5e10, 'f_2': 50, 'ERL_ONLY': 1, 'zero_pad': 0, 'A_ft': 0.5, 'f_r': 0.75, 'TR_TDR': 0.008,
 'A_nt': 0.5, 'c(0)': 0.6, 'N': 3500, 'ffe_pre_tap_len': 5, 'c(-1)': '[-0.3:.02:0]', 'TDR_Butterworth': 0,
 'ffe_post_tap_len': 12, 'c(-2)': '[0:.02:0.1]', 'beta_x': 1.7e9, 'ffe_pre_tap1_max': 0.5,
 'board_tl_gamma0_a1_a2': '[0 3.8206e-04 9.5909e-05]', 'c(-3)': '[-0.06:.02:0]', 'rho_x': 0.25, 'ffe_post_tap1_max': 0.5,
 'board_tl_tau': 0.006, 'c(-4)': '[0:.01:0.02]', 'TDR_W_TXPKG': 1, 'ffe_tapn_max': 0.5, 'board_Z_c': 100,
 'c(1)': '[-0.1:.02:0]', 'N_bx': 24, 'num_ui_RXFF_noise': 2048, 'z_bp (TX)': 12, 'N_b': 2,
 'fixture delay time': '[ 0.1 0.1 ]', 'z_bp (NEXT)': 12, 'b_max(1)': 0.75, 'Tukey_Window': 0, 'N_bg': 3, 'z_bp (FEXT)': 12,
 'b_max(2..N_b)': 0.2, 'Z_t': 50, 'N_bf': 6, 'z_bp (RX)': 12, 'b_min(1)': -0.1, 'N_f': 60, 'C_0': '[0.29e-4 0.29e-4]',
 'b_min(2..N_b)': -0.15, 'sigma_RJ': 0.02, 'bmaxg': 0.1, 'C_1': '[0.19e-4 0.19e-4]', 'g_DC': '[-16:1:0]', 'A_DD': 0.025,
 'N_tail_start': 12, 'Include PCB': 1, 'f_z': 40, 'eta_0': 1e-8, 'TS_SRCH_MODE': 'middle', 'f_p1': 40, 'SNR_TX': 32.5,
 'Non-zero Local Search Method': 0, 'f_p2': 100, 'R_LM': 0.97, 'Clip Method': 'Fast', 'g_DC_HP': '[-4:1:0]', 'N_qb': 7,
 'P_peak': 1e-6, 'f_HP_PZ': 1.5, 'P_qc': 1e-6, 'c(2)': '[-0.1:.02:0]',
 # package parameters
 'package_tl_gamma0_a1_a2': '[0 0.0009909 0.0002772]', 'package_tl_tau': 0.0065,
 'package_Z_c': '[90 90 ; 95 95 ; 100 100 ; 100 100]', 'z_p (TX)': '[12 31 31 31 ; 1.8 1.8 1.8 1.8 ; 0 0 0 0 ; 0 0 0 0]',
 'z_p (NEXT)': '[12 31 31 31 ; 1.8 1.8 1.8 1.8 ; 0 0 0 0 ; 0 0 0 0]', 'z_p (FEXT)': '[12 31 31 31 ; 1.8 1.8 1.8 1.8 ; 0 0 0 0 ; 0 0 0 0]',
 'z_p (RX)': '[12 31 31 31 ; 1.8 1.8 1.8 1.8 ; 0 0 0 0 ; 0 0 0 0]', 'C_p': '[0.3e-4 0.3e-4]', 'A_v': 0.4, 'A_fe': 0.4, 'A_ne': 0.6,
}
SAY = {'DER_0': '1e-4', 'flim': '5e10', 'beta_x': '1.7e9', 'eta_0': '1e-8', 'P_peak': '1e-6', 'P_qc': '1e-6', 'z_p select': '2'}

def test_value(name, current):
    if name == 'R_d': return '[45 45]' if '50' in str(current) else '[50 50]'
    if name in TEST_VALUES: return TEST_VALUES[name]
    if isinstance(current, int) and not isinstance(current, bool): return current + 1
    if isinstance(current, float): return round(current * 1.1, 6) if current else 1.0
    return None                                                     # text or matrix with no test value: skipped

def val(v):
    if v is None: return ''
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        f = float(v)
        if f == int(f) and abs(f) < 1e6: return str(int(f))
        return ('%.10g' % f).replace('e-0', 'e-').replace('e+', 'e')
    return re.sub(r'\s+', ' ', str(v)).strip()

def norm(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool): return repr(float(v))
    s = re.sub(r'\s+', ' ', str(v)).strip(); s = re.sub(r'\[\s*', '[', s); s = re.sub(r'\s*\]', ']', s)
    try: return repr(float(s))
    except ValueError: return s

# ----------------------------------------------------------------------------- classification rules
NOTES = {
 'COM Pass threshold': 'add_rx_noise defaults to the pass threshold (MATLAB 4p15p0 L9955, same in the port), so this can move results under MLSE, not only the verdict; announce it',
 'RX_CALIBRATION': 'switches to receiver calibration; with MMSE Rx FFE the loader maps it to PSDRXCAL; announce that this is a different analysis',
 'ERL_ONLY': 'ERL-only run: announce that COM is not computed', 'RESULT_DIR': 'the loader expands {date}',
 'L': '2 is NRZ, 4 is PAM4', 'flim': 'stored in Hz', 'RUNTAG': 'prefix on result file names',
 'Include PCB': 'adds the board model on each side, so z_bp, board_tl_gamma0_a1_a2, board_tl_tau, board_Z_c, C_0 and C_1 start to matter; announce it',
}
DECIDED = {   # policies settled on 3 Oct 2026; the text goes into the note of the first row that shows each one
    'pairboth': 'decided: a range sweep on a [TX RX] pair whose two values are equal writes each point to both and announces it',
    'search': 'decided: a range sweep on a parameter that is already a search range asks, one run or a batch',
    'unselected': 'decided: an edit to an unselected package block asks',
    'noeffect': 'decided: a change with no effect on engine inputs stops, and runs anyway on request',
    'same': 'decided: a request for the value already set stops and says so',
    'linked': 'decided: an explicit value on a linked parameter is applied, the link is removed in the working config and that is announced',
    'add': 'decided: a known keyword missing from the config is added, and the loader default that was in use is announced'}
BOARD = {'z_bp (TX)', 'z_bp (NEXT)', 'z_bp (FEXT)', 'z_bp (RX)', 'board_tl_gamma0_a1_a2', 'board_tl_tau', 'board_Z_c', 'C_0', 'C_1'}
KW_MAIN = set()

def hint(cfg, name):
    """Notes that depend on the state of this config, not on any particular sheet."""
    if name in NOTES: return NOTES[name]
    if name in BOARD and cfg['main'].get('Include PCB') in (0, None): return 'recorded by the loader but unused while Include PCB = 0'
    if KW_MAIN and name.lower() not in KW_MAIN and not re.fullmatch(r'c\(-?\d+\)', name.lower()):
        users = sorted(k for k, e in cfg['links'].items() if '{%s}' % name in e)
        return 'not a loader keyword; reaches the engine only through the %s link' % ', '.join(users) if users else 'not a loader keyword'
    return ''

def WHY_NOT(cfg, name):
    if name in ('b_max(2..N_b)', 'b_min(2..N_b)') and cfg['main'].get('N_b') == 1: return 'unused while N_b = 1'
    if name == 'COM_CONTRIBUTION' and cfg['main'].get('MLSE'): return 'the loader forces the contribution plot off while MLSE = 1'
    return ''

FB_RUN = 'if validate_config fails, STOP with the loader message and keep the last good config'

def linktxt(cfg, key):
    return re.sub(r'\{(?:packages\.[^.}]+\.)?([^}]+)\}', r'\1', cfg['links'].get(key, ''))

def jptr(key):
    return '/' + key.replace('.', '/', 2) if key.startswith('packages.') else '/main/' + key

def find_key(scope, name):
    """The key as this config spells it. The loader matches keywords ignoring case and surrounding spaces."""
    if name in scope: return name
    k = name.strip().lower()
    return next((x for x in scope if x.strip().lower() == k), None)

def classify(S, cfg, base, spec, new, after_txt=None):
    """spec: ('main', name) | ('pkg', name) | ('named', block, name). Returns the expected outcome for this config."""
    sel = selected_blocks(cfg); blocks = list(cfg['packages'])
    o = dict(target_param=spec[-1], target_block='', json_path='', value_before='', value_after=after_txt or val(new),
             linked_changes='', loader_effect='', verified='', note='', fallback='', kind='plan')
    name = spec[-1]
    if spec[0] == 'named':
        if spec[1] not in blocks:
            o.update(expected_type='QUESTION', kind='noblock', target_block=spec[1], fallback='cancel',
                     expected_detail='no package block named %s in this config: options %s' % (spec[1], ', '.join(blocks) or 'none defined'))
            return o
        name = find_key(cfg['packages'][spec[1]], name) or name
        key = 'packages.%s.%s' % (spec[1], name); blk = spec[1]
    elif spec[0] == 'pkg' and blocks:
        if len(set(sel)) > 1:
            o.update(expected_type='QUESTION', kind='sides', fallback='cancel',
                     expected_detail='TX and RX use different package blocks: options %s (TX), %s (RX), both' % (sel[0], sel[-1]))
            return o
        blk = sel[0] if sel else blocks[0]
        name = find_key(cfg['packages'][blk], name) or name; key = 'packages.%s.%s' % (blk, name)
    else:
        name = find_key(cfg['main'], name) or name; key = name; blk = 'main'
    o['target_param'] = name
    # (c) a package selection must name blocks this config defines
    if name.strip().lower() == 'pkg_name':
        unknown = [t for t in str(new).split() if t not in blocks]
        if unknown and not blocks:
            o.update(expected_type='STOP', kind='noblock', target_block='main', json_path='/main/PKG_NAME', fallback='none', value_before='(absent)',
                     expected_detail='this config defines no package blocks, so a package cannot be selected by name; say so and list the package parameters in the main area')
            return o
        if unknown:
            order = blocks
            o.update(expected_type='QUESTION', kind='noblock', target_block='main', json_path='/main/PKG_NAME', fallback='cancel',
                     value_before=val(cfg['main'].get(name)), expected_detail='no package block named %s in this config: options %s' % (unknown[0], ', '.join(order)))
            return o
    scope, _ = _scope(cfg, key)
    present = name in scope
    cur = scope.get(name)
    lt = linktxt(cfg, key)
    bef = (val(cur) + (' (linked to %s)' % lt if lt else '')) if present else '(absent)'
    o.update(target_block=blk, json_path=jptr(key), value_before=bef,
             note=hint(cfg, name) if blk == 'main' else '', fallback=FB_RUN)
    if present and norm(new) == norm(cur):
        o.update(expected_type='STOP', kind='same', fallback='offer to run anyway on request', loader_effect='no loaded field changes', verified='loader',
                 expected_detail='%s is already %s%s: report that nothing changed and results equal the last run' % (name, val(cur), '' if blk == 'main' else ' in ' + blk))
        return o
    newcfg, ch = apply(cfg, {key: new})
    fields, err = load_fields(S, cfg=newcfg)
    lk = '; '.join('%s %s -> %s' % (k.rsplit('.', 1)[-1], val(b), val(a)) for k, (b, a) in ch.items())
    eff = diff(base, fields) if fields else None
    o.update(linked_changes=lk, loader_effect=('loader error: ' + err) if err else ' | '.join(eff), verified='loader-error' if err else 'loader')
    verb = 'set_param' if present else 'add_param'
    setp = '%s(%s: %s -> %s)' % (verb, name, bef, o['value_after']); rl = ' > recompute_linked(%s)' % lk if lk else ''
    if err:
        o.update(expected_type='STOP', kind='error', expected_detail='%s%s > validate_config fails: report the loader message and keep the last good config' % (setp, rl))
    elif blk != 'main' and blk not in sel and sel:
        o.update(expected_type='QUESTION', kind='unselected', fallback='cancel: the working config is unchanged',
                 expected_detail='state that %s is not the package selected by PKG_NAME, so this edit alone leaves the run inputs unchanged; options: also switch TX and RX to %s and run, save the edit and run with the current package, cancel. Edit itself: %s%s' % (blk, blk, setp, rl))
    elif not eff:
        o.update(expected_type='STOP', kind='noeffect', fallback='offer to run anyway on request', loader_effect='',
                 note='no effect because ' + (WHY_NOT(cfg, name) or 'no loaded value changes in this config'), expected_detail='%s > validate_config finds no engine input changed; report that results would equal the last run and do not run' % setp)
    else:
        extra = '; announce old and new value'
        if not present:
            extra = '; announce that %s was not in this config, so the loader default was in use, and that an entry was added' % name; o['kind'] = 'add'
        if spec[0] == 'pkg' and blocks:
            nb = sum(name in b for b in cfg['packages'].values())
            extra += '; the name occurs in %s package blocks and resolves to %s because PKG_NAME selects it for TX and RX' % ({2: 'two', 3: 'three', 4: 'four'}.get(nb, nb), blk)
        if lt: extra += '; the parameter was linked (%s); an explicit value removes the link in the working config, so say so' % lt; o['kind'] = 'linked'
        if name == 'MLSE' and 'COM_CONTRIBUTION_CURVES' in o['loader_effect']: o['note'] = 'turning MLSE off also re-enables the COM contribution plot, because COM_CONTRIBUTION = 1 in this config'
        o.update(expected_type='PLAN', expected_detail=setp + rl + ' > validate_config > run_com > summarize' + extra)
    return o

# ----------------------------------------------------------------------------- sweeps
SWITCHES = {'save_config2mat', 'diagnostics', 'display_window', 'rx_calibration', 'csv_report', 'save_figures', 'force_tr', 'ew',
            'com_contribution', 'mlse', 'tdr', 'erl', 'erl_only', 'zero_pad', 'tdr_butterworth', 'tdr_w_txpkg', 'tukey_window', 'non-zero local search method'}
PORT_ORDERS = ['1 3 2 4', '1 2 3 4', '3 1 2 4', '2 1 3 4', '1 3 4 2', '1 2 4 3', '3 1 4 2', '2 1 4 3']
FIXED_CHOICES = {'clip method': ['Slow', 'Fast'], 'ts_srch_mode': ['full-sweep', 'middle'], 'pmd_type': ['C2C', 'C2M'], 'ts_anchor': [0, 1, 2], 'include pcb': [0, 1, 2],
                 'l': [2, 4, 6, 8, 16]}
# These three tables mirror sicopr_param_schema.csv (kind, legal_values, shape); check_against_schema() reports any drift.
TXRX = {'c_p', 'c_d', 'l_s', 'c_b', 'fixture delay time', 'c_v', 'c_0', 'c_1', 'package_z_c', 'board_z_c', 'r_d'}   # the loader's make_length2 fields
SWEEP_LIMIT = 30                                                    # default of the max_new_runs setting: more new runs than this and the tool asks first

def check_against_schema(path):
    """Compare the built-in tables with the parameter schema file; returns a list of differences (empty when they agree)."""
    rows = list(csv.DictReader(open(path, newline='', encoding='utf-8'))); out = []
    main = {r['keyword'].lower(): r for r in rows if r['scope'] == 'main'}
    for k in sorted(SWITCHES):
        if main.get(k, {}).get('kind') != 'switch': out.append('%s: switch here, %s in the schema' % (k, main.get(k, {}).get('kind')))
    for k, v in FIXED_CHOICES.items():
        r = main.get(k, {})
        if r.get('kind') != 'choice' or [x.strip() for x in r.get('legal_values', '').split(',')] != [str(x) for x in v]:
            out.append('%s: %s here, %s (%s) in the schema' % (k, v, r.get('legal_values'), r.get('kind')))
    pair = {r['keyword'].lower() for r in rows if 'TX RX' in r['shape']}
    if pair != TXRX: out.append('[TX RX] keywords differ: %s' % sorted(pair ^ TXRX))
    counts = {r['keyword'].lower() for r in rows if r['whole_number'].startswith('yes')}
    return out

def numbers_in(s):
    try: return [float(x) for x in re.sub(r'[\[\];]', ' ', str(s)).split()]
    except ValueError: return None

def legal_values(cfg, name):
    """[(as spoken, as stored)] for a setting with a closed list of values, or None."""
    k = name.strip().lower()
    if k in SWITCHES: return [('0', 0), ('1', 1)]
    if k in FIXED_CHOICES: return [(str(v), v) for v in FIXED_CHOICES[k]]
    if k == 'port order': return [('[%s]' % o, '[ %s ]' % o) for o in PORT_ORDERS]
    if k == 'pkg_name': return [(b, '%s %s' % (b, b)) for b in cfg['packages']]
    if k == 'z_p select':
        sel = selected_blocks(cfg)
        src = cfg['packages'][sel[0]] if sel and sel[0] in cfg['packages'] else cfg['main']
        zk = find_key(src, 'z_p (TX)')
        if zk is None: return None
        ncases = len(str(src[zk]).strip(' []').split(';')[0].split())
        return [(str(i), '[ %d ]' % i) for i in range(1, ncases + 1)]
    return None

def kind_of(S, cfg, name, value):
    k = name.strip().lower()
    if k in SWITCHES: return 'switch'
    if legal_values(cfg, name) is not None: return 'choice'
    if (re.fullmatch(r'c\(-?\d+\)', k) and k != 'c(0)') or k in ('g_dc', 'g_dc_hp'): return 'search'
    if isinstance(value, (int, float)) and not isinstance(value, bool): return 'count' if k in getattr(S, '_COUNT_KEYWORDS', ()) else 'number'
    s = str(value).strip()
    if s.startswith('['):
        nums = numbers_in(s)
        if nums is None or ':' in s: return 'vector'
        if ';' in s.strip('[] ;'): return 'matrix'
        if len(nums) == 1: return 'number'
        return 'pair' if k in TXRX and len(nums) == 2 else 'vector'
    return 'text'

def sweep_points(a, b, step):
    """start:step:end as MATLAB counts it: the last point is the last one not past the end."""
    if step == 0: return None
    if (b - a) * step < 0: step = -step                              # a descending sweep
    n = int((b - a) / step + 1e-9) + 1
    return [float('%.12g' % (a + i * step)) for i in range(max(n, 0))]

# ----------------------------------------------------------------------------- behaviour limits
# Every limit has a built-in default and can be set in a text file. Later files win: built-in < user file < project file.
LIMITS = {   # name: (default, lowest legal value, type, meaning)
    'max_new_runs':         (SWEEP_LIMIT, 0, int,   'a batch with more new runs than this asks before it starts (0 = always ask)'),
    'tutor_until_commands': (5,           0, int,   'the tutor shows on its own while history holds fewer successful commands than this (0 = never)'),
    'max_download_mb':      (500,         0, int,   'a download larger than this asks before it starts'),
    'download_retries':     (1,           0, int,   'automatic retries of an interrupted download'),
    'name_match_cutoff':    (0.8,         0, float, 'similarity from 0 to 1 that a name needs to be offered as a near match'),
    'max_options':          (8,           1, int,   'most options listed in one question'),
    'trouble_misses':       (3,           0, int,   'this many misses within the trouble window bring up the tutor (0 = never)'),
    'trouble_window':       (5,           1, int,   'how many of the most recent commands are looked at for misses'),
    'feedback_notice_gap':  (20,          0, int,   'the line about Export feedback is shown at most once in this many commands'),
    'min_correlation_points': (3,         2, int,   'with fewer results than this a correlation request gives the pairs and the difference, with no coefficient'),
}

FEEDBACK_ADDRESS = 'sicopr_assist.showman405@simplelogin.com'   # where an exported feedback file is emailed; a forwarding alias, kept in this one place so it can be replaced
OPTIONS = {  # on/off settings, section [feedback]: name: (default, meaning)
    'mask_names':   (True,  'the feedback export replaces paths, file names and names that are not standard COM keywords with placeholders'),
    'mask_numbers': (False, 'the feedback export replaces numbers with <num>, keeping their units'),
    'feedback_notice': (True, 'when the tutor comes up because of trouble, add one line about Export feedback'),
}
_BOOL = {'on': True, 'yes': True, 'true': True, '1': True, 'off': False, 'no': False, 'false': False, '0': False}

def read_limits(paths):
    """paths: settings files in rising priority, e.g. [user file, project file]; missing files are skipped.
    Returns ({name: value}, {name: source}, [warnings]). A bad file, key or value never stops a run: the default stays in force and a warning says so."""
    import configparser
    vals = {k: v[0] for k, v in list(LIMITS.items()) + list(OPTIONS.items())}; src = {k: 'default' for k in vals}; warn = []
    for p in paths:
        if not os.path.exists(p): continue
        cp = configparser.ConfigParser(inline_comment_prefixes=('#', ';'))
        try:
            text = open(p, encoding='utf-8').read()
            if not re.search(r'^\s*\[', text, re.M): text = '[limits]\n' + text      # a bare list of name = value lines is accepted
            cp.read_string(text, source=p)
        except configparser.Error as e:
            warn.append('%s could not be read (%s): its settings are ignored' % (p, str(e).splitlines()[0])); continue
        for sec in cp.sections():
            if sec.lower() not in ('limits', 'feedback'): warn.append('%s: section [%s] is not known and was ignored' % (p, sec)); continue
            for k, raw in cp.items(sec):                                   # a setting is accepted under either heading
                if k not in LIMITS and k not in OPTIONS:
                    near = difflib.get_close_matches(k, list(LIMITS) + list(OPTIONS), 1, 0.5)
                    warn.append('%s: %s is not a setting and was ignored%s' % (p, k, '; did you mean %s' % near[0] if near else '')); continue
                if k in OPTIONS:
                    if raw.strip().lower() not in _BOOL:
                        warn.append('%s: %s = %s is not on or off, so %s stays in force' % (p, k, raw, 'on' if vals[k] else 'off')); continue
                    vals[k] = _BOOL[raw.strip().lower()]; src[k] = p; continue
                d, lo, typ, _ = LIMITS[k]
                try:
                    v = typ(raw)
                    if v < lo or (k == 'name_match_cutoff' and v > 1): raise ValueError
                except ValueError:
                    warn.append('%s: %s = %s is not a usable value, so %s stays in force' % (p, k, raw, vals[k])); continue
                vals[k] = v; src[k] = p
    return vals, src, warn

# ----------------------------------------------------------------------------- names: COM keywords only, no alias table
# A parameter is named by its COM keyword (case and surrounding spaces ignored, as the loader does). Any other wording
# asks, and the options come from the workbook itself: names that resemble the words, then keyword descriptions that
# contain them. Nothing is remembered from the answer.
_SKIP = {'the', 'a', 'an', 'of', 'to', 'and', 'in', 'for', 'is', 'as', 'on', 'all', 'every', 'with', 'without', 'at'}

def keyword_info(path):
    """{keyword: description} from the workbook's keywords sheet (any sheet whose name starts with 'keywords'); {} if there is none."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True); out = {}
    for name in wb.sheetnames:
        if not name.lower().startswith('keywords'): continue
        for row in wb[name].iter_rows(values_only=True):
            if len(row) >= 4 and isinstance(row[1], str) and isinstance(row[3], str) and row[1].strip() != 'keyword':
                out.setdefault(row[1].strip(), row[3].strip())
    return out

def name_options(cfg, info, words, limit=None):
    """Options for wording that is not a COM keyword and does not resemble one (a mistyped name is handled as a near match instead).
    Only plain words are looked up; a token with an underscore, a digit or a bracket is a name attempt and is not searched.
    Order: keyword descriptions holding every word, then those holding some (most words first); sheet order within a tie."""
    limit = LIMITS['max_options'][0] if limit is None else limit
    names = list(dict.fromkeys(list(cfg['main']) + [k for b in cfg['packages'].values() for k in b]))
    toks = [t.lower() for t in words.split() if re.fullmatch(r'[A-Za-z]+', t) and t.lower() not in _SKIP]
    forms = lambda t: {t, t[:-1] if t.endswith('s') else t + 's'}
    low = {k.lower(): v for k, v in info.items()}
    scored = []
    for i, n in enumerate(names):
        desc = set(re.findall(r'[a-z0-9]+', low.get(n.strip().lower(), '').lower()))
        hit = sum(1 for t in toks if forms(t) & desc)
        if hit: scored.append((0 if hit == len(toks) else 1, -hit, i, n))
    return [n for *_, n in sorted(scored)][:limit]

def steps_rule(a, b, x):
    """What 'from a to b in x steps' means. A whole number is a count of points that covers both ends,
    step = (b - a) / (x - 1); anything else is a step size counted as MATLAB start:step:end (may stop short).
    Returns ('count' | 'step', points) or ('ask', None) for a count below 2. 'in steps of x' and 'by x' are always a step size."""
    if float(x) == int(x):
        n = int(x)
        if n < 2: return 'ask', None
        return 'count', [float('%.12g' % (a + i * (b - a) / (n - 1))) for i in range(n)]
    return 'step', sweep_points(a, b, x)

def sweep_outcome(S, cfg, base, spec, points, stored, how):
    """Aggregate one batch: points are (spoken, stored value). Points that leave every loaded value alone reuse the last run."""
    per = [classify(S, cfg, base, spec, sv, None) for _, sv in points]
    first = per[0]
    o = dict(target_param=first['target_param'], target_block=first['target_block'], json_path=first['json_path'], value_before=first['value_before'],
             value_after='; '.join(sp for sp, _ in points), linked_changes='', loader_effect='', verified='loader', note='', fallback='', kind='sweep')
    if first['kind'] in ('sides', 'noblock', 'unselected'):
        o.update(expected_type=first['expected_type'], kind=first['kind'], fallback='cancel', expected_detail=first['expected_detail'].split('. Edit itself')[0]); return o
    new = [(p, r) for p, r in zip(points, per) if r['kind'] in ('plan', 'linked', 'add')]
    reuse = [p for p, r in zip(points, per) if r['kind'] in ('same', 'noeffect')]
    bad = [(p, r) for p, r in zip(points, per) if r['kind'] == 'error']
    o['loader_effect'] = ' || '.join('%s: %s' % (p[0], r['loader_effect'] or 'same as the last run') for p, r in zip(points, per))[:1500]
    name = o['target_param']
    if not new and not bad:
        why = 'the value given changes no engine input' if len(points) == 1 else 'none of the %d values changes an engine input' % len(points)
        o.update(expected_type='STOP', kind='noeffect', fallback='offer to run anyway on request',
                 expected_detail='%s: report that every result would equal the last run and do not run' % why); return o
    if not new:
        o.update(expected_type='STOP', kind='error', fallback='none', verified='loader-error',
                 expected_detail='no value in the sweep passes validate_config: report the loader message for each'); return o
    lk = sorted({x.split(' ')[0] for _, r in new for x in r['linked_changes'].split('; ') if x})
    def n_(k, word): return '%d %s%s' % (k, word, '' if k == 1 else 's')
    parts = ['build_matrix(%s = %s)' % (name, ', '.join(p[0] for p in points)), n_(len(new), 'working config'), 'validate each', 'run_batch', 'batch_summary with one row per value']
    if how == 'values': parts.append('diff_results against the last run')
    tail = '; announce %s: %s' % (n_(len(points), 'value'), n_(len(new), 'new run'))
    if reuse: tail += ', %d reused from the last run (%s)' % (len(reuse), ', '.join(p[0] for p in reuse))
    if bad: tail += ', %d skipped because validate_config fails (%s)' % (len(bad), ', '.join(p[0] for p, _ in bad))
    if lk: tail += '; linked parameters follow at each point (%s)' % ', '.join(lk)
    if any(r['kind'] == 'linked' for _, r in new): tail += '; %s was linked, and the explicit values remove the link in these configs' % name
    o.update(expected_type='PLAN', expected_detail=' > '.join(parts) + tail, fallback='a run that fails is skipped and listed in batch_summary')
    return o

def sweep_rows(S, cfgs, reg, bases, emit):
    """Two phrases per parameter: a numeric range, and a bare 'sweep X' that means every legal value."""
    def cur_of(c, spec):
        name = spec[-1]
        if spec[0] == 'pkg' and c['packages']:
            sel = selected_blocks(c); blk = sel[0] if sel else list(c['packages'])[0]
            src = c['packages'].get(blk, {})
        else: src = c['main']
        k = find_key(src, name)
        return (src[k] if k else None)
    for i, spec in enumerate([s for s in reg if s[0] in ('main', 'pkg')]):
        name = spec[-1]; ref = next((c for c in cfgs if cur_of(c, spec) is not None), None)
        if ref is None: continue
        cur = cur_of(ref, spec); kind = kind_of(S, ref, name, cur)
        # ---- phrase 1: a numeric range
        if kind in ('number', 'pair'):
            c0 = cur if kind == 'number' and not isinstance(cur, str) else numbers_in(cur)[0]
            if c0 == 0: a, b, st = 0.0, 0.04, 0.01
            else:
                a, b, st = sorted([0.8 * c0, 1.2 * c0]) + [abs(0.1 * c0)]
        elif kind == 'count': a, b, st = float(cur), float(cur) + 2, 1.0
        elif kind == 'switch': a, b, st = 0.0, 1.0, 1.0
        elif kind == 'choice' and all(isinstance(v, int) for _, v in legal_values(ref, name)) or name.strip().lower() == 'z_p select':
            lv = legal_values(ref, name); nums = [int(sp) for sp, _ in lv]; a, b, st = float(min(nums)), float(max(nums)), 1.0
        elif kind == 'search': a, b, st = -0.1, 0.0, 0.02
        else: a, b, st = 1.0, 3.0, 1.0
        A, B, ST = val(float('%.6g' % a)), val(float('%.6g' % b)), val(float('%.6g' % st))
        a, b, st = float(A), float(B), float(ST)
        said = ST if st != int(st) else str(len(sweep_points(a, b, st)))   # a whole number after 'in' is a count, so a whole step is said as its count
        how_said, rule_pts = steps_rule(a, b, float(said))
        assert how_said in ('count', 'step') and rule_pts == sweep_points(a, b, st), (name, A, B, ST, said)
        phrase = 'Re-run the setup, but sweep %s from %s to %s in %s steps.' % (name, A, B, said)
        para = 'sweep %s from %s to %s by %s | run %s from %s to %s in steps of %s' % (name, A, B, ST, name, A, B, ST)
        per = {}
        for c in cfgs:
            cv = cur_of(c, spec); kd = kind_of(S, c, name, cv) if cv is not None else kind
            pts = rule_pts
            def q(detail, t='QUESTION', kindname='ask'):
                return dict(target_param=name, target_block='', json_path='', value_before=val(cv), value_after='', linked_changes='', loader_effect='', verified='',
                            note='', fallback='cancel', kind=kindname, expected_type=t, expected_detail=detail)
            if kd == 'number':
                keep = isinstance(cv, str) and cv.strip().startswith('[')        # a number the sheet writes as [46.25] stays in that form
                o = sweep_outcome(S, c, bases[c['label']], spec, [(val(p), '[%s]' % val(p) if keep else p) for p in pts], cv, 'range')
            elif kd == 'count':
                o = sweep_outcome(S, c, bases[c['label']], spec, [(val(p), int(p)) for p in pts], cv, 'range')
            elif kd == 'pair':
                two = numbers_in(cv)
                if two[0] == two[1]:
                    o = sweep_outcome(S, c, bases[c['label']], spec, [(val(p), '[%s %s]' % (val(p), val(p))) for p in pts], cv, 'range')
                    if o['expected_type'] == 'PLAN': o['expected_detail'] += '; %s is a [TX RX] pair with equal values, so each point is written to both and that is announced' % name; o['kind'] = 'pairboth'
                else:
                    o = q('%s is a [TX RX] pair with different values (%s): options sweep TX, sweep RX, sweep both together' % (name, val(cv)), kindname='pairsides')
            elif kd in ('switch', 'choice'):
                lv = legal_values(c, name); spoken = {norm(sp): (sp, sv) for sp, sv in lv}
                asked = [norm(val(p)) for p in pts]
                if all(isinstance(sv, (int,)) or name.strip().lower() == 'z_p select' for _, sv in lv) and all(x in spoken for x in asked):
                    o = sweep_outcome(S, c, bases[c['label']], spec, [spoken[x] for x in asked], cv, 'range')
                elif all(isinstance(sv, int) or name.strip().lower() == 'z_p select' for _, sv in lv):
                    o = q('%s only takes the values %s, and the range asks for %s: options sweep the legal values, cancel' % (name, ', '.join(sp for sp, _ in lv), ', '.join(val(p) for p in pts)), kindname='illegal')
                else:
                    o = q('%s takes named values (%s), not a numeric range: options sweep all of them, cancel' % (name, ', '.join(sp for sp, _ in lv)), kindname='named')
            elif kd == 'search':
                o = q('%s is itself a search range inside one run: options one run with %s searched over [%s:%s:%s], a batch of runs with %s held at each value' % (name, name, A, ST, B, name), kindname='search')
            elif kd in ('matrix', 'vector'):
                o = q('%s holds %d values (%s): a sweep needs to know which of them; options name the element, or sweep all of them together' % (name, len(numbers_in(cv) or []), val(cv)), kindname='element')
            else:
                o = q('%s is text, not a number, and has no list of legal values: it cannot be swept' % name, t='STOP', kindname='text'); o['fallback'] = 'none'
            per[c['label']] = o
        emit('SW%03d' % (i + 1), 'sweep', phrase, para, per)
        # ---- phrase 2: no range given
        phrase = 'Sweep %s and give me a summary of the results and differences.' % name
        para = 'try every %s option and compare | sweep %s' % (name, name)
        per = {}
        for c in cfgs:
            cv = cur_of(c, spec); kd = kind_of(S, c, name, cv) if cv is not None else kind
            if kd in ('switch', 'choice'):
                lv = legal_values(c, name)
                o = sweep_outcome(S, c, bases[c['label']], spec, lv, cv, 'values')
            elif kd == 'text':
                o = dict(target_param=name, target_block='', json_path='', value_before=val(cv), value_after='', linked_changes='', loader_effect='', verified='', note='',
                         fallback='none', kind='text', expected_type='STOP', expected_detail='%s is text with no list of legal values: there is nothing to sweep' % name)
            else:
                what = {'search': 'already holds a search range for one run', 'matrix': 'holds several values', 'vector': 'holds several values', 'pair': 'is a [TX RX] pair of numbers'}.get(kd, 'is a number')
                o = dict(target_param=name, target_block='', json_path='', value_before=val(cv), value_after='', linked_changes='', loader_effect='', verified='', note='',
                         fallback='cancel', kind='needrange', expected_type='QUESTION', expected_detail='%s %s, with no closed list of values: ask for a range (from, to, step) or a list of values' % (name, what))
            per[c['label']] = o
        emit('SV%03d' % (i + 1), 'sweep', phrase, para, per)

# ----------------------------------------------------------------------------- phrases
T = ["Re-run with {n} = {v}.", "Change {n} to {v} and run again.", "Same run, but set {n} to {v}.", "Run it again with {n} at {v}.", "Now try {n} = {v}."]
TB = ["Re-run with {n} = {v} in {b}.", "Change {n} in {b} to {v} and run again.", "Same run, but set {n} in {b} to {v}.", "Run it again with {b} {n} at {v}.", "Now try {n} = {v} for {b}."]

def phrase_registry(cfgs):
    """Ordered phrase specs. The first config fixes the order, so ids stay stable; later configs append."""
    reg = []
    def add(s):
        k = tuple(x.strip().lower() if i == len(s) - 1 else x for i, x in enumerate(s))
        if k not in seen: seen.add(k); reg.append(s)
    seen = set()
    for n, cfg in enumerate(cfgs):
        blocks = list(cfg['packages']); pk = []
        for b in blocks:
            for name in cfg['packages'][b]:
                if name not in pk: pk.append(name)
        for name in cfg['main']: add(('main', name)) if not (name in pk and blocks) else None
        if blocks:
            for name in pk: add(('pkg', name))
            for b in blocks[1:] + blocks[:1]:
                for name in cfg['packages'][b]: add(('named', b, name))
    return reg

def generate(S, cfgs):
    reg = phrase_registry(cfgs); rows = []; seen_open = set(); skipped = []
    bases = {c['label']: load_fields(S, cfg=c)[0] for c in cfgs}
    def emit(rid, group, phrase, para, per):
        groups = collections.OrderedDict()
        for c in cfgs:
            o = per[c['label']]; groups.setdefault(json.dumps(o, sort_keys=True), (o, []))[1].append(c['label'])
        for o, labels in groups.values():
            r = {x: '' for x in COLS}; r.update({x: o[x] for x in o if x in COLS})
            if o['kind'] in DECIDED and o['kind'] not in seen_open:
                r['note'] = (r['note'] + '; ' if r['note'] else '') + DECIDED[o['kind']]; seen_open.add(o['kind'])
            r.update(id=rid, group=group, priority='must', starting_config='; '.join(labels), prior_state=PRIOR, phrase=phrase, paraphrases=para, open_decision='')
            rows.append(r)
    for i, spec in enumerate(reg):
        name = spec[-1]; k = i % 5
        cur = None
        for c in cfgs:                                             # current value in the first config that has it, to pick a test value
            try:
                if spec[0] == 'named': cur = c['packages'][spec[1]][find_key(c['packages'][spec[1]], name)]
                elif spec[0] == 'pkg': cur = next(b[find_key(b, name)] for b in c['packages'].values() if find_key(b, name))
                else: cur = c['main'][find_key(c['main'], name)]
                break
            except (KeyError, StopIteration): continue
        new = test_value(name, cur)
        if new is None: skipped.append(spec); continue
        say = SAY.get(name, val(new)) if spec[0] == 'main' else val(new)
        ph = [t.format(n=name, v=say, b=spec[1]) for t in TB] if spec[0] == 'named' else [t.format(n=name, v=say) for t in T]
        after = '[ 2 ]' if name == 'z_p select' else val(new)
        emit('PA%03d' % (i + 1), 'param', ph[k], ' | '.join([ph[(k + 1) % 5], ph[(k + 2) % 5]]),
             {c['label']: classify(S, c, bases[c['label']], spec, new, after) for c in cfgs})
    # enumerations: package pairs come from the blocks each config defines
    n = 0
    def EN(phrase, para, name, new, after, note):
        nonlocal n; n += 1; per = {}
        for c in cfgs:
            o = classify(S, c, bases[c['label']], ('main', name), new, after)
            o['note'] = note; per[c['label']] = o
        emit('EN%02d' % n, 'enum', phrase, para, per)
    allblocks = []
    for c in cfgs:
        for b in c['packages']:
            if b not in allblocks: allblocks.append(b)
    for a in allblocks:
        for b in allblocks:
            if a == b: ph = 'Re-run with package %s.' % a; para = 'use %s for both ends and run again | same run with %s on TX and RX' % (a, a)
            else: ph = 'Re-run with TX package %s and RX package %s.' % (a, b); para = 'run again, %s on the transmit side and %s on the receive side' % (a, b)
            EN(ph, para, 'PKG_NAME', '%s %s' % (a, b), '%s %s' % (a, b), 'first name is TX, second is RX; a single name applies to both (MATLAB 4p15p0 L10355)')
    for v in ['1', '3', '4', '1 2', '2 4', '1 2 3 4']:
        s = v if ' ' not in v else '[%s]' % v
        EN('Re-run with z_p select = %s.' % s, 'run again with z_p select %s | same run, z_p select %s' % (s, s), 'z_p select', '[ %s ]' % v, '[ %s ]' % v,
           'one run returns one result per selected package case' if ' ' in v else '')
    for v in ['1 3 2 4', '3 1 2 4', '2 1 3 4', '1 3 4 2', '1 2 4 3', '3 1 4 2', '2 1 4 3']:
        EN('Re-run with Port Order = [%s].' % v, 'run again with port order %s | same run, ports ordered %s' % (v, v), 'Port Order', '[ %s ]' % v, '[ %s ]' % v,
           'one of the eight orders listed on the selections_and_tables sheet')
    for v in (1, 2): EN('Re-run with ts_anchor = %d.' % v, 'run again with ts_anchor %d' % v, 'ts_anchor', v, str(v), '0 = MM, 1 = peak, 2 = max DV (keywords sheet)')
    for v in (0, 1, 2): EN('Re-run with N_bg = %d.' % v, 'run again with N_bg %d' % v, 'N_bg', v, str(v), '0 switches floating taps off; legal values 0 to 3')
    sweep_rows(S, cfgs, reg, bases, emit)
    return rows, skipped

# ----------------------------------------------------------------------------- config check and variants
def check(S, path, KW):
    cfg = discover(path, KW)
    x, xerr = load_fields(S, path=path); j, jerr = load_fields(S, cfg=cfg)
    rep = dict(file=os.path.basename(path), main=len(cfg['main']), blocks={b: len(p) for b, p in cfg['packages'].items()}, links=len(cfg['links']),
               selected=selected_blocks(cfg), issues=list(cfg['issues']), xlsx_load=xerr or '%d fields' % len(x), json_load=jerr or '%d fields' % len(j))
    if x and j: rep['xlsx_vs_json_differing_fields'] = len(diff(x, j))
    if xerr or jerr:
        have = {k.lower() for k in cfg['main']}
        miss = sorted(k for k in MANDATORY if k not in have)
        if miss: rep['issues'].append('keywords the loader reads without a default are missing: ' + ', '.join(miss))
    return cfg, rep

def compare(S, ref_path, other_path, KW):
    """How another sheet differs from a reference sheet, by keyword and by what the loader ends up with."""
    a = discover(ref_path, KW); b = discover(other_path, KW)
    out = dict(only_in_reference=[k for k in a['main'] if k not in b['main']], only_in_other=[k for k in b['main'] if k not in a['main']],
               blocks_only_in_reference=[k for k in a['packages'] if k not in b['packages']], blocks_only_in_other=[k for k in b['packages'] if k not in a['packages']],
               values_differ=['%s: %s -> %s' % (k, val(a['main'][k]), val(b['main'][k])) for k in a['main'] if k in b['main'] and norm(a['main'][k]) != norm(b['main'][k])],
               moved=['%s: %s -> %s' % (k, a['cells'][k], b['cells'][k]) for k in a['main'] if k in b['main'] and a['cells'][k] != b['cells'][k]],
               links=(len(a['links']), len(b['links'])), issues=b['issues'])
    fa, ea = load_fields(S, cfg=a); fb, eb = load_fields(S, cfg=b)
    out['loaded_fields_differ'] = diff(fa, fb) if fa and fb else 'load failed: %s' % (ea or eb)
    return out

def _restore_stored(path, sheet_index, stored):
    """openpyxl drops the stored result of every formula on save. Put them back so the file reads as Excel left it."""
    import zipfile, shutil
    from xml.sax.saxutils import escape
    name = 'xl/worksheets/sheet%d.xml' % sheet_index
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == name:
                x = data.decode('utf-8')
                def put(m):
                    c, attrs, f = m.group(1), m.group(2), m.group(3)
                    v = stored.get(c)
                    if v is None: return m.group(0)
                    if isinstance(v, str): return '<c r="%s"%s t="str"><f>%s</f><v>%s</v></c>' % (c, re.sub(r'\s+t="[^"]*"', '', attrs), f, escape(v))
                    return '<c r="%s"%s><f>%s</f><v>%r</v></c>' % (c, re.sub(r'\s+t="[^"]*"', '', attrs), f, v)
                x = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>([^<]*)</f><v\s*/>(?:</v>)?</c>|<c r="([A-Z]+\d+)"([^>]*)><f>([^<]*)</f><v></v></c>',
                           lambda m: put(m) if m.group(1) else put(_Shift(m)), x)
                data = x.encode('utf-8')
            zout.writestr(item, data)
    shutil.move(tmp, path)

class _Shift:
    def __init__(s, m): s.m = m
    def group(s, i): return s.m.group(0) if i == 0 else s.m.group(i + 3)

def make_variants(base, outdir):
    """Write config sheets that differ from the base the way other people's sheets do."""
    os.makedirs(outdir, exist_ok=True); made = []
    stored_base = {c.coordinate: c.value for row in openpyxl.load_workbook(base, data_only=True)[SHEET].iter_rows() for c in row if c.value is not None}
    def find(ws, label):
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip().lower() == label.lower(): return c
        raise KeyError(label)
    def save(name, desc, edit, mode='keep'):
        """mode: keep = formulas with their stored results; flat = values pasted over formulas; lost = formulas, no stored results"""
        wb = openpyxl.load_workbook(base); ws = wb[SHEET]
        edit(wb, ws)
        if mode == 'flat':
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith('='): c.value = stored_base.get(c.coordinate)
        p = os.path.join(outdir, name + '.xlsx'); wb.save(p)
        if mode == 'keep':
            forms = {c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')}
            _restore_stored(p, wb.sheetnames.index(SHEET) + 1, {c: stored_base[c] for c in forms if c in stored_base})
        made.append((name, desc, p))
    def omit(label):
        def f(wb, ws):
            c = find(ws, label); ws.cell(c.row, c.column + 1).value = None; c.value = None
        return f
    def rename(old, new):
        def f(wb, ws): find(ws, old).value = new
        return f
    def move(label, to):
        """as a cut and paste in Excel: formulas that pointed at the old cell follow it"""
        def f(wb, ws):
            c = find(ws, label); v = ws.cell(c.row, c.column + 1); t = ws[to]; nv = ws.cell(t.row, t.column + 1)
            old, new = v.coordinate, nv.coordinate
            t.value = c.value; nv.value = v.value; c.value = None; v.value = None
            for row in ws.iter_rows():
                for x in row:
                    if isinstance(x.value, str) and x.value.startswith('='):
                        x.value = REF.sub(lambda m: m.group(0).replace(m.group(1), re.match('[A-Z]+', new).group()).replace(m.group(2), re.search(r'\d+', new).group())
                                          if m.group(1) + m.group(2) == old else m.group(0), x.value)
        return f
    def add(label, value, at):
        def f(wb, ws): ws[at] = label; ws.cell(ws[at].row, ws[at].column + 1).value = value
        return f
    def dup(label, at):
        def f(wb, ws):
            c = find(ws, label); ws[at] = c.value; ws.cell(ws[at].row, ws[at].column + 1).value = ws.cell(c.row, c.column + 1).value
        return f
    def no_blocks(wb, ws):
        """legacy single-package layout: the selected block's parameters sit in the main area, no PKG_NAME"""
        pn = find(ws, 'PKG_NAME'); sel = str(ws.cell(pn.row, pn.column + 1).value).split()[0]
        rows = [c.row for row in ws.iter_rows(min_col=1, max_col=1) for c in row if c.value in ('.START', '.END')]
        pairs = list(zip(rows[0::2], rows[1::2])); keep = next((s, e) for s, e in pairs if ws.cell(s, 2).value == sel)
        for s, e in pairs:
            for r in range(s, e + 1):
                if keep[0] < r < keep[1]: continue
                for col in range(1, 6): ws.cell(r, col).value = None
        omit('PKG_NAME')(wb, ws)
    def rename_blocks(wb, ws):
        m = {'PKG_LowR_CLASSA': 'PkgA', 'PKG_HiR_CLASSB': 'PkgB'}
        for row in ws.iter_rows(min_col=1, max_col=1):
            for c in row:
                if c.value == '.START' and ws.cell(c.row, 2).value in m: ws.cell(c.row, 2).value = m[ws.cell(c.row, 2).value]
        c = find(ws, 'PKG_NAME'); v = ws.cell(c.row, c.column + 1)
        v.value = ' '.join(m.get(t, t) for t in str(v.value).split())
    def mixed(wb, ws):
        c = find(ws, 'PKG_NAME'); ws.cell(c.row, c.column + 1).value = 'PKG_LowR_CLASSA PKG_Module'
    def both(*fs):
        def f(wb, ws):
            for g in fs: g(wb, ws)
        return f
    save('V01_P_peak_omitted', 'P_peak row removed', omit('P_peak'))
    save('V02_c2_and_N_v_added', 'two extra keywords the loader reads: c(2) and N_v', both(add('c(2)', '[-0.06:.02:0]', 'A33'), add('N_v', 3, 'F33')))
    save('V03_keyword_typo', 'P_peak mistyped as P_peek, so the loader ignores it', rename('P_peak', 'P_peek'))
    save('V04_rows_moved', 'DER_0 and SNR_TX moved to other cells, f_b moved to column S', both(move('DER_0', 'J33'), move('SNR_TX', 'F34'), move('f_b', 'S3')))
    save('V05_values_pasted', 'every formula replaced by its value, so there are no links', lambda wb, ws: None, mode='flat')
    save('V06_no_package_blocks', 'legacy layout: one package in the main area, no .START blocks, no PKG_NAME', no_blocks)
    save('V07_blocks_renamed', 'package blocks renamed PkgA and PkgB', rename_blocks)
    save('V08_tx_rx_differ', 'TX uses PKG_LowR_CLASSA, RX uses PKG_Module', mixed)
    save('V09_keyword_case_spaces', 'keywords typed as " der_0 " and "F_B"', both(rename('DER_0', ' der_0 '), rename('f_b', 'F_B')))
    save('V10_duplicate_keyword', 'DER_0 appears twice', dup('DER_0', 'J33'))
    save('V11_f_b_omitted', 'mandatory f_b row removed', omit('f_b'), mode='flat')
    save('V12_stored_results_lost', 'unchanged sheet re-saved by a tool that does not calculate: formulas kept, stored results gone', lambda wb, ws: None, mode='lost')
    return made

def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--sicopr', default='sicopr.py'); ap.add_argument('--out'); ap.add_argument('--check', nargs='*'); ap.add_argument('--variants', nargs=2)
    ap.add_argument('--compare', nargs=2, metavar=('REFERENCE', 'OTHER'))
    ap.add_argument('configs', nargs='*', help='Label=path.xlsx')
    a = ap.parse_args()
    if a.variants:
        for name, desc, p in make_variants(*a.variants): print(name, '|', desc)
        return
    S = load_sicopr(a.sicopr); KW = known_keywords(a.sicopr); KW_MAIN.update(KW[0])
    if a.compare: print(json.dumps(compare(S, a.compare[0], a.compare[1], KW), indent=1))
    for p in a.check or []:
        cfg, rep = check(S, p, KW); print(json.dumps(rep, indent=1))
    if a.configs:
        cfgs = [discover(p, KW, label) for label, p in (c.split('=', 1) for c in a.configs)]
        rows, skipped = generate(S, cfgs)
        with open(a.out or 'rows.csv', 'w', newline='', encoding='utf-8') as f:
            w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(rows)
        print('rows', len(rows), 'phrases', len({r['id'] for r in rows}), 'skipped (no test value)', len(skipped))

if __name__ == '__main__':
    main()
