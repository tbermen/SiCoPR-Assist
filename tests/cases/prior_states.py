"""One builder per prior_state string of the use-case table (sicopr_prior_states.csv).

A prior state is built by seeding files, never by running COM: a working config, a history record
and a result file with made-up numbers stand in for "run complete". Every file is written through
sicopr_assist.project, so the formats are the product's own.

    ctx = Ctx(fx, 'Case1', row, rows_by_id)
    build(row['prior_state'], ctx)
    ctx.root        the project folder
    ctx.prior_wc    the working config the next command starts from
    ctx.env         environment variables to set while the core is called: by default a stand-in Rscript is
                    found and nothing else (fx.r_env), so a dashboard plans the same on every machine; row D3
                    hides it, from PATH and from Program Files

Readings the table leaves open are written down where they are made, and in the Phase 1 report:
every project folder starts from the fixture layout (data\\config.xlsx is the row's starting
workbook, data\\thru.s4p the synthetic THRU), and a state adds to it or takes from it.
"""
# Copyright 2026 Todd Bermensolo
# SPDX-License-Identifier: BSD-3-Clause
import copy
import hashlib
import os
import zipfile

import fixtures
import l4
from sicopr_assist import core, engine
from sicopr_assist.project import Project, history_record, result_record

EXAMPLE = os.path.join(fixtures.ROOT, 'sicopr_assist', 'data', 'sicopr_config_example.json')
RESULT_NAMES = ('channel_operating_margin_dB', 'FOM', 'ERL', 'VEC_dB', 'VEO_mV', 'DER_thresh',
                'available_signal_after_eq_mV')
STANDARD_PLAN = ('import_config', 'validate_inputs', 'run_com', 'summarize')
CHANGE_PLAN = ('set_param', 'validate_config', 'run_com', 'summarize')
FETCH_INDEX_ENV = 'SICOPR_ASSIST_FETCH_INDEX'       # where a test points the version 2 fetch at its stand-in page


def sha(path):
    with open(path, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


class Ctx:
    def __init__(self, fx, label, row=None, rows_by_id=None):
        self.fx = fx
        self.label = label
        self.row = row or {}
        self.rows_by_id = rows_by_id or {}
        self.root = fx.project(label)
        self.project = Project(self.root)
        user_file = _user()
        if os.path.exists(user_file):                  # the user settings folder is shared by every project
            os.remove(user_file)
        self.env = fx.r_env()
        self.prior_wc = fx.working_config(label) if label in fx.configs else None
        self._k = 0

    def data(self, name):
        return os.path.join(self.root, 'data', name)

    def rel(self, path):
        return os.path.relpath(path, self.root)

    def numbers(self):
        """Made-up result values, different for every run in a project."""
        self._k += 1
        k = self._k
        return {'channel_operating_margin_dB': round(3.5 + 0.37 * k, 4), 'FOM': round(14.2 + 0.11 * k, 4),
                'ERL': round(12.0 + 0.05 * k, 4), 'VEC_dB': round(9.0 - 0.2 * k, 4),
                'VEO_mV': round(8.5 + 0.3 * k, 4), 'DER_thresh': 2e-4,
                'available_signal_after_eq_mV': round(10.0 + 0.4 * k, 4)}


# ----------------------------------------------------------------------------- working-config edits
def key_of(json_path):
    """'/main/DER_0' -> 'DER_0'; '/packages/B/C_p' -> 'packages.B.C_p' (the oracle's key form)."""
    parts = json_path.strip('/').split('/', 2)
    if parts[0] == 'main':
        return parts[1]
    return 'packages.%s.%s' % (parts[1], parts[2])


def json_path_of(key):
    if key.startswith('packages.'):
        _, blk, name = key.split('.', 2)
        return '/packages/%s/%s' % (blk, name)
    return '/main/' + key


def with_values(wc, values):
    """wc with {key: value} applied as the oracle applies a change: links downstream recomputed."""
    new, _ = fixtures.oracle().apply(wc, values)
    return fixtures.working_config(dict(new))


def overrides_of(before, values):
    out = []
    for key, v in values.items():
        scope, name = fixtures.oracle()._scope(before, key)
        out.append({'json_path': json_path_of(key), 'before': scope.get(name), 'after': v,
                    'kind': 'set' if name in scope else 'add'})
    return out


def plan(steps, **fields):
    return core.outcome('PLAN', steps=[{'op': s, 'args': {}} for s in steps], pattern='seeded', **fields)


# ----------------------------------------------------------------------------- seeding
def seed_run(ctx, phrase='Run COM on data\\thru.s4p with data\\config.xlsx', wc=None, overrides=(),
             steps=STANDARD_PLAN, channels=None, cases=None, record=True, source=None, numbers=None):
    """One complete run: working config, run CSV, result file and (unless record is False) the history
    record. Returns the run id. The new working config becomes the prior working config."""
    p = ctx.project
    wc = copy.deepcopy(wc if wc is not None else ctx.prior_wc)
    rid = p.next_run_id()
    p.write_working_config(rid, wc)
    fixtures.oracle().to_csv(wc, p.run_csv_path(rid))
    chans = channels if channels is not None else [('data\\thru.s4p', 'thru')]
    chan_recs = [{'path': os.path.join(*c.split('\\')), 'sha256': sha(os.path.join(ctx.root, *c.split('\\'))),
                  'role': role} for c, role in chans]
    if cases is None:
        cases = [numbers or ctx.numbers()]
    p.write_result(rid, result_record(run_id=rid, status='complete', release=engine.default_release(),
                                      sicopr_sha256=engine.engine_sha256(), config_csv=ctx.rel(p.run_csv_path(rid)),
                                      channels=chan_recs, cases=cases))
    if source is None:
        src = ctx.data('config.xlsx')
        source = {'path': ctx.rel(src), 'sha256': sha(src)} if os.path.exists(src) else None
    if record:
        p.append_history(history_record(
            phrase=phrase, pattern='seeded', outcome='PLAN', status='ok',
            plan=plan(steps, changes=[dict(o, applied=True) for o in overrides]),
            source=source, channels=chan_recs, overrides=list(overrides),
            working_config=ctx.rel(p.working_config_path(rid)), run_ids=[rid],
            sicopr={'sha256': engine.engine_sha256(), 'release': engine.default_release()}))
    ctx.prior_wc = wc
    return rid


def seed_override_run(ctx, values, phrase=None):
    before = ctx.prior_wc
    wc = with_values(before, values)
    ov = overrides_of(before, values)
    phrase = phrase or 'Re-run with ' + ' and '.join('%s = %s' % (k.rsplit('.', 1)[-1], v) for k, v in values.items())
    return seed_run(ctx, phrase=phrase, wc=wc, overrides=ov, steps=CHANGE_PLAN)


def seed_batch(ctx, phrase, points, skipped=(), reuse_first=False):
    """A finished batch. points: [{key: value}] in order; a point in skipped failed validation and has
    no run. The first point that equals the prior config reuses the prior run when reuse_first."""
    p = ctx.project
    base = ctx.prior_wc
    prior_rid = None
    if reuse_first:
        prior_rid = seed_run(ctx)
    recs, rids = [], []
    for values in points:
        jvals = {json_path_of(k): v for k, v in values.items()}
        if values in skipped:
            recs.append({'values': jvals, 'run_id': None, 'new': False, 'reused_from': None,
                         'skipped': 'validate_config failed'})
            continue
        wc = with_values(base, values)
        if prior_rid and wc == base:
            recs.append({'values': jvals, 'run_id': prior_rid, 'new': False, 'reused_from': prior_rid, 'skipped': None})
            rids.append(prior_rid)
            continue
        rid = seed_run(ctx, wc=wc, record=False)
        recs.append({'values': jvals, 'run_id': rid, 'new': True, 'reused_from': None, 'skipped': None})
        rids.append(rid)
    src = ctx.data('config.xlsx')
    p.append_history(history_record(
        phrase=phrase, pattern='seeded', outcome='PLAN', status='ok',
        plan=plan(('build_matrix', 'validate_config', 'run_batch', 'batch_summary'), points=recs,
                  new_runs=sum(1 for r in recs if r['new'])),
        source={'path': ctx.rel(src), 'sha256': sha(src)},
        channels=[{'path': os.path.join('data', 'thru.s4p'), 'sha256': sha(ctx.data('thru.s4p')), 'role': 'thru'}],
        run_ids=rids, points=recs, sicopr={'sha256': engine.engine_sha256(), 'release': engine.default_release()}))
    ctx.prior_wc = base
    return rids


def seed_commands(ctx, sequence, version=None):
    """History records: sequence of 'ok', 'miss', 'cancelled', 'failed', 'answered' or 'shown' (an ok command
    after which the tutor and the feedback line were shown).

    The first 'ok' is a seeded run of the starting config on data\\thru.s4p (a working config, a result file and
    its history record), and the rest are history records only. Successful commands in a project folder that
    holds both files include a run, and the phrases that follow these states (Re-run with ..., rows TU14 to TU19)
    need one: without it they would be read as a re-run with no run to start from (Phase V2-0 finding)."""
    phrases = ['Re-run with DER_0 = 1e-4.', 'Show settings.', 'What did I run last?',
               'Re-run with SNR_TX = 32.5.', 'Show me the summary again.']
    ran = False
    for i, kind in enumerate(sequence):
        ph = phrases[i % len(phrases)]
        if kind == 'ok' and not ran:
            seed_run(ctx)
            ran = True
            continue
        if kind == 'ok':
            rec = history_record(phrase=ph, pattern='seeded', outcome='PLAN', status='ok')
        elif kind == 'shown':
            rec = history_record(phrase=ph, pattern='seeded', outcome='PLAN', status='ok',
                                 shown=['tutor', 'feedback_notice'])
        elif kind == 'answered':
            rec = history_record(phrase='Re-run the setup, but sweep C_b from 0.1e-4 to 0.5e-4.', pattern='seeded',
                                 outcome='PLAN', status='ok', options=['give a step', 'give a number of points', 'cancel'],
                                 chosen='give a step')
        elif kind == 'miss':
            rec = history_record(phrase='xyzzy the frobnicator', pattern='nothing matched', unplaced=['xyzzy', 'frobnicator'],
                                 outcome='STOP', status='stopped', reason='not a supported request')
        elif kind == 'cancelled':
            rec = history_record(phrase='Re-run with zp select 2', pattern='seeded', outcome='QUESTION',
                                 status='cancelled', unplaced=['zp'], options=['z_p select', 'none of these'], chosen='cancel')
        elif kind == 'failed':
            rec = history_record(phrase='Re-run with f_b = -5.', pattern='seeded', outcome='STOP', status='failed',
                                 reason='f_b must be positive')
        else:
            raise ValueError(kind)
        if version:
            rec['tool_version'] = version
        ctx.project.append_history(rec)


def seed_phrases(ctx, phrases):
    for ph in phrases:
        ctx.project.append_history(history_record(phrase=ph, pattern='seeded', outcome='PLAN', status='ok'))


def write_settings(path, lines):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8', newline='\n') as f:
        f.write('\n'.join(lines) + '\n')


def project_settings(ctx, *lines):
    write_settings(ctx.project.settings_path, ['[limits]'] + list(lines))


def user_settings(ctx, *lines):
    from sicopr_assist.project import user_settings_path
    write_settings(user_settings_path(), ['[limits]'] + list(lines))


def seed_question(ctx, rid_row, cancelled=False):
    """The question a row asks, open (pending) or cancelled, after that row's own prior state."""
    row = ctx.rows_by_id[rid_row][0]
    build(row['prior_state'], ctx)
    opts, _ = l4.options_of(row['expected_detail'])
    q = core.outcome('QUESTION', question={'text': row['expected_detail'], 'options': list(opts or [])},
                     pattern='seeded')
    rec = ctx.project.append_history(history_record(
        phrase=row['phrase'], pattern='seeded', outcome='QUESTION', status='cancelled' if cancelled else 'question',
        options=list(opts or []), chosen='cancel' if cancelled else None, plan=q))
    if not cancelled:
        ctx.project.write_pending(row['phrase'], q, seq=rec['seq'])


def replace_file(path, data):
    """Write new bytes at path without touching the file a hard link may share."""
    if os.path.exists(path):
        os.remove(path)
    with open(path, 'wb') as f:
        f.write(data)


def edit_workbook(src, dst, edits=None, drop=None):
    """A copy of a workbook with keyword values changed or a keyword row removed, saved the way Excel
    leaves a file: every formula keeps its stored result (the oracle's _restore_stored)."""
    import openpyxl
    O = fixtures.oracle()
    stored = {c.coordinate: c.value for row in openpyxl.load_workbook(src, data_only=True)[O.SHEET].iter_rows()
              for c in row if c.value is not None}
    wb = openpyxl.load_workbook(src)
    ws = wb[O.SHEET]
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str):
                continue
            k = c.value.strip()
            if edits and k in edits:
                ws.cell(c.row, c.column + 1).value = edits[k]
            if drop and k == drop:
                ws.cell(c.row, c.column + 1).value = None
                c.value = None
    if os.path.exists(dst):
        os.remove(dst)
    wb.save(dst)
    forms = {c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')}
    O._restore_stored(dst, wb.sheetnames.index(O.SHEET) + 1, {c: stored[c] for c in forms if c in stored})


# ----------------------------------------------------------------------------- files-only states
def base_layout(ctx):
    """The fixture layout and nothing else: no history, no results, no settings."""


def two_configs(ctx):
    other = 'Case3' if ctx.label != 'Case3' else 'Case1'
    fixtures.link_or_copy(ctx.fx.configs[other], ctx.data('config_2.xlsx'))


def truncated_s4p(ctx):
    with open(ctx.fx.thru, 'rb') as f:
        data = f.read()
    replace_file(ctx.data('bad.s4p'), data[:len(data) // 3 + 17])          # cut in the middle of a line


def low_frequency_thru(ctx):
    import contextlib
    import io
    with contextlib.redirect_stdout(io.StringIO()):
        fixtures.make_s4p().write(ctx.data('lowf.s4p'), fmax=20e9)            # 20 GHz; f_b 106.25 needs about 53


def legacy_xls(ctx):
    """A true Excel 97-2003 file: the OLE2 compound-file signature and an empty header sector."""
    replace_file(ctx.data('config.xls'), bytes.fromhex('d0cf11e0a1b11ae1') + bytes(504))


def config_json(ctx, rename=None):
    import json
    wc = copy.deepcopy(ctx.prior_wc)
    if rename:
        old, new = rename
        wc['main'] = {(new if k == old else k): v for k, v in wc['main'].items()}
    with open(ctx.data('config.json'), 'w', encoding='utf-8', newline='\n') as f:
        json.dump(wc, f, indent=1)


def both_files(ctx):
    """The two files the phrase names exist. The fixture layout holds data\\config.xlsx and
    data\\thru.s4p; a phrase that names data\\config.csv also gets that file, written from the
    working config as the run CSV of Decision 2."""
    if 'config.csv' in ctx.row.get('phrase', ''):
        fixtures.oracle().to_csv(ctx.prior_wc, ctx.data('config.csv'))


def case1_beside_variant(ctx):
    """data\\config.xlsx is the Case1 sheet; the row's own config (this config) sits beside it under
    its own file name."""
    os.remove(ctx.data('config.xlsx'))
    fixtures.link_or_copy(ctx.fx.configs['Case1'], ctx.data('config.xlsx'))
    fixtures.link_or_copy(ctx.fx.configs[ctx.label], ctx.data(os.path.basename(ctx.fx.configs[ctx.label])))


def no_f_z_row(ctx):
    edit_workbook(ctx.fx.configs[ctx.label], ctx.data('config.xlsx'), drop='f_z')
    ctx.prior_wc = None                                                     # the import is part of the test


# ----------------------------------------------------------------------------- run states
def r1_complete(ctx):
    seed_run(ctx, phrase='Run SiCoPR for me on the example setup.',
             source={'path': EXAMPLE, 'sha256': sha(EXAMPLE)})


def r2_complete(ctx):
    seed_run(ctx, phrase='Run COM on data\\thru.s4p with data\\config.xlsx')


def d1_complete(ctx):
    seed_run(ctx, phrase='Now give me the results in a dynamic dashboard.',
             steps=('export_mat', 'render_dashboard', 'open_report'))


def no_rscript(ctx):
    """No Rscript for the call to the core: PATH without any folder that holds one, and ProgramFiles pointing at an
    empty folder, because a Windows R install is found there when it is not on PATH (as on this workspace)."""
    ctx.env.update(ctx.fx.r_env(found=False))


def four_package_cases(ctx):
    wc = ctx.prior_wc
    blk = str(wc['main'].get('PKG_NAME', '')).split()[0]
    zp = str(wc['packages'][blk]['z_p (TX)'])
    if len(zp.strip(' []').split(';')[0].split()) != 4:
        raise AssertionError('z_p (TX) of %s does not define 4 package cases: %s' % (ctx.label, zp))


def three_thru_files(ctx):
    for n in ('thru_2.s4p', 'thru_3.s4p'):
        fixtures.link_or_copy(ctx.fx.thru, ctx.data(n))


def thru_modified(ctx):
    with open(ctx.data('thru.s4p'), 'rb') as f:
        data = f.read()
    replace_file(ctx.data('thru.s4p'), data + b'! modified after the last run\n')


def thru_replaced(ctx):
    import contextlib
    import io
    os.remove(ctx.data('thru.s4p'))
    with contextlib.redirect_stdout(io.StringIO()):
        fixtures.make_s4p().write(ctx.data('thru.s4p'), 300.0)


THRU2_MM = 300.0                                    # data\thru2.s4p: a THRU of another length than the 400 mm one


def second_thru(ctx):
    fixtures.link_or_copy(ctx.fx.thru_of(THRU2_MM), ctx.data('thru2.s4p'))


def fext_file(ctx):
    fixtures.link_or_copy(ctx.fx.thru, ctx.data('fext1.s4p'))


def crosstalk_run(ctx):
    for n in ('fext_1.s4p', 'next_1.s4p'):
        fixtures.link_or_copy(ctx.fx.thru, ctx.data(n))
    seed_run(ctx, channels=[('data\\thru.s4p', 'thru'), ('data\\fext_1.s4p', 'fext'), ('data\\next_1.s4p', 'next')])


# after the maintainer's second session (7 October): THRUs with crosstalk files named after them (Decision 19).
# Every file is a hard link to the synthetic THRU: the tests read the names and never run COM on them.
XT_FEXT = ['b300_xtalk1_Fext.s4p', 'b300_xtalk2_Fext.s4p', 'b300_xtalk3_Fext.s4p']
XT_NEXT = ['b300_xtalk4_Next.s4p', 'b300_xtalk5_Next.s4p', 'b300_xtalk6_Next.s4p', 'b300_xtalk7_Next.s4p']


def _channel_folder(ctx, folder, names):
    for n in names:
        fixtures.link_or_copy(ctx.fx.thru, ctx.data(os.path.join(folder, n)))


def xtalk_folder(ctx, folder='ch'):
    """data\\ch: two THRUs, each with three FEXT and four NEXT files named after it."""
    names = []
    for stem in ('b300', 'b500'):
        names += [n.replace('b300', stem) for n in ['b300_thru1.s4p'] + XT_FEXT + XT_NEXT]
    _channel_folder(ctx, folder, names)


def xtalk_thru_only(ctx):
    xtalk_folder(ctx)
    seed_run(ctx, phrase='Run COM on data\\ch\\b300_thru1.s4p with data\\config.xlsx',
             channels=[('data\\ch\\b300_thru1.s4p', 'thru')])


def xtalk_done(ctx):
    xtalk_folder(ctx)
    seed_run(ctx, phrase='Re-run with crosstalk.',
             channels=[('data\\ch\\b300_thru1.s4p', 'thru')] + [('data\\ch\\' + n, 'fext') for n in XT_FEXT]
             + [('data\\ch\\' + n, 'next') for n in XT_NEXT])


def xtalk_one_thru(ctx):
    _channel_folder(ctx, 'ch1', ['x_thru_Tx7_to_Rx7.s4p', 'x_FEXT_Tx5_to_Rx7.s4p', 'x_FEXT_Tx6_to_Rx7.s4p',
                                 'x_FEXT_Tx8_to_Rx7.s4p', 'x_NEXT_Tx7_to_Rx7.s4p', 'x_NEXT_Tx8_to_Tx7.s4p'])
    seed_run(ctx, phrase='Run COM on data\\ch1\\x_thru_Tx7_to_Rx7.s4p with data\\config.xlsx',
             channels=[('data\\ch1\\x_thru_Tx7_to_Rx7.s4p', 'thru')])


def xtalk_no_match(ctx):
    _channel_folder(ctx, 'ch2', ['b300_thru1.s4p', 'b500_thru1.s4p', 'c700_xtalk1_Fext.s4p', 'c700_xtalk2_Fext.s4p',
                                 'c700_xtalk4_Next.s4p'])
    seed_run(ctx, phrase='Run COM on data\\ch2\\b300_thru1.s4p with data\\config.xlsx',
             channels=[('data\\ch2\\b300_thru1.s4p', 'thru')])


def config_edited_in_excel(ctx):
    edit_workbook(ctx.fx.configs[ctx.label], ctx.data('config.xlsx'), edits={'SNR_TX': 34})


def config_locked(ctx):
    owner = b'\x0bTest User' + bytes(54) + b'\x0b\x00T\x00e\x00s\x00t\x00' + bytes(100)
    replace_file(ctx.data('~$config.xlsx'), owner[:165])


def below_threshold(ctx):
    n = ctx.numbers()
    n['channel_operating_margin_dB'] = 2.1                                   # COM Pass threshold is 3 dB in the starting configs
    seed_run(ctx, numbers=n)


def zp_select_all(ctx):
    wc = with_values(ctx.prior_wc, {'z_p select': '[ 1 2 3 4 ]'})
    return seed_run(ctx, phrase='Re-run with z_p select = [1 2 3 4].', wc=wc,
                   overrides=overrides_of(ctx.prior_wc, {'z_p select': '[ 1 2 3 4 ]'}), steps=CHANGE_PLAN,
                   cases=[ctx.numbers() for _ in range(4)])


# ----------------------------------------------------------------------------- fetch environment (version 2)
def fetch_env(ctx, confirmed=True, network='up', cached=False, superseded=False, unpaired=False,
              invalid=False, no_touchstone=False, interrupted=False, thru_only=False, multi=False, tie=False):
    """A local stand-in for the IEEE P802.3dj tools page and its zips. Tests never touch the network:
    the page is a file, and FETCH_INDEX_ENV points at it while the core is called.

    A minimal stand-in until Phase V2-3 writes the page in the recorded markup: multi gives the newest
    entry a second zip under the same date (row F11), tie gives a second entry the newest date (row F12)."""
    import contextlib
    import io
    replace_file(ctx.data('config.xls'), open(ctx.fx.configs[ctx.label], 'rb').read())   # an .xlsx named .xls
    d = os.path.join(ctx.root, 'ieee_standin')
    os.makedirs(d)
    small = os.path.join(d, '_small.s4p')
    with contextlib.redirect_stdout(io.StringIO()):
        fixtures.make_s4p().write(small, df=1e9)
    good = open(small, 'rb').read()
    os.remove(small)

    def zip_of(name, files):
        path = os.path.join(d, name)
        with zipfile.ZipFile(path, 'w') as z:
            for n, data in files.items():
                z.writestr(n, data)
        return path

    files = {}
    for ch in ('ch1', 'ch2', 'ch3'):
        files['%s_thru.s4p' % ch] = good
        if not thru_only:
            files['%s_fext1.s4p' % ch] = good
            files['%s_next1.s4p' % ch] = good
    if unpaired:
        files = {('alpha_%d.s4p' % i): good for i in range(len(files))}
    if invalid:
        files['ch2_thru.s4p'] = good[:len(good) // 2]
    if no_touchstone:
        files = {'readme.txt': b'channel data to follow\n'}
    newest = zip_of('dj_kr_channels_2026_09.zip', files)
    older = zip_of('dj_kr_channels_2026_06.zip', {'ch1_thru.s4p': good, 'ch1_fext1.s4p': good, 'ch1_next1.s4p': good})
    rows = [('2026-09-20', os.path.basename(newest), 'superseded' if superseded else 'current', 'KR'),
            ('2026-06-11', os.path.basename(older), 'current', 'KR')]
    if multi or tie:
        second = zip_of('dj_kr_channels_2026_09_b.zip', files)
        rows.insert(1, ('2026-09-20', os.path.basename(second), 'same entry' if multi else 'current', 'KR'))
    size = os.path.getsize(newest)
    html = ['<html><body><table>']
    for date, name, status, iface in rows:
        listed = size * 3 if (interrupted and name == os.path.basename(newest)) else os.path.getsize(os.path.join(d, name))
        html.append('<tr><td>%s</td><td><a href="%s" data-size="%d">%s</a></td><td>%s</td><td>%s</td></tr>'
                    % (date, name, listed, name, status, iface))
    html.append('</table></body></html>')
    with open(os.path.join(d, 'index.html'), 'w', encoding='utf-8', newline='\n') as f:
        f.write('\n'.join(html) + '\n')
    with open(os.path.join(d, 'network'), 'w', encoding='utf-8', newline='\n') as f:
        f.write(network + '\n')
    ctx.env[FETCH_INDEX_ENV] = os.path.join(d, 'index.html')
    if cached:
        cache = os.path.join(ctx.project.dir, 'cache')
        os.makedirs(cache, exist_ok=True)
        fixtures.link_or_copy(newest, os.path.join(cache, os.path.basename(newest)))
    if confirmed:
        ctx.project.append_history(history_record(
            phrase='KR', pattern='confirm_interface', outcome='PLAN', status='ok',
            options=['KR', 'CR', 'C2C', 'C2M'], chosen='KR'))
    return d


def f1_complete(ctx):
    fetch_env(ctx)
    points = []
    for ch in ('ch1', 'ch2', 'ch3'):
        for mode in ('thru_only', 'xtalk'):
            points.append((ch, mode))
    recs, rids = [], []
    for ch, mode in points:
        rid = seed_run(ctx, record=False)
        recs.append({'values': {'channel': ch, 'mode': mode}, 'run_id': rid, 'new': True, 'reused_from': None,
                     'skipped': None})
        rids.append(rid)
    ctx.project.append_history(history_record(
        phrase="Here's my config spreadsheet at data\\config.xls. Download the latest 802.3dj channels and run them.",
        pattern='seeded', outcome='PLAN', status='ok', run_ids=rids, points=recs,
        plan=plan(('import_config', 'fetch_index', 'download', 'build_matrix', 'run_batch', 'batch_summary'),
                  points=recs, new_runs=len(recs))))


# ----------------------------------------------------------------------------- the registry
FALCON = ['Run COM on D:\\acme\\falcon_rev2\\thru.s4p with D:\\acme\\cfg_falcon.xlsx',
          'Re-run with package PKG_AcmeBGA']


def _cb_points(values):
    return [{'C_b': '[%s %s]' % (v, v)} for v in values]


BUILDERS = {
    'none': base_layout,
    r'none; data\config.xlsx and data\thru.s4p exist': base_layout,
    'none; project folder holds exactly one config and one THRU file': base_layout,
    'none; project folder holds two configs': two_configs,
    'none; no 448G preset exists': base_layout,
    r'none; data\config.xlsx exists, data\config.xls does not': base_layout,
    r'none; data\bad.s4p is truncated': truncated_s4p,
    'none; THRU file max frequency is well below what the baud rate needs': low_frequency_thru,
    'R1 complete': r1_complete,
    'no prior run in this project': base_layout,
    'R1 complete; Rscript not on PATH': lambda c: (r1_complete(c), no_rscript(c)),
    'R1 then D1 complete': lambda c: (r1_complete(c), d1_complete(c)),
    'R1 complete; z_p (TX) defines 4 package cases': lambda c: (four_package_cases(c), r1_complete(c)),
    'R2 complete; inventory holds three THRU files': lambda c: (r2_complete(c), three_thru_files(c)),
    'C2 complete': lambda c: (r1_complete(c), d1_complete(c),
                              seed_override_run(c, {'z_p select': '[ 2 ]'}, 'Re-run with z_p select = 2')),
    r'R2 complete; data\thru.s4p modified on disk since that run': lambda c: (r2_complete(c), thru_modified(c)),
    r'none; data\config.xls exists; interface type KR already confirmed for this project; network up': fetch_env,
    'as F1 but interface type not yet confirmed': lambda c: fetch_env(c, confirmed=False),
    'as F1 but no network': lambda c: fetch_env(c, network='down'),
    'as F1 but the selected zip is already cached': lambda c: fetch_env(c, cached=True),
    'as F1 but the newest entry is marked superseded or do not use': lambda c: fetch_env(c, superseded=True),
    'as F1 but FEXT and NEXT files cannot be paired to THRU files with confidence': lambda c: fetch_env(c, unpaired=True),
    'as F1 but one Touchstone file in the zip is invalid': lambda c: fetch_env(c, invalid=True),
    'as F1 but the zip contains no Touchstone files': lambda c: fetch_env(c, no_touchstone=True),
    'as F1 but the download is interrupted': lambda c: fetch_env(c, interrupted=True),
    'as F1 but the zip holds THRU files only': lambda c: fetch_env(c, thru_only=True),
    # rows F11 and F12 (7 October, from the page itself): minimal stand-ins until Phase V2-3
    'as F1 but the newest entry for the interface holds several zips under one date': lambda c: fetch_env(c, multi=True),
    'as F1 but two entries for the interface share the newest date': lambda c: fetch_env(c, tie=True),
    'new terminal; history file from an earlier session exists': r2_complete,
    'history exists': r2_complete,
    'project folder with no history': base_layout,
    'last run has COM below the pass threshold': below_threshold,
    'any': base_layout,
    r'run complete with the starting config and data\thru.s4p': lambda c: seed_run(c),
    'run complete with one active override: DER_0 = 1e-4': lambda c: (seed_run(c), seed_override_run(c, {'DER_0': 1e-4})),
    'run complete with two active overrides: DER_0 = 1e-4, SNR_TX = 32.5':
        lambda c: (seed_run(c), seed_override_run(c, {'DER_0': 1e-4}), seed_override_run(c, {'SNR_TX': 32.5})),
    r'run complete with data\config.xlsx, data\thru.s4p and crosstalk files': crosstalk_run,
    'none; both files exist': both_files,
    'run complete with PKG_NAME = PKG_LowR_CLASSA PKG_Module':
        lambda c: (seed_run(c), seed_override_run(c, {'PKG_NAME': 'PKG_LowR_CLASSA PKG_Module'})),
    r'none; data\config.xls is a legacy Excel 97-2003 file': legacy_xls,
    # the table does not say which override; DER_0 = 1e-4 is the one the other rows use
    'run complete with one active override': lambda c: (seed_run(c), seed_override_run(c, {'DER_0': 1e-4})),
    'run complete with c(2) added (RS26)': lambda c: (seed_run(c), seed_override_run(c, {'c(2)': '[-0.1:.02:0]'})),
    r'run complete; data\config.xlsx edited and saved in Excel since then':
        lambda c: (seed_run(c), config_edited_in_excel(c)),
    r'run complete; data\config.xlsx is open in Excel (lock file present)': lambda c: (seed_run(c), config_locked(c)),
    'run complete with two active overrides':
        lambda c: (seed_run(c), seed_override_run(c, {'DER_0': 1e-4}), seed_override_run(c, {'SNR_TX': 32.5})),
    r'none; data\config.json was written by an earlier session': config_json,
    r'none; data\config.json was edited by hand and has f_bb in place of f_b': lambda c: config_json(c, ('f_b', 'f_bb')),
    r'none; data\config.xlsx is the Case1 sheet': case1_beside_variant,
    r'none; the config and data\thru.s4p exist': base_layout,
    'run complete with C_b = [0.2e-4 0.3e-4]': lambda c: (seed_run(c), seed_override_run(c, {'C_b': '[0.2e-4 0.3e-4]'})),
    'runs complete for Clip Method = Slow and for Clip Method = Fast':
        lambda c: seed_batch(c, 'Sweep clip method and give me a summary of the results and differences.',
                             [{'Clip Method': 'Slow'}, {'Clip Method': 'Fast'}], reuse_first=True),
    r'run complete; data\thru.s4p was replaced on disk since then': lambda c: (seed_run(c), thru_replaced(c)),
    'run complete with one active override: N_qb = 0': lambda c: (seed_run(c), seed_override_run(c, {'N_qb': 0})),
    'the question of SX34 is open': lambda c: seed_question(c, 'SX34'),
    'new project folder: no history yet; both files exist': base_layout,
    'history holds 20 successful commands': lambda c: seed_commands(c, ['ok'] * 20),
    'no settings file; history holds 4 successful commands': lambda c: seed_commands(c, ['ok'] * 4),
    'no settings file; history holds 5 successful commands': lambda c: seed_commands(c, ['ok'] * 5),
    'no settings file; history holds 3 successful commands and 6 that were cancelled or failed':
        lambda c: seed_commands(c, ['ok', 'cancelled', 'failed', 'ok', 'cancelled', 'failed', 'ok', 'cancelled', 'failed']),
    'new project folder: no history yet; the user settings file sets tutor_until_commands = 0':
        lambda c: user_settings(c, 'tutor_until_commands = 0'),
    'the project settings file sets tutor_until_commands = 20; history holds 10 successful commands':
        lambda c: (project_settings(c, 'tutor_until_commands = 20'), seed_commands(c, ['ok'] * 10)),
    'history holds 20 successful commands; 2 of the last 4 commands were misses':
        lambda c: seed_commands(c, ['ok'] * 18 + ['ok', 'miss', 'ok', 'miss']),
    'history holds 20 successful commands; 1 of the last 4 commands was a miss':
        lambda c: seed_commands(c, ['ok'] * 17 + ['ok', 'ok', 'miss', 'ok']),
    'history holds 20 successful commands; the last four commands were sweep questions that were answered '
    '(the run limit, which element, a range for a number, a step size)':
        lambda c: seed_commands(c, ['ok'] * 16 + ['answered'] * 4),
    # "6 commands ago" read with the "4 commands since then" it is given with: the shown record, then 4 more
    'history holds 20 successful commands; the tutor and the feedback line were shown 6 commands ago; '
    '2 of the 4 commands since then were misses':
        lambda c: seed_commands(c, ['ok'] * 17 + ['shown', 'miss', 'ok', 'miss', 'ok']),
    'history holds 20 successful commands; 2 of the last 4 commands were misses; the user settings file sets '
    'feedback_notice = off':
        lambda c: (seed_commands(c, ['ok'] * 18 + ['ok', 'miss', 'ok', 'miss']),
                   write_settings(_user(), ['[feedback]', 'feedback_notice = off'])),
    'history holds 20 successful commands; the last 4 commands were misses; the user settings file sets '
    'trouble_misses = 0':
        lambda c: (seed_commands(c, ['ok'] * 20 + ['miss'] * 4), user_settings(c, 'trouble_misses = 0')),
    'no settings file anywhere': base_layout,
    'the user settings file sets max_new_runs = 100 and tutor_until_commands = 0; the project settings file '
    'sets max_new_runs = 10':
        lambda c: (user_settings(c, 'max_new_runs = 100', 'tutor_until_commands = 0'),
                   project_settings(c, 'max_new_runs = 10')),
    r'run complete with the starting config and data\thru.s4p; the project settings file sets max_new_runs = 50':
        lambda c: (seed_run(c), project_settings(c, 'max_new_runs = 50')),
    r'run complete with the starting config and data\thru.s4p; the project settings file sets max_new_runs = 5':
        lambda c: (seed_run(c), project_settings(c, 'max_new_runs = 5')),
    r'run complete with the starting config and data\thru.s4p; the project settings file sets max_new_runs = 0':
        lambda c: (seed_run(c), project_settings(c, 'max_new_runs = 0')),
    'the project settings file holds max_runs = 50': lambda c: project_settings(c, 'max_runs = 50'),
    'the project settings file holds max_new_runs = many': lambda c: project_settings(c, 'max_new_runs = many'),
    'the project settings file cannot be parsed':
        lambda c: write_settings(c.project.settings_path, ['[limits', 'max_new_runs == = 3', '[[feedback']),
    'runs complete for Clip Method = Slow and for Clip Method = Fast; max_new_runs was then changed from 30 to 5':
        lambda c: (seed_batch(c, 'Sweep clip method and give me a summary of the results and differences.',
                              [{'Clip Method': 'Slow'}, {'Clip Method': 'Fast'}], reuse_first=True),
                   project_settings(c, 'max_new_runs = 5')),
    'the question of C3 is open': lambda c: seed_question(c, 'C3'),
    'the question of PH12 was cancelled': lambda c: seed_question(c, 'PH12', cancelled=True),
    'history holds 40 commands': lambda c: seed_commands(c, ['ok'] * 40),
    'history holds Run COM on D:\\acme\\falcon_rev2\\thru.s4p with D:\\acme\\cfg_falcon.xlsx, then Re-run with '
    'package PKG_AcmeBGA, then Re-run with zp select 2':
        lambda c: seed_phrases(c, FALCON + ['Re-run with zp select 2']),
    'history written by an earlier SiCoPR-Assist version; its last command was Re-run with zp select 2, '
    'answered z_p select': lambda c: earlier_version(c),
    'the project settings file sets mask_names = off; history holds Run COM on D:\\acme\\falcon_rev2\\thru.s4p '
    'with D:\\acme\\cfg_falcon.xlsx, then Re-run with package PKG_AcmeBGA, then Re-run with C_d = [0.4e-4 0.9e-4]':
        lambda c: (write_settings(c.project.settings_path, ['[feedback]', 'mask_names = off']),
                   seed_phrases(c, FALCON + ['Re-run with C_d = [0.4e-4 0.9e-4]'])),
    'the project settings file sets mask_numbers = on; history holds Run COM on D:\\acme\\falcon_rev2\\thru.s4p '
    'with D:\\acme\\cfg_falcon.xlsx, then Re-run with package PKG_AcmeBGA, then Re-run with C_d = [0.4e-4 0.9e-4]':
        lambda c: (write_settings(c.project.settings_path, ['[feedback]', 'mask_numbers = on']),
                   seed_phrases(c, FALCON + ['Re-run with C_d = [0.4e-4 0.9e-4]'])),
    'the project settings file holds mask_names = maybe; history holds 40 commands':
        lambda c: (write_settings(c.project.settings_path, ['[feedback]', 'mask_names = maybe']),
                   seed_commands(c, ['ok'] * 40)),
    'the C_b sweep of SX02 is complete: 9 results':
        lambda c: seed_batch(c, 'Re-run the setup, but sweep C_b from 0.1e-4 to 0.5e-4 in 0.5e-5 steps.',
                             _cb_points(['1e-05', '1.5e-05', '2e-05', '2.5e-05', '3e-05', '3.5e-05', '4e-05',
                                         '4.5e-05', '5e-05']), reuse_first=True),
    'the Clip Method sweep of SX17 is complete: 2 results':
        lambda c: seed_batch(c, 'Sweep clip method and give me a summary of the results and differences.',
                             [{'Clip Method': 'Slow'}, {'Clip Method': 'Fast'}], reuse_first=True),
    'the C_b sweep of SX27 is complete: 2 results':
        lambda c: seed_batch(c, 'Re-run the setup, but sweep C_b from 0.1e-4 to 0.5e-4 in 2 steps.',
                             _cb_points(['1e-05', '5e-05'])),
    'the grid of SX15 is complete: C_b at 3 values by DER_0 at 2 values, 6 results':
        lambda c: seed_batch(c, 'Sweep C_b from 0.1e-4 to 0.5e-4 in 2e-5 steps and run DER_0 at 1e-4 and 2e-4.',
                             [{'C_b': '[%s %s]' % (v, v), 'DER_0': d} for v in ('1e-05', '3e-05', '5e-05')
                              for d in (1e-4, 2e-4)], reuse_first=True),
    'run complete with z_p select = [ 1 2 3 4 ]: 4 results': zp_select_all,
    r'run complete with the starting config and data\thru.s4p: 1 result': lambda c: seed_run(c),
    'the N_b sweep of SX23 is complete: 2 results, 1 point skipped':
        lambda c: seed_batch(c, 'Re-run the setup, but sweep N_b from 0 to 2 in 3 steps.',
                             [{'N_b': 0}, {'N_b': 1}, {'N_b': 2}], skipped=[{'N_b': 0}], reuse_first=True),
    'F1 complete: every channel run thru-only and with crosstalk': f1_complete,
    'none; both files exist; the config has no f_z row': no_f_z_row,
    # after the maintainer's first use (6 October): a second THRU, or a FEXT file, beside a finished run
    r'run complete with the starting config and data\thru.s4p; data\thru2.s4p also exists':
        lambda c: (seed_run(c), second_thru(c)),
    r'run complete with one active override: DER_0 = 1e-4; data\thru2.s4p also exists':
        lambda c: (seed_run(c), seed_override_run(c, {'DER_0': 1e-4}), second_thru(c)),
    r'run complete with the starting config and data\thru.s4p; data\fext1.s4p also exists':
        lambda c: (seed_run(c), fext_file(c)),
    # after the maintainer's second session (7 October): adding crosstalk (Decision 19)
    r'run complete thru-only with the starting config and data\ch\b300_thru1.s4p; data\ch holds two THRUs, '
    r'b300_thru1.s4p and b500_thru1.s4p, each with three FEXT (b300_xtalk1_Fext.s4p to b300_xtalk3_Fext.s4p) and '
    r'four NEXT files (b300_xtalk4_Next.s4p to b300_xtalk7_Next.s4p) named after it': xtalk_thru_only,
    r'run complete with the starting config, data\ch\b300_thru1.s4p and its three FEXT and four NEXT files; data\ch '
    r'as in the thru-only state': xtalk_done,
    r'run complete thru-only with the starting config and data\ch1\x_thru_Tx7_to_Rx7.s4p; data\ch1 holds that one '
    r'THRU, three FEXT files (x_FEXT_Tx5_to_Rx7.s4p, x_FEXT_Tx6_to_Rx7.s4p, x_FEXT_Tx8_to_Rx7.s4p) and two NEXT '
    r'files (x_NEXT_Tx7_to_Rx7.s4p, x_NEXT_Tx8_to_Tx7.s4p)': xtalk_one_thru,
    r'run complete thru-only with the starting config and data\ch2\b300_thru1.s4p; data\ch2 holds two THRUs, '
    r'b300_thru1.s4p and b500_thru1.s4p, and FEXT and NEXT files whose names match neither': xtalk_no_match,
}


def _user():
    from sicopr_assist.project import user_settings_path
    return user_settings_path()


def earlier_version(ctx):
    r1_complete(ctx)
    before = ctx.prior_wc
    wc = with_values(before, {'z_p select': '[ 2 ]'})
    ov = overrides_of(before, {'z_p select': '[ 2 ]'})
    p = ctx.project
    rid = p.next_run_id()
    seed_run(ctx, phrase='Re-run with zp select 2', wc=wc, overrides=ov, steps=CHANGE_PLAN, record=False)
    rec = history_record(phrase='Re-run with zp select 2', pattern='seeded', outcome='PLAN', status='ok',
                         unplaced=['zp'], options=['z_p select', 'none of these'], chosen='z_p select',
                         plan=plan(CHANGE_PLAN, changes=[dict(o, applied=True) for o in ov]),
                         overrides=ov, working_config=ctx.rel(p.working_config_path(rid)), run_ids=[rid])
    rec['tool_version'] = '0.0.1'
    p.append_history(rec)


def build(state, ctx):
    """Seed the prior state named by state into ctx.root. A state with no builder raises KeyError."""
    if state not in BUILDERS:
        raise KeyError('no builder for prior state %r' % state)
    BUILDERS[state](ctx)
    return ctx


def check_built(ctx):
    """Read back everything the builder wrote through the product's own readers."""
    p = ctx.project
    hist = p.read_history()
    for rec in hist:
        for rid in rec.get('run_ids') or []:
            p.read_working_config(rid)
            p.read_result(rid)
    p.read_pending()
    return len(hist)


def strings():
    import csv
    return [r['prior_state'] for r in csv.DictReader(open(os.path.join(fixtures.HERE, 'sicopr_prior_states.csv'),
                                                          newline='', encoding='utf-8'))]

