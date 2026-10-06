"""Phase 2: the config model, checked by calling config_model.py and xlsx_io.py directly.

    python tests\\test_config.py

Gates (prompt section 8, Phase 2):
    G2.1  every workbook SiCoPR ships and the four starting workbooks: the import, written as the run CSV,
          loads to the same param and OP as the workbook, field for field
    G2.2  18 links per starting workbook; f_b = 53.125 recomputes the seven linked values of row ST06; a
          link outside the grammar of Decision 15 is refused with a warning and nothing evaluates it
    G2.3  every row outside the sweep group verified by the loader, with one json_path and one value,
          applied directly, reproduces its loader_effect
    G2.4  export: loads like the working config, opens in openpyxl, unchanged cells keep formula and
          stored result, a changed linked cell becomes a value; an added keyword (CP07) in column order
    G2.5  links and recompute (CP03, CP04, CP05), import problems on V03, V10, V11, V12 and V06, a .json
          with a wrong key (CP15), export refusing the source path
    G2.6  the source workbooks' SHA-256 never changes

The checks that need the four starting workbooks print SKIP: and are left out when they are absent
(or hidden by SICOPR_ASSIST_HIDE_CONFIGS=1); the rest run on the workbooks SiCoPR ships.
"""
# Copyright 2026 Todd Bermensolo
# SPDX-License-Identifier: BSD-3-Clause
import ast
import builtins
import copy
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import time
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, 'cases'))

from check import check, finish  # noqa: E402
import fixtures  # noqa: E402

import openpyxl  # noqa: E402

from sicopr_assist import config_model as cm  # noqa: E402
from sicopr_assist import engine, xlsx_io  # noqa: E402

LINKED_BY_F_B = ('f_z', 'f_p1', 'f_p2', 'f_HP_PZ', 'f_v', 'f_f', 'f_n')
SOURCES = {}                                # path -> SHA-256 when first used (G2.6)
# Rows recorded under disputes in the build state (Rule 9): a miss on one is excused and named.
DISPUTED = {'IV04': 'loader_effect is the effect of the lower-case spelling as typed, not of its own value_after',
            'IV13': 'loader_effect is the effect of the text yes as typed, not of its own value_after 1'}


def sha(path):
    with open(path, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def used(path):
    SOURCES.setdefault(os.path.abspath(path), sha(path))
    return path


def fields_of_file(path, tmp):
    """(fields, error) of a config file loaded from a temporary copy (a load can write beside the file)."""
    d = tempfile.mkdtemp(dir=tmp)
    q = os.path.join(d, os.path.basename(path))
    shutil.copyfile(path, q)
    try:
        return engine.load_fields(q)
    finally:
        shutil.rmtree(d, ignore_errors=True)


def table_value(text):
    """A value_after of the table as the change would carry it: a number when it reads as one."""
    t = text.strip()
    if re.fullmatch(r'[+-]?[0-9]+', t):
        return int(t)
    if re.fullmatch(r'[+-]?([0-9]+\.?[0-9]*|\.[0-9]+)([eE][+-]?[0-9]+)?', t):
        return float(t)
    return text


def sheet_cells(path):
    """{cell: (formula or value as openpyxl reads it, stored value)} of COM_Settings."""
    wf = openpyxl.load_workbook(path)['COM_Settings']
    wv = openpyxl.load_workbook(path, data_only=True)['COM_Settings']
    return {c.coordinate: (c.value, wv[c.coordinate].value) for row in wf.iter_rows() for c in row}


def edit_sheet_xml(src, dst, cell, formula, keep_stored=True):
    """A copy of src whose cell holds formula (raw text, written into the XML as it is)."""
    sheet = xlsx_io._sheet_xml_name(src, 'COM_Settings')
    with zipfile.ZipFile(src) as z:
        names = z.namelist()
        blobs = {n: z.read(n) for n in names}
    x = blobs[sheet].decode('utf-8')
    pat = re.compile(r'(<c r="%s"[^>]*>)<f>[^<]*</f>(<v>[^<]*</v>)?' % cell)
    y = pat.sub(lambda m: m.group(1) + '<f>%s</f>' % xlsx_io._escape(formula) + (m.group(2) or '' if keep_stored else ''),
                x, count=1)
    assert y != x, 'no formula in %s to replace' % cell
    blobs[sheet] = y.encode('utf-8')
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, blobs[n])
    return dst


class NoEval:
    """Records any call of eval or exec, and of compile other than ast.parse's, while it is active."""
    def __enter__(self):
        self.calls = []
        self.saved = builtins.eval, builtins.exec, builtins.compile
        ev, ex, co = self.saved

        def rec_eval(*a, **k):
            self.calls.append('eval')
            return ev(*a, **k)

        def rec_exec(*a, **k):
            self.calls.append('exec')
            return ex(*a, **k)

        def rec_compile(source, filename, mode, flags=0, *a, **k):
            if not flags & ast.PyCF_ONLY_AST:
                self.calls.append('compile')
            return co(source, filename, mode, flags, *a, **k)
        builtins.eval, builtins.exec, builtins.compile = rec_eval, rec_exec, rec_compile
        return self

    def __exit__(self, *exc):
        builtins.eval, builtins.exec, builtins.compile = self.saved
        return False


# ----------------------------------------------------------------------------- checks without the starting workbooks
def check_grammar():
    good = {'{f_b}/2.5': None, '1/{T_t}/{f_b}*0.2365*1000': None, '{ffe_post_tap_len}+1': None, '-({f_b}+1)': None,
            '{z_p (TX)}': None, '{packages.PKG_LowR_CLASSA.z_p (TX)}': None, '0.005': None, '(1+2)*3': None}
    bad = ['{f_b}**2', '{f_b}^2', 'POWER({f_b},2)', 'SUM({f_b}:{f_p1})', "__import__('os').getcwd()", '{f_b}.real',
           '[1, 2]', '"text"', 'True', '{f_b} if 1 else 2', 'lambda: 1', '{f_b}%2', '2j', '{f_b}//2', '+{f_b}',
           '{f_b} < 3', 'Sheet2!{f_b}', '{f_b', '_ref0', 'x', '']
    ok = []
    for t in good:
        try:
            cm.parse_link(t)
        except cm.LinkRefused as e:
            ok.append('%s refused: %s' % (t, e))
    check('G2.2 grammar: the link forms the workbooks hold are accepted (%d)' % len(good), not ok, ok)
    let = []
    for t in bad:
        try:
            cm.parse_link(t)
            let.append(t)
        except cm.LinkRefused:
            pass
    check('G2.2 grammar: %d links outside Decision 15 are refused' % len(bad), not let, 'accepted: %s' % let)
    vals = {'f_b': 106.25, 'T_t': 6, 'ffe_post_tap_len': 8, 'z_p (TX)': '[ 1 2 ]', 'txt': 'abc'}
    got = {t: cm.evaluate_link(cm.parse_link(t), vals.__getitem__) for t in
           ('{f_b}/2.5', '1/{T_t}/{f_b}*0.2365*1000', '{ffe_post_tap_len}+1', '-({f_b}+1)', '{z_p (TX)}', '8-2-1', '8/2/2')}
    want = {'{f_b}/2.5': 42.5, '1/{T_t}/{f_b}*0.2365*1000': 1 / 6 / 106.25 * 0.2365 * 1000, '{ffe_post_tap_len}+1': 9,
            '-({f_b}+1)': -107.25, '{z_p (TX)}': '[ 1 2 ]', '8-2-1': 5, '8/2/2': 2}
    check('G2.2 grammar: values, left to right, and a single reference copies text', got == want, got)
    try:
        cm.evaluate_link(cm.parse_link('{txt}*2'), vals.__getitem__)
        check('G2.2 grammar: arithmetic on a text value is an error', False)
    except cm.LinkError:
        check('G2.2 grammar: arithmetic on a text value is an error', True)
    # nothing in the config model can run text: no eval, exec or compile call anywhere in its source
    calls = []
    for mod in (cm, xlsx_io):
        tree = ast.parse(open(mod.__file__, encoding='utf-8').read())
        calls += ['%s: %s' % (os.path.basename(mod.__file__), n.func.id) for n in ast.walk(tree)
                  if isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id in ('eval', 'exec', 'compile')]
    check('G2.2 no eval, exec or compile in config_model.py and xlsx_io.py', not calls, calls)


def check_csv_writer(tmp):
    wc = {'schema': 'sicopr-config/1', 'source': {'file': 'x', 'sheet': None, 'sha256': ''},
          'main': {'f_b': 106.25, 'L': 4, 'note': 'a, "quoted" text', 'none': None, 'big': 67000000000.0},
          'packages': {'P': {'R_d': '[ 46.25 46.25]'}}, 'links': {'L': '{f_b}'}}
    text = cm.run_csv_text(wc)
    want = ('f_b,106.25\r\nL,4\r\nnote,"a, ""quoted"" text"\r\nnone,\r\nbig,67000000000.0\r\n.START,P\r\n'
            'R_d,[ 46.25 46.25]\r\n.END\r\n')
    check('Decision 2: the run CSV holds main in order, then .START, pairs, .END; numbers by repr, null empty, '
          'a comma or quote quoted, links not written', text == want, repr(text))
    p = cm.write_run_csv(wc, os.path.join(tmp, 'w.csv'))
    with open(p, 'rb') as f:
        check('Decision 2: write_run_csv writes exactly run_csv_text, in UTF-8', f.read() == want.encode('utf-8'))


def check_shipped(tmp):
    """G2.1 on the workbooks SiCoPR ships (copied first: Rule 1)."""
    books = engine.shipped_workbooks()
    bad = []
    for i, p in enumerate(books):
        q = os.path.join(tmp, 'shipped_%02d_%s' % (i, os.path.basename(p)))
        shutil.copyfile(used(p), q)
        used(q)
        wc, issues = cm.import_config(q)
        f, err = cm.loaded_fields(wc)
        x, xerr = fields_of_file(q, tmp)
        diff = cm.moved_fields(x, f) if f and x else {'load': (xerr, err)}
        if diff or cm.stops(issues):
            bad.append('%s: %d fields differ %s' % (os.path.basename(p), len(diff), list(diff)[:3]))
    return len(books), bad


# ----------------------------------------------------------------------------- the starting workbooks
def check_g21_starting(fx, tmp):
    bad = []
    counts = []
    for label in fixtures.CASES:
        wc, issues = cm.import_config(used(fx.configs[label]))
        f, err = cm.loaded_fields(wc)
        x, xerr = fields_of_file(fx.configs[label], tmp)
        diff = cm.moved_fields(x, f) if f and x else {'load': (xerr, err)}
        counts.append(len(f or {}))
        if diff or issues:
            bad.append('%s: %s %s' % (label, list(diff)[:3], [i['text'] for i in issues]))
    return counts, bad


def st06_linked(rows):
    row = next(r for r in rows if r['id'] == 'ST06')
    out = {}
    for item in row['linked_changes'].split('; '):
        name, rest = item.split(' ', 1)
        out[name] = rest.split(' -> ')[1]
    return out


def check_g22(fx, rows, tmp):
    import scoring
    want = st06_linked(rows)
    check('G2.2 row ST06 lists the seven values f_b moves', sorted(want) == sorted(LINKED_BY_F_B), sorted(want))
    for label in fixtures.CASES:
        wc, issues = cm.import_config(used(fx.configs[label]))
        check('G2.2 %s imports with 18 links' % label, len(wc['links']) == 18, len(wc['links']))
        res = cm.apply_changes(wc, {'/main/f_b': 53.125})
        got = {c['json_path'].rsplit('/', 1)[1]: c['after'] for c in res['changes'] if c['kind'] == 'linked'}
        wrong = [n for n in want if n not in got or not scoring.same(got[n], want[n])]
        check('G2.2 %s: f_b = 53.125 recomputes %s to the values of row ST06' % (label, ', '.join(LINKED_BY_F_B)),
              sorted(got) == sorted(want) and not wrong, 'got %s' % got)

    # A formula outside the grammar, in a workbook: refused at import with a warning naming it; f_z keeps its
    # stored value; nothing evaluates it, even when it has no stored result and would have to be recomputed.
    src = used(fx.configs['Case1'])
    caret = edit_sheet_xml(src, os.path.join(tmp, 'caret.xlsx'), 'B28', 'B3^2')
    with NoEval() as ne:
        wc, issues = cm.import_config(caret)
        res = cm.apply_changes(wc, {'/main/f_b': 53.125})
    warn = [i for i in issues if i['code'] == 'link_refused' and i['level'] == 'warn']
    check('G2.2 a workbook formula outside the grammar (=B3^2) is refused with a warning naming it',
          'f_z' not in wc['links'] and len(warn) == 1 and 'B3^2' in warn[0]['text'] and len(wc['links']) == 17,
          [i['text'] for i in issues])
    check('G2.2 the refused f_z keeps its stored value and does not follow f_b',
          wc['main']['f_z'] == 42.5 and res['config']['main']['f_z'] == 42.5)
    check('G2.2 nothing called eval, exec or compile during that import and apply', not ne.calls, ne.calls)

    marker = os.path.join(tmp, 'pwned')
    hostile = "__import__('os').mkdir('%s'+B3)" % marker.replace('\\', '/')
    lost = edit_sheet_xml(src, os.path.join(tmp, 'hostile.xlsx'), 'B28', hostile, keep_stored=False)
    with NoEval() as ne:
        wc, issues = cm.import_config(lost)
    codes = sorted(i['code'] for i in issues)
    check('G2.2 a hostile formula with no stored result is refused and the import stops: %s' % codes,
          codes == ['link_refused', 'not_recomputable'] and 'f_z' not in wc['links'])
    check('G2.2 the hostile formula never ran (no folder made, no eval, exec or compile)',
          not os.path.exists(marker) and not any(n.startswith('pwned') for n in os.listdir(tmp)) and not ne.calls,
          ne.calls)

    # The same from a hand-edited working config: the link is dropped with a warning, the value kept.
    wc, _ = cm.import_config(src)
    wc['links']['f_z'] = "__import__('os').mkdir('%s')" % marker.replace('\\', '/')
    jp = os.path.join(tmp, 'hostile.json')
    with open(jp, 'w', encoding='utf-8') as f:
        json.dump(wc, f)
    with NoEval() as ne:
        wj, issues = cm.import_config(jp)
        cm.apply_changes(wj, {'/main/f_b': 53.125})
    check('G2.2 a link outside the grammar in a .json is refused with a warning, and nothing evaluates it',
          'f_z' not in wj['links'] and [i['code'] for i in issues] == ['link_refused'] and not ne.calls
          and not os.path.exists(marker), [i['text'] for i in issues] + ne.calls)


def g23_rows(rows):
    import l4
    out = []
    for r in rows:
        if r['group'] == 'sweep' or 'loader' not in r['verified']:
            continue
        paths = [p for p in r['json_path'].split(';') if p.strip()]
        if len(paths) != 1 or ' or ' in paths[0] or not r['value_after'].strip():       # RS28 names two paths
            continue
        if len([x for x in l4.split_top(r['value_after'], '; ') if x.strip()]) != 1:
            continue
        out.append(r)
    return out


def expected_effect(row):
    exp = row['loader_effect'].strip()
    exp = exp.replace(' (accepted by the loader)', '')          # a remark, not part of the effect (C5, IV01, IV02)
    if row['id'] == 'SC14':
        exp = re.sub(r'(matlab_version: )\S+( ->)', r'\g<1>%s\g<2>' % engine.default_release(), exp)
    if exp == 'no loaded field changes':
        exp = ''
    if exp.startswith('the loader stops: '):
        exp = 'loader error: ' + exp[len('the loader stops: '):]
    return exp


def norm_effect(row, text):
    if row['id'] == 'PA013':
        return re.sub(r'[0-9]{2}-[A-Z][a-z]{2}-[0-9]{4}', '<date>', text).replace('\\', '/')
    return text


def effect_text(fx, base, new):
    O = fixtures.oracle()
    nf, nerr = fx.load_csv_bytes(cm.run_csv_text(new).encode('utf-8'))
    if nerr:
        return 'loader error: ' + nerr
    bf, berr = fx.load_csv_bytes(cm.run_csv_text(base).encode('utf-8'))
    if berr:
        return 'base does not load: ' + berr
    return ' | '.join(O.diff(bf, nf))


def check_g23(fx, rows):
    import prior_states
    import scoring
    sel = g23_rows(rows)
    by_id = {}
    for r in rows:
        by_id.setdefault(r['id'], []).append(r)
    imports = {}
    pairs = matched = 0
    misses, excused = [], []
    t0 = time.time()
    for r in sel:
        if not scoring.applies(r)[0]:
            continue
        exp = expected_effect(r)
        for label in scoring.labels_of(r):
            pairs += 1
            ctx = prior_states.Ctx(fx, label, r, by_id)
            try:
                prior_states.build(r['prior_state'], ctx)
                bases = [('prior state', ctx.prior_wc)]
            finally:
                fx.remove_project(ctx.root)
            if r['prior_state'] in ('run complete with the starting config and data\\thru.s4p', 'none; both files exist'):
                if label not in imports:
                    imports[label] = cm.import_config(fx.configs[label])[0]
                bases.append(('product import', imports[label]))
            ok = True
            for how, base in bases:
                res = cm.apply_changes(base, {r['json_path'].strip(): table_value(r['value_after'])})
                got = effect_text(fx, base, res['config'])
                if norm_effect(r, got) != norm_effect(r, exp):
                    ok = False
                    miss = '%s %s (%s): expected %r, got %r' % (r['id'], label, how, exp[:120], got[:120])
                    (excused if r['id'] in DISPUTED else misses).append(miss)
            matched += ok
    print('G2.3: %d rows (%d apply to this engine), %d row and config pairs, %d matched (%.0f s)'
          % (len(sel), sum(1 for r in sel if scoring.applies(r)[0]), pairs, matched, time.time() - t0))
    for e in excused:
        print('G2.3 excused, row disputed: %s' % e)
    return sel, pairs, matched, misses, excused


def check_g24(fx, tmp):
    for label in fixtures.CASES:
        src = os.path.join(tmp, 'g24_%s.xlsx' % label)
        shutil.copyfile(used(fx.configs[label]), src)
        used(src)
        wc, _ = cm.import_config(src)
        res = cm.apply_changes(wc, {'/main/f_b': 53.125})
        new = res['config']
        out = os.path.join(tmp, 'g24_%s_out.xlsx' % label)
        rep = xlsx_io.export_workbook(new, src, out)
        x, xerr = fields_of_file(out, tmp)
        f, ferr = cm.loaded_fields(new)
        check('G2.4 %s: the exported workbook loads to the same fields as the working config' % label,
              x is not None and f is not None and not cm.moved_fields(f, x), xerr or list(cm.moved_fields(f, x))[:5])
        try:
            before, after = sheet_cells(src), sheet_cells(out)
            opened = True
        except Exception as e:                                          # the check is that it opens
            opened, before, after = repr(e), {}, {}
        check('G2.4 %s: the exported workbook opens through openpyxl' % label, opened is True, opened)
        written = {w['cell'] for w in rep['written']}
        kept = [c for c in before if c not in written and before[c] != after.get(c)]
        check('G2.4 %s: every cell not changed keeps its formula and stored result (%d cells)' % (label, len(before)),
              not kept, kept[:5])
        cells = {w['json_path']: w for w in rep['written']}
        linked = [c for c in res['changes'] if c['kind'] == 'linked']
        bad = [c['json_path'] for c in linked if c['json_path'] not in cells
               or isinstance(after[cells[c['json_path']]['cell']][0], str) and after[cells[c['json_path']]['cell']][0].startswith('=')
               or not cm.same_value(after[cells[c['json_path']]['cell']][1], c['after'])]
        check('G2.4 %s: each linked cell whose value changed (%d) is written as that value, not a formula'
              % (label, len(linked)), len(linked) == 7 and not bad and cells['/main/f_z']['was_formula'], bad)
        with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
            sheet = xlsx_io._sheet_xml_name(src, 'COM_Settings')
            other = [n for n in a.namelist() if n != sheet and a.read(n) != (b.read(n) if n in b.namelist() else None)]
        check('G2.4 %s: every other part of the workbook is byte for byte the source (License Notice included)' % label,
              not other, other)

    # CP07 (stretch): an added keyword loads with its field, and its row's cells are in column order. Case1
    # holds empty styled cells in row 33, so the same is done with those two cells taken out of the sheet XML
    # (the new cells go into an existing row) and with row 33 taken out (a new row goes in).
    base = os.path.join(tmp, 'g24_Case1.xlsx')
    sheet = xlsx_io._sheet_xml_name(base, 'COM_Settings')
    with zipfile.ZipFile(base) as z:
        names = z.namelist()
        blobs = {n: z.read(n) for n in names}
    xml0 = blobs[sheet].decode('utf-8')
    variants = {'as it is': xml0,
                'A33 and B33 absent': re.sub(r'<c r="[AB]33"[^>]*/>', '', xml0),
                'row 33 absent': re.sub(r'<row[^>]*\sr="33"[^>]*>.*?</row>', '', xml0, flags=re.S)}
    for how, xml_src in variants.items():
        src = os.path.join(tmp, 'g24_c2_src_%d.xlsx' % len(xml_src))
        with zipfile.ZipFile(src, 'w', zipfile.ZIP_DEFLATED) as z:
            for n in names:
                z.writestr(n, xml_src.encode('utf-8') if n == sheet else blobs[n])
        used(src)
        wc, _ = cm.import_config(src)
        res = cm.apply_changes(wc, {'/main/c(2)': '[-0.1:.02:0]'})
        out = os.path.join(tmp, 'g24_c2_out_%d.xlsx' % len(xml_src))
        rep = xlsx_io.export_workbook(res['config'], src, out)
        x, xerr = fields_of_file(out, tmp)
        s, _ = fields_of_file(src, tmp)
        d = cm.moved_fields(s, x) if x else {}
        check('G2.4 CP07 (%s): c(2) is added as a keyword and value above the first .START row, and loads with its '
              'field' % how, [c['kind'] for c in res['changes']] == ['add'] and len(rep['added']) == 1
              and list(d) == ['param.tx_ffe_cp2_values'] and d['param.tx_ffe_cp2_values'][0] is cm.ABSENT,
              (rep['added'], list(d), xerr))
        a = rep['added'][0] if rep['added'] else {'label_cell': 'A1', 'cell': 'B1'}
        row = xlsx_io.place(a['cell'])[0] + 1
        with zipfile.ZipFile(out) as z:
            xml = z.read(sheet).decode('utf-8')
        m = re.search(r'<row[^>]*\sr="%d"[^>]*>(.*?)</row>' % row, xml, re.S)
        cols = [xlsx_io._col_index(c) for c in re.findall(r'<c r="([A-Z]+)[0-9]+"', m.group(1))] if m else []
        rows_order = [int(n) for n in re.findall(r'<row[^>]*\sr="([0-9]+)"', xml)]
        check('G2.4 CP07 (%s): the cells of row %d are in column order, and the rows in row order' % (how, row),
              cols and cols == sorted(cols) and len(set(cols)) == len(cols) and rows_order == sorted(rows_order),
              (cols[:6], rows_order[28:36]))
        wb = openpyxl.load_workbook(out, data_only=True)['COM_Settings']
        check('G2.4 CP07 (%s): the new pair reads back as c(2) and its value' % how, wb[a['label_cell']].value == 'c(2)'
              and wb[a['cell']].value == '[-0.1:.02:0]', (wb[a['label_cell']].value, wb[a['cell']].value))


def check_export_shared_and_chain(fx, tmp):
    """A shared formula whose cell is overwritten keeps its other cells as formulas; a workbook with a
    calculation chain loses the chain when a formula becomes a value. Both still load like the working
    config."""
    # f_z and f_p1 of Case1 made one shared formula (=$B$3/2.5 over B28:B29, first cell B28); f_z then set by
    # hand. f_p1 still follows f_b, so it must stay a formula, now written out in full.
    src = os.path.join(tmp, 'shared_src.xlsx')
    sheet = xlsx_io._sheet_xml_name(used(fx.configs['Case1']), 'COM_Settings')
    with zipfile.ZipFile(fx.configs['Case1']) as z:
        names = z.namelist()
        blobs = {n: z.read(n) for n in names}
    x = blobs[sheet].decode('utf-8')
    x2 = x.replace('<f>B3/2.5</f><v>42.5</v></c><c r="C28"', '<f t="shared" ref="B28:B29" si="90">$B$3/2.5</f><v>42.5</v></c><c r="C28"', 1)
    x2 = re.sub(r'(<c r="B29"[^>]*>)<f>B3/2.5</f>', lambda m: m.group(1) + '<f t="shared" si="90"/>', x2, count=1)
    blobs[sheet] = x2.encode('utf-8')
    with zipfile.ZipFile(src, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, blobs[n])
    used(src)
    check('export: the shared-formula workbook was made (B28:B29 share =$B$3/2.5)',
          x2.count('si="90"') == 2 and sheet_cells(src)['B29'][0] == '=$B$3/2.5')
    wc, _ = cm.import_config(src)
    res = cm.apply_changes(wc, {'/main/f_z': 40})
    out = os.path.join(tmp, 'shared_out.xlsx')
    xlsx_io.export_workbook(res['config'], src, out)
    after = sheet_cells(out)
    x, xerr = fields_of_file(out, tmp)
    f, _ = cm.loaded_fields(res['config'])
    check('export: overwriting the first cell of a shared formula keeps its other cell a formula with its stored result',
          after['B28'] == (40, 40) and after['B29'] == ('=$B$3/2.5', 42.5) and x and not cm.moved_fields(f, x),
          (after['B28'], after['B29'], xerr))

    books = [p for p in engine.shipped_workbooks() if os.path.basename(p).endswith('Case1.xlsx')]
    src = os.path.join(tmp, 'chain_src.xlsx')
    shutil.copyfile(books[0], src)
    used(src)
    with zipfile.ZipFile(src) as z:
        had = 'xl/calcChain.xml' in z.namelist()
    wc, _ = cm.import_config(src)
    res = cm.apply_changes(wc, {'/main/f_b': 53.125})
    out = os.path.join(tmp, 'chain_out.xlsx')
    xlsx_io.export_workbook(res['config'], src, out)
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        ct = z.read('[Content_Types].xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    x, xerr = fields_of_file(out, tmp)
    f, _ = cm.loaded_fields(res['config'])
    check('export: a calculation chain is dropped with its references when a formula becomes a value',
          had and 'xl/calcChain.xml' not in names and 'calcChain' not in ct and 'calcChain' not in rels
          and x and not cm.moved_fields(f, x), (had, xerr))


def check_g25(fx, rows, tmp):
    src = used(fx.configs['Case1'])
    wc, issues = cm.import_config(src)
    stored = {c: v for c, (_, v) in sheet_cells(src).items()}
    cells = xlsx_io.scan(xlsx_io.read_workbook(src).grid, xlsx_io.read_workbook(src).formulas)
    formula_keys = {k: cl.cell for k, cl in cells.main.items() if cl.formula}
    check('G2.5 CP03: each formula cell of a keyword keeps its stored value exactly (f_v = %r)' % wc['main']['f_v'],
          formula_keys and all(wc['main'][k] == xlsx_io.stored_value(stored[c]) for k, c in formula_keys.items())
          and wc['main']['f_v'] == 0.3709803922, formula_keys)
    check('G2.5 CP03: 18 links by keyword name, such as f_z = {f_b}/2.5',
          len(wc['links']) == 18 and wc['links']['f_z'] == '{f_b}/2.5'
          and wc['links']['packages.PKG_LowR_CLASSA.z_p (NEXT)'] == '{z_p (TX)}')
    with open(os.path.join(fixtures.ROOT, 'sicopr_assist', 'data', 'sicopr_config_example.json'), encoding='utf-8') as f:
        example = json.load(f)
    check('G2.5 the Case1 import equals sicopr_config_example.json, key order included',
          wc == example and list(wc['main']) == list(example['main']))

    res = cm.apply_changes(wc, {'/main/f_b': 53.125})
    moved = sorted(c['json_path'] for c in res['changes'])
    others = [k for k in wc['main'] if k not in LINKED_BY_F_B + ('f_b',) and res['config']['main'][k] != wc['main'][k]]
    check('G2.5 CP04: f_b recomputes the seven links downstream, every other value stays as imported',
          moved == sorted(['/main/f_b'] + ['/main/' + n for n in LINKED_BY_F_B]) and not others
          and res['config']['packages'] == wc['packages'] and res['config']['links'] == wc['links'], moved)

    r5 = cm.apply_changes(wc, {'/main/f_z': 40})
    r6 = cm.apply_changes(r5['config'], {'/main/f_b': 53.125})
    check('G2.5 CP05: f_z = 40 removes its link (and says so); a later f_b change no longer moves f_z',
          'f_z' not in r5['config']['links'] and r5['links_removed'] == [{'json_path': '/main/f_z', 'link': '{f_b}/2.5'}]
          and r6['config']['main']['f_z'] == 40 and r6['config']['main']['f_p1'] == 21.25)

    same = cm.apply_changes(wc, {'/main/DER_0': 0.0002, '/main/f_z': 42.5})
    check('Decision 1: a value equal to the current one is not a change, and does not drop a link',
          not same['changes'] and len(same['unchanged']) == 2 and same['config'] == wc)

    nb = cm.apply_changes(wc, {'/packages/NoSuchBlock/C_p': 1})
    check('apply: a change in a package block that does not exist is refused, not added',
          nb['refused'] and nb['config'] == wc)

    lf13 = cm.apply_changes(wc, {'/main/C_p': '[0.2e-4 0.2e-4]'}, check=True)
    lf14 = cm.apply_changes(wc, {'/main/b_max(2..N_b)': 0.2}, check=True)
    t_t = cm.apply_changes(wc, {'/main/T_t': 5}, check=True)
    check('LF13: a main-area C_p with PKG_NAME set has no effect, and the rule says why',
          (lf13['no_effect'] or {}).get('rule') == 'LF13' and lf13['changes'][0]['kind'] == 'add', lf13.get('no_effect'))
    check('LF14: b_max(2..N_b) with N_b = 1 has no effect, and the rule says why',
          (lf14['no_effect'] or {}).get('rule') == 'LF14', lf14.get('no_effect'))
    check('a change to T_t, which the loader does not read, moves f_v through its link (row PA018)',
          t_t['no_effect'] is None and 'param.f_v' in t_t['effect']['moved'], t_t.get('effect'))
    # VR01: T_dev is read only under the newest release; under an earlier one a change to it moves nothing.
    # The rule is named when the engine's keyword list lacks the keyword (an engine without T_dev, simulated).
    older = [r for r in engine.supported_releases() if r != engine.default_release()]
    w_old = cm.apply_changes(wc, {'/main/COM Version': older[0]})['config']
    plain = cm.apply_changes(w_old, {'/main/T_dev': 2}, check=True)
    real = engine.reads_keyword
    engine.reads_keyword = lambda name, scope=None: False if name == 'T_dev' else real(name, scope)
    try:
        vr = cm.apply_changes(w_old, {'/main/T_dev': 2}, check=True)
    finally:
        engine.reads_keyword = real
    check('VR01: T_dev under %s moves nothing and is reported as no effect; with an engine that lacks T_dev the rule '
          'says why' % older[0], plain['no_effect'] and plain['no_effect']['rule'] is None
          and (vr['no_effect'] or {}).get('rule') == 'VR01', (plain.get('no_effect'), vr.get('no_effect')))

    def issues_of(label):
        return cm.import_config(used(fx.configs[label]))

    w, iss = issues_of('V03')
    check('G2.5 V03 (ST05): the import asks about the label P_peek in J31, which is nearly the absent P_peak',
          [(i['level'], i['code'], i['label'], i['cell'], i['keyword']) for i in iss]
          == [('ask', 'near_keyword', 'P_peek', 'J31', 'P_peak')], iss)
    w, iss = issues_of('V10')
    check('G2.5 V10 (ST14): the import stops on DER_0 twice and names both places',
          [(i['level'], i['code'], i['keyword'], i['cells']) for i in iss] == [('stop', 'duplicate_keyword', 'DER_0', ['J4', 'J33'])],
          iss)
    w, iss = issues_of('V11')
    check('G2.5 V11 (ST15): the import stops on the missing f_b and names it',
          [(i['level'], i['code'], i['keyword']) for i in iss] == [('stop', 'missing_required', 'f_b')], iss)
    w, iss = issues_of('V12')
    f12, e12 = cm.loaded_fields(w)
    rec = [i for i in iss if i['code'] == 'recomputed']
    close = all(isinstance(w['main'][k], (int, float)) and abs(w['main'][k] - wc['main'][k]) <= 1e-9 * abs(wc['main'][k])
                for k in LINKED_BY_F_B)
    check('G2.5 V12 (ST16): the 18 linked values with no stored result are recomputed and announced, and the import loads',
          [i['level'] for i in iss] == ['announce'] and len(rec[0]['linked']) == 18 and close and f12 and len(f12) == 315,
          (iss, e12))
    w, iss = issues_of('V06')
    check('G2.5 V06 (ST08, ST09): no package blocks is noted, the package parameters are main parameters, and A_fe '
          'follows A_v', [i['code'] for i in iss] == ['no_package_blocks'] and w['packages'] == {}
          and w['links'].get('A_fe') == '{A_v}' and 'C_p' in iss[0]['keywords'], iss)
    r = cm.apply_changes(w, {'/main/A_v': 0.4})
    check('G2.5 V06 (ST08): A_v = 0.4 recomputes A_fe', [(c['json_path'], c['after']) for c in r['changes']]
          == [('/main/A_v', 0.4), ('/main/A_fe', 0.4)])

    # .json (CP14, CP15) and .csv (CP16)
    jp = os.path.join(tmp, 'config.json')
    with open(jp, 'w', encoding='utf-8') as f:
        json.dump(wc, f)
    w, iss = cm.import_config(jp)
    check('G2.5 CP14: a working-config .json is accepted as it is', w == wc and not iss, iss)
    bad = copy.deepcopy(wc)
    bad['main'] = {('f_bb' if k == 'f_b' else k): v for k, v in bad['main'].items()}
    with open(jp, 'w', encoding='utf-8') as f:
        json.dump(bad, f)
    w, iss = cm.import_config(jp)
    unk = [i for i in iss if i['code'] == 'unknown_key']
    check('G2.5 CP15: a .json with f_bb for f_b names the unknown key and offers f_b; the links that name f_b are kept',
          len(unk) == 1 and unk[0]['json_path'] == '/main/f_bb' and unk[0]['near'][:1] == ['f_b']
          and unk[0]['level'] == 'ask' and w['links'] == wc['links']
          and sorted(i['code'] for i in iss) == ['link_unresolved', 'missing_required', 'unknown_key'], iss)
    bad = dict(copy.deepcopy(wc), schema='sicopr-config/0')
    with open(jp, 'w', encoding='utf-8') as f:
        json.dump(bad, f)
    w, iss = cm.import_config(jp)
    check('G2.5 CP14: a .json of another schema version stops', w is None and [i['code'] for i in iss] == ['json_form'], iss)
    cp = cm.write_run_csv(wc, os.path.join(tmp, 'config.csv'))
    w, iss = cm.import_config(cp)
    read = {k: v for k, v in wc['main'].items() if engine.reads_keyword(k)}
    fw, _ = cm.loaded_fields(w)
    f0, _ = cm.loaded_fields(wc)
    check('G2.5 CP16: a .csv config imports to the same form and order, and loads the same; it has no formulas, so no '
          'links, and labels only a formula used (T_t, T_ft, T_nt) are not keys',
          w['main'] == read and list(w['main']) == list(read) and w['packages'] == wc['packages'] and w['links'] == {}
          and not iss and fw and not cm.moved_fields(f0, fw), (iss, sorted(set(wc['main']) - set(w['main']))))

    # .xls: an .xlsx under that name is read (F1); a true Excel 97-2003 file stops (IV21)
    renamed = os.path.join(tmp, 'config.xls')
    shutil.copyfile(src, renamed)
    w, iss = cm.import_config(renamed)
    check('G2.5 F1: an .xlsx named .xls is read from a temporary .xlsx copy, with a note',
          w is not None and w['main'] == wc['main'] and w['links'] == wc['links'] and [i['code'] for i in iss] == ['xls_is_xlsx'], iss)
    legacy = os.path.join(tmp, 'legacy.xls')
    with open(legacy, 'wb') as f:
        f.write(bytes.fromhex('d0cf11e0a1b11ae1') + bytes(504))
    w, iss = cm.import_config(legacy)
    check('G2.5 IV21: a true Excel 97-2003 file stops and asks for an .xlsx',
          w is None and [i['code'] for i in iss] == ['legacy_xls'] and '.xlsx' in iss[0]['text'], iss)

    # export refuses the source path, however it is spelled, and an existing file
    copy_src = os.path.join(tmp, 'refuse.xlsx')
    shutil.copyfile(src, copy_src)
    used(copy_src)
    refused, other_errors = [], []
    for target in (copy_src, os.path.join(tmp, '.', 'refuse.xlsx'), os.path.join(tmp, 'sub', '..', 'refuse.xlsx'),
                   copy_src.upper() if os.name == 'nt' else copy_src):
        try:
            xlsx_io.export_workbook(wc, copy_src, target, overwrite=True)
        except ValueError:
            refused.append(target)
        except Exception as e:                                          # anything else is not the refusal
            other_errors.append('%s: %r' % (target, e))
    check('G2.5 export refuses to write over the source config (4 spellings of its path, even with overwrite set)',
          len(refused) == 4 and sha(copy_src) == SOURCES[os.path.abspath(copy_src)], (refused, other_errors))
    other = os.path.join(tmp, 'exists.xlsx')
    shutil.copyfile(src, other)
    try:
        xlsx_io.export_workbook(wc, copy_src, other)
        check('export refuses an existing file unless told to overwrite it', False)
    except FileExistsError:
        check('export refuses an existing file unless told to overwrite it', True)


def check_harness_hooks(fx, rows):
    """The Phase 1 harness's own checks of CP02, CP03, CP06, CP07 and ST16 call the product by name."""
    import scoring
    from sicopr_assist import core
    sc = scoring.Scorer(fx, rows, core.handle)
    out = []
    for rid, label in (('CP02', 'Case1'), ('CP03', 'Case2'), ('CP06', 'Case3'), ('CP07', 'Case1'), ('ST16', 'V12')):
        row = next(r for r in rows if r['id'] == rid)
        ctx = sc.context(row, label)
        try:
            ok, note = scoring.SPECIAL_L3[rid](sc, row, ctx, core.outcome('PLAN'))
        finally:
            fx.remove_project(ctx.root)
        out.append(rid)
        check('harness L3 check of %s on %s passes with the product' % (rid, label), ok, note)


PARTS = ('grammar', 'csv', 'shipped', 'g21', 'g22', 'g23', 'g24', 'export', 'g25', 'hooks')


def main(argv=None):
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('--only', help='comma-separated parts to run, of: %s (default: all)' % ', '.join(PARTS))
    args = ap.parse_args(argv)
    parts = set(args.only.split(',')) if args.only else set(PARTS)
    t0 = time.time()
    tmp = tempfile.mkdtemp(prefix='sicopr_assist_test_config_')
    try:
        if 'grammar' in parts:
            check_grammar()
        if 'csv' in parts:
            check_csv_writer(tmp)
        if 'shipped' in parts:
            n, bad = check_shipped(tmp)
            check('G2.1 the %d workbooks SiCoPR ships: import, run CSV, load; no field differs from the workbook' % n,
                  n >= 9 and not bad, bad)
        if not parts & {'g21', 'g22', 'g23', 'g24', 'export', 'g25', 'hooks'}:
            pass
        elif not fixtures.find_starting_configs():
            print('SKIP: the four starting workbooks (by SHA-256, in tests\\configs or SICOPR_ASSIST_CONFIGS) are absent '
                  'or hidden; G2.1 on them, G2.2 to G2.5 need them')
        else:
            import l4
            rows = l4.load_table()
            with fixtures.Fixtures() as fx:
                for p in fx.configs.values():
                    used(p)
                if 'g21' in parts:
                    counts, bad = check_g21_starting(fx, tmp)
                    check('G2.1 the four starting workbooks: no field differs, no import issue (%s fields)' % counts,
                          not bad and counts == [315] * 4, bad)
                if 'g22' in parts:
                    check_g22(fx, rows, tmp)
                if 'g23' in parts:
                    sel, pairs, matched, misses, excused = check_g23(fx, rows)
                    n_exc = len({e.split(' (')[0] for e in excused})
                    check('G2.3 %d of %d row and config pairs reproduce their loader_effect; %d pairs of disputed rows '
                          '(%s) excused' % (matched, pairs, n_exc, ', '.join(sorted(DISPUTED))),
                          pairs > 0 and matched + n_exc == pairs and not misses, misses[:8])
                if 'g24' in parts:
                    check_g24(fx, tmp)
                if 'export' in parts:
                    check_export_shared_and_chain(fx, tmp)
                if 'g25' in parts:
                    check_g25(fx, rows, tmp)
                if 'hooks' in parts:
                    check_harness_hooks(fx, rows)
        changed = [p for p, h in SOURCES.items() if os.path.exists(p) and sha(p) != h]
        check('G2.6 the SHA-256 of every source workbook used (%d, the workbooks SiCoPR ships included) is unchanged'
              % len(SOURCES), not changed, changed)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print('%.0f s' % (time.time() - t0))
    finish()


if __name__ == '__main__':
    sys.exit(main())
