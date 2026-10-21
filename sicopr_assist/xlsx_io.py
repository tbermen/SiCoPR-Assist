"""Reading a COM workbook, and writing one for export.

This module starts from SiCoPR's gui/config_io.py at commit 55c3f717bf8e17daffe65ecf20215a4fbcdb0833
(SiCoPR 1.1.1, BSD-3-Clause, Copyright 2026 Todd Bermensolo). read_config, write_config and the helpers
they need (_col_letters, _region_keywords, _read_schema, _sheet_xml_name, _fmt, _rewrite_cells) were
copied and then adapted. The gui package is not imported.

What the engine requires, as config_io.py explains it: the loader scans a region of the COM_Settings
grid for a cell whose stripped, lower-cased text equals the keyword and takes the cell to its right as
the value. The main region is everything above the first .START row in column A; each .START ... .END
block is a region of its own. A keyword found twice in one region makes the loader raise.

Why export does not save through openpyxl: nine value cells of a COM workbook are formulas, and openpyxl
can keep a formula or its stored result, never both. The loader reads stored results, so a workbook
saved by openpyxl no longer loads. Export therefore copies the source and edits the sheet XML in the
copy: only the cells that changed are rewritten, and every other byte of the cell data stays as it was.

Adapted from the original:
  - reading reports a formula with no stored result, a keyword found twice in a region, and a label
    that is nearly a keyword and is not one;
  - only keywords the engine's loader reads count as keywords (the list comes from engine.py), and a
    label that a formula depends on is taken in by its label;
  - a file named .xls is looked at before it is opened: an .xlsx under another name is read from a
    temporary copy, a true Excel 97-2003 file is refused;
  - the grid is read from row 1 and column A, so a cell's place in the grid is its place in the sheet;
  - a cell added to the sheet is put in column order within its row (and a missing row in row order),
    so Excel opens the file without a repair prompt;
  - overwriting a cell that belongs to a shared formula writes the other cells of that formula out in
    full, and overwriting a formula drops the calculation chain, which Excel rebuilds.
"""
# Copyright 2026 Todd Bermensolo
# SPDX-License-Identifier: BSD-3-Clause
import difflib
import os
import re
import shutil
import tempfile
import zipfile
from types import SimpleNamespace

import openpyxl

from sicopr_assist import engine

SHEET = 'COM_Settings'
ZIP_SIGNATURE = b'PK\x03\x04'
OLE_SIGNATURE = bytes.fromhex('d0cf11e0a1b11ae1')        # Excel 97-2003 (.xls) compound file

# A cell reference in a formula: A1, $A$1. Not part of a name (LOG10), a sheet reference (Sheet2!A1)
# or a number (1E3).
REF = re.compile(r"(?<![A-Za-z0-9_$.!'])\$?([A-Z]{1,3})\$?([0-9]+)(?![A-Za-z0-9_(!])")


class LegacyXls(ValueError):
    """A true Excel 97-2003 file: neither the engine nor this tool can read it."""


class NotAWorkbook(ValueError):
    """A file that is not an .xlsx workbook."""


class StrictOOXML(ValueError):
    """An .xlsx saved as Excel's Strict Open XML Spreadsheet: openpyxl, and so this tool, and SiCoPR cannot read it
    (core fix 4, row CF11; the IEEE COM 3.70 workbooks are such files)."""


STRICT_NAMESPACE = b'purl.oclc.org/ooxml'


def is_strict_ooxml(path):
    """Whether an .xlsx package is strict OOXML: its workbook part uses the purl.oclc.org/ooxml namespaces."""
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            part = 'xl/workbook.xml' if 'xl/workbook.xml' in names else next(
                (n for n in names if n.endswith('workbook.xml')), None)
            return part is not None and STRICT_NAMESPACE in z.read(part)[:4000]
    except (zipfile.BadZipFile, OSError, KeyError):
        return False


def _col_letters(idx0):
    """0-based column index -> spreadsheet letters."""
    s, n = '', idx0 + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _col_index(letters):
    """Spreadsheet letters -> 0-based column index."""
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n - 1


def coord(r, c):
    """0-based row and column -> 'B3'."""
    return '%s%d' % (_col_letters(c), r + 1)


def place(cell):
    """'B3' -> (2, 1), 0-based row and column."""
    m = re.fullmatch(r'([A-Z]+)([0-9]+)', cell)
    return int(m.group(2)) - 1, _col_index(m.group(1))


def stored_value(v):
    """A cell value as the working config keeps it: a whole float as an int (as the loader reads it),
    anything else exactly as stored."""
    if isinstance(v, float) and v.is_integer() and abs(v) < 2 ** 53:
        return int(v)
    return v


# ----------------------------------------------------------------------------- reading
def _readable_copy(path):
    """(path to open, temporary folder or None, note or None). An .xlsx under another name is read from
    a temporary copy named .xlsx; a true .xls is refused."""
    with open(path, 'rb') as f:
        head = f.read(8)
    if head.startswith(OLE_SIGNATURE):
        raise LegacyXls('%s is an Excel 97-2003 (.xls) file, which cannot be read; open it in Excel and '
                        'save it as an .xlsx workbook' % os.path.basename(path))
    if not head.startswith(ZIP_SIGNATURE):
        raise NotAWorkbook('%s is not an .xlsx workbook' % os.path.basename(path))
    if os.path.splitext(path)[1].lower() in ('.xlsx', '.xlsm'):
        return path, None, None
    tmp = tempfile.mkdtemp(prefix='sicopr_assist_xlsx_')
    copy = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + '.xlsx')
    shutil.copyfile(path, copy)
    return copy, tmp, '%s is an .xlsx workbook under another name; it was read as one' % os.path.basename(path)


def read_workbook(path):
    """The COM_Settings sheet as two grids: stored values and formulas.

    Returns a namespace with path, sheet, sheetnames, grid (stored values, grid[r][c] is the cell in sheet
    row r+1 and column c+1), formulas ({'B28': 'B3/2.5'}, without the '='), schema (the keywords sheet)
    and note (set when an .xlsx was read under another name). Raises LegacyXls or NotAWorkbook."""
    readable, tmp, note = _readable_copy(path)
    if is_strict_ooxml(readable):
        if tmp:
            shutil.rmtree(tmp, ignore_errors=True)
        raise StrictOOXML("%s is a strict OOXML workbook (Excel's Strict Open XML Spreadsheet), which SiCoPR-Assist "
                          "and SiCoPR cannot read: save it as an ordinary Excel Workbook (.xlsx) and run again"
                          % os.path.basename(path))
    try:
        wb_val = openpyxl.load_workbook(readable, data_only=True)
        wb_fml = openpyxl.load_workbook(readable, data_only=False)
    finally:
        if tmp:
            shutil.rmtree(tmp, ignore_errors=True)
    name = SHEET if SHEET in wb_val.sheetnames else wb_val.sheetnames[0]
    ws_val, ws_fml = wb_val[name], wb_fml[name]
    n_rows, n_cols = ws_val.max_row, ws_val.max_column
    grid = [list(r) for r in ws_val.iter_rows(min_row=1, min_col=1, max_row=n_rows, max_col=n_cols,
                                               values_only=True)]
    formulas = {}
    for row in ws_fml.iter_rows(min_row=1, min_col=1, max_row=n_rows, max_col=n_cols):
        for cell in row:
            v = cell.value
            text = v if isinstance(v, str) else getattr(v, 'text', None)    # an array formula has .text
            if isinstance(text, str) and text.startswith('='):
                formulas[cell.coordinate] = text[1:]
    return SimpleNamespace(path=path, sheet=name, sheetnames=list(wb_val.sheetnames), grid=grid,
                           formulas=formulas, schema=_read_schema(wb_val), note=note)


def _read_schema(wb):
    """The keywords_* sheet: the config's own documentation of itself.

    Columns are `matlab variable | keyword | default | info`."""
    sheets = [s for s in wb.sheetnames if s.lower().startswith('keywords')]
    if not sheets:
        return {}
    out = {}
    for row in wb[sheets[0]].iter_rows(values_only=True):
        if not row or len(row) < 2 or not row[1]:
            continue
        kw = str(row[1]).strip()
        if kw.lower() == 'keyword':
            continue
        out[kw] = {
            'matlab': str(row[0]).strip() if row[0] else '',
            'default': row[2] if len(row) > 2 else None,
            'info': str(row[3]).strip() if len(row) > 3 and row[3] else '',
        }
    return out


def keyword_descriptions(path):
    """{keyword: description} from the workbook's keywords sheet; {} when it has none."""
    return {k: v['info'] for k, v in read_workbook(path).schema.items()}


def selection_lists(path):
    """{column heading: [the values below it]} from the workbook's selections sheet, the lists a workbook
    offers for some keywords (the 802.3dj sheets list eight port orders under 'ports'); {} when it has
    none. Validation offers these as options (rule NC03)."""
    readable, tmp, _ = _readable_copy(path)
    try:
        wb = openpyxl.load_workbook(readable, data_only=True)
    finally:
        if tmp:
            shutil.rmtree(tmp, ignore_errors=True)
    sheets = [s for s in wb.sheetnames if s.lower().startswith('selections')]
    if not sheets:
        return {}
    rows = [list(r) for r in wb[sheets[0]].iter_rows(values_only=True)]
    out = {}
    for c, head in enumerate(rows[0] if rows else []):
        if head is None or not str(head).strip():
            continue
        vals = [r[c] for r in rows[1:] if c < len(r) and r[c] is not None and str(r[c]).strip()]
        out[str(head).strip()] = [v.strip() if isinstance(v, str) else v for v in vals]
    return out


def _cell(grid, r, c):
    if 0 <= r < len(grid) and 0 <= c < len(grid[r]):
        return grid[r][c]
    return None


def _issue(level, code, text, **details):
    return dict(level=level, code=code, text=text, **details)


def _region_keywords(grid, formulas, lo, hi, keywords, where, issues):
    """Keyword cells in grid rows lo..hi-1, found the way the loader finds them: the label anywhere in the
    region (stripped, any case), the value in the cell to its right. A keyword found more than once is
    reported (the loader raises on it) and its first place is used."""
    seen = {}
    for r in range(lo, hi):
        for c, cell in enumerate(grid[r]):
            if isinstance(cell, str) and cell.strip():
                seen.setdefault(cell.strip().lower(), []).append((r, c))
    spelled = {k.lower(): k for k in keywords}
    out = {}
    for low, places in seen.items():
        if low not in spelled:
            continue
        r, c = places[0]
        label = grid[r][c].strip()
        if len(places) > 1:
            cells = [coord(*p) for p in places]
            issues.append(_issue('stop', 'duplicate_keyword',
                                 '%s appears %d times %s (%s); the loader stops on this, so remove all but one'
                                 % (label, len(places), where, ' and '.join(cells)), keyword=label, cells=cells))
        out[label] = _keyword(grid, formulas, r, c, label)
    return out


def _keyword(grid, formulas, r, c, label):
    vc = coord(r, c + 1)
    f = formulas.get(vc)
    v = _cell(grid, r, c + 1)
    lost = f is not None and (v is None or v == '')
    return SimpleNamespace(name=label, label_cell=coord(r, c), cell=vc, value=None if lost else stored_value(v),
                           formula=f, lost=lost)


def scan(grid, formulas, main_keywords=None, package_keywords=None, near_cutoff=0.8):
    """Find the keywords of a COM_Settings grid as the loader would, plus the labels formulas depend on.

    Returns a namespace:
      main       {key: cell} in sheet order (column by column), a cell being a namespace with name,
                 label_cell, cell (the value cell), value, formula (text or None) and lost (a formula with
                 no stored result)
      packages   {block: {key: cell}}
      links      {link key: link text}, a formula rewritten with {name} for each cell it refers to; the
                 link key is the key for the main area and packages.<block>.<key> in a block
      blocks     [(block, .START row, .END row)], 0-based rows
      main_rows  (0, first .START row)
      issues     [{level, code, text, ...}], level one of stop, ask, announce, warn, note
    """
    main_keywords = engine.keyword_names('main') if main_keywords is None else main_keywords
    package_keywords = engine.keyword_names('package block') if package_keywords is None else package_keywords
    issues = []
    col_a = [row[0] if row else None for row in grid]
    starts = [i for i, v in enumerate(col_a) if v == '.START']
    ends = [i for i, v in enumerate(col_a) if v == '.END']
    if len(starts) != len(ends):
        issues.append(_issue('stop', 'start_end', 'the sheet has %d .START rows and %d .END rows; the loader stops '
                                                  'on this' % (len(starts), len(ends))))
        return SimpleNamespace(main={}, packages={}, links={}, blocks=[], main_rows=(0, len(grid)), issues=issues)
    main_hi = starts[0] if starts else len(grid)
    blocks = [(str(_cell(grid, s, 1)), s, e) for s, e in zip(starts, ends)]
    main = _region_keywords(grid, formulas, 0, main_hi, main_keywords, 'in the main area', issues)
    packages = {}
    for name, s, e in blocks:
        packages[name] = _region_keywords(grid, formulas, s + 1, e, package_keywords,
                                          'in package block %s' % name, issues)

    def region_of(r):
        if r < main_hi:
            return None
        return next((b for b, s, e in blocks if s < r < e), False)          # False: between blocks

    def scope(block):
        return main if block is None else packages[block]

    by_cell = {cl.cell: (None, k) for k, cl in main.items()}
    for b, cells in packages.items():
        by_cell.update({cl.cell: (b, k) for k, cl in cells.items()})

    # Formulas become links by name. A cell a formula refers to that is not a keyword is taken in by the
    # label to its left (T_t, T_ft and T_nt in the 802.3dj sheets), and its own formula is followed too.
    links = {}
    todo = [(b, k) for b, k in by_cell.values()]
    while todo:
        block, key = todo.pop(0)
        cl = scope(block)[key]
        if cl.formula is None or not REF.search(cl.formula):
            continue                                       # a value, or a constant such as =0.005
        link_key = key if block is None else 'packages.%s.%s' % (block, key)
        bad = []

        def name_of(m):
            ref = m.group(1) + m.group(2)
            if ref not in by_cell:
                r, c = place(ref)
                label = _cell(grid, r, c - 1)
                tb = region_of(r)
                if (tb is not False and isinstance(label, str) and label.strip() and coord(r, c - 1) not in formulas
                        and label.strip() not in scope(tb)):
                    scope(tb)[label.strip()] = _keyword(grid, formulas, r, c - 1, label.strip())
                    by_cell[ref] = (tb, label.strip())
                    todo.append((tb, label.strip()))
                else:
                    bad.append(ref)
                    return m.group(0)
            tb, tk = by_cell[ref]
            if tb == block or (tb is None and (block is None or tk not in packages[block])):
                return '{%s}' % tk
            if tb is not None:
                return '{packages.%s.%s}' % (tb, tk)
            bad.append(ref)
            return m.group(0)

        text = REF.sub(name_of, cl.formula)
        if bad:
            issues.append(_issue('warn', 'link_unlabelled',
                                 'the formula of %s (=%s in %s) refers to %s, which has no label; %s keeps its '
                                 'stored value and is not linked' % (key, cl.formula, cl.cell, ', '.join(bad), key),
                                 key=link_key, formula=cl.formula))
        else:
            links[link_key] = text

    def order(cells):
        return dict(sorted(cells.items(), key=lambda kv: place(kv[1].label_cell)[::-1]))   # column by column
    main = order(main)
    packages = {b: order(cells) for b, cells in packages.items()}

    # A label that is nearly a keyword the loader reads, and is not one: the loader ignores that row.
    known = {k.lower() for k in main}
    absent = {k.lower(): k for k in main_keywords if k.lower() not in known}
    labels = {cl.label_cell for cl in main.values()}
    all_main = {k.lower() for k in main_keywords}
    for r in range(main_hi):
        for c, cell in enumerate(grid[r]):
            if not (isinstance(cell, str) and len(cell.strip()) >= 4) or coord(r, c) in labels:
                continue
            if _cell(grid, r, c + 1) is None or cell.strip().lower() in all_main:
                continue
            near = difflib.get_close_matches(cell.strip().lower(), list(absent), n=1, cutoff=near_cutoff)
            if near:
                kw = absent[near[0]]
                issues.append(_issue('ask', 'near_keyword',
                                     'label "%s" in %s is not a keyword the loader reads, and %s is not in this '
                                     'config; did you mean "%s"?' % (cell.strip(), coord(r, c), kw, kw),
                                     label=cell.strip(), cell=coord(r, c), keyword=kw))
    return SimpleNamespace(main=main, packages=packages, links=links, blocks=blocks, main_rows=(0, main_hi),
                           issues=issues)


# ----------------------------------------------------------------------------- export
def _same_cell_value(a, b):
    if isinstance(a, bool) or isinstance(b, bool):
        return type(a) is type(b) and a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float(a) == float(b)
    return type(a) is type(b) and a == b


def export_workbook(working_config, source, out, overwrite=False):
    """Write the working config as a copy of its source workbook, at out.

    Only the cells whose value differs from the source are rewritten; a formula cell whose value changed is
    written as that value. A key the sheet does not hold is added as a keyword and value pair: in the main
    area in free cells above the first .START row, in a package block inside the block. The source is never
    written to: out naming the source is refused, and so is an existing out unless overwrite is set.

    Returns {'path', 'written': [{json_path, cell, before, after, was_formula}], 'added': [{json_path,
    label_cell, cell, after}], 'warnings': [text]}."""
    if os.path.exists(out) and os.path.samefile(out, source) or \
            os.path.normcase(os.path.abspath(out)) == os.path.normcase(os.path.abspath(source)):
        raise ValueError('refusing to write over the source config %s: name another file' % source)
    if os.path.exists(out) and not overwrite:
        raise FileExistsError('%s exists already' % out)
    readable, tmp, _ = _readable_copy(source)
    try:
        doc = read_workbook(readable)
        found = scan(doc.grid, doc.formulas)
        edits, written, added, warnings = {}, [], [], []
        used = set()
        for block, values in [(None, working_config['main'])] + list(working_config['packages'].items()):
            if block is not None and block not in found.packages:
                warnings.append('package block %s is not in %s and was not written' % (block, os.path.basename(source)))
                continue
            cells = found.main if block is None else found.packages[block]
            for key, value in values.items():
                jp = '/main/%s' % key if block is None else '/packages/%s/%s' % (block, key)
                cl = cells.get(key)
                if cl is not None:
                    if cl.lost or not _same_cell_value(cl.value, value):
                        edits[cl.cell] = value
                        written.append(dict(json_path=jp, cell=cl.cell, before=cl.value, after=value,
                                            was_formula=cl.formula is not None))
                        if cl.formula is not None:
                            warnings.append('%s was a formula (=%s); it is now the value %r' % (key, cl.formula, value))
                    continue
                if block is None:
                    lo, hi = found.main_rows
                else:
                    _, s, e = next(b for b in found.blocks if b[0] == block)
                    lo, hi = s + 1, e
                label_cols = sorted({place(c.label_cell)[1] for c in cells.values()})
                slot = _free_pair(doc.grid, lo, hi, label_cols, used)
                if slot is None:
                    raise ValueError('no free pair of cells for %s %s' % (key, 'above the first .START row' if block is None
                                                                          else 'inside package block %s' % block))
                r, c = slot
                used.update({(r, c), (r, c + 1)})
                edits[coord(r, c)] = key
                edits[coord(r, c + 1)] = value
                added.append(dict(json_path=jp, label_cell=coord(r, c), cell=coord(r, c + 1), after=value))
        for key in [k for k in found.main if k not in working_config['main']]:
            warnings.append('%s is in the sheet but not in the working config; the sheet keeps it' % key)
        # A formula that refers to a rewritten cell keeps a stale stored result until the workbook is recalculated.
        for cell, f in doc.formulas.items():
            hit = sorted({m.group(1) + m.group(2) for m in REF.finditer(f)} & set(edits))
            if cell not in edits and hit:
                warnings.append('%s (=%s) refers to %s, which this export changes; its stored result is stale until '
                                'the workbook is recalculated' % (cell, f, ', '.join(hit)))
        _rewrite_cells(readable, out, doc.sheet, edits, doc.formulas)
    finally:
        if tmp:
            shutil.rmtree(tmp, ignore_errors=True)
    return {'path': out, 'written': written, 'added': added, 'warnings': warnings}


def _free_pair(grid, lo, hi, label_cols, used):
    """(row, column) of a free label cell in rows lo..hi-1 whose right neighbour is free too, and whose left
    neighbour is empty (else the new label would be read as the value of the cell to its left). Columns that
    already hold labels are tried first, top to bottom; then any column."""
    width = max((len(row) for row in grid), default=0)

    def free(r, c):
        return (r, c) not in used and _cell(grid, r, c) in (None, '')

    def ok(r, c):
        return c + 1 < width and free(r, c) and free(r, c + 1) and (c == 0 or free(r, c - 1))

    for c in label_cols:
        for r in range(lo, hi):
            if ok(r, c):
                return r, c
    for r in range(lo, hi):
        for c in range(width):
            if ok(r, c):
                return r, c
    return None


def _sheet_xml_name(path, sheet_name):
    """Which xl/worksheets/sheetN.xml holds `sheet_name`."""
    with zipfile.ZipFile(path) as z:
        wb = z.read('xl/workbook.xml').decode('utf-8')
        rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    m = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(sheet_name), wb)
    if not m:
        m = re.search(r'<sheet[^>]*r:id="([^"]+)"', wb)
    rid = m.group(1)
    m2 = re.search(r'<Relationship[^>]*Id="%s"[^>]*Target="([^"]+)"' % re.escape(rid), rels)
    target = m2.group(1)
    if target.startswith('/'):
        return target.lstrip('/')
    return 'xl/' + target.lstrip('./')


def _escape(text):
    for a, b in (('&', '&amp;'), ('<', '&lt;'), ('>', '&gt;')):
        text = text.replace(a, b)
    return text


def _fmt(value):
    """(xml_attrs, inner_xml) for a cell value."""
    if value is None:
        return '', ''
    if isinstance(value, bool):
        return ' t="b"', '<v>%d</v>' % int(value)
    if isinstance(value, (int, float)):
        return '', '<v>%s</v>' % repr(value)
    # inlineStr avoids having to touch the shared-strings table at all
    return ' t="inlineStr"', '<is><t xml:space="preserve">%s</t></is>' % _escape(str(value))


def _cell_re(cell):
    return re.compile(r'<c r="%s"([^>]*?)(/>|>(.*?)</c>)' % re.escape(cell), flags=re.S)


def _put_new_cell(xml, cell, cell_xml):
    """Insert a cell the sheet XML does not hold: in column order within its row, and the row in row order."""
    rown, col = place(cell)[0] + 1, place(cell)[1]
    rm = re.search(r'<row[^>]*\sr="%d"[^>]*?(/>|>)' % rown, xml)
    if rm:
        if rm.group(1) == '/>':                         # an empty row written as <row .../>
            return xml[:rm.start()] + rm.group(0)[:-2] + '>' + cell_xml + '</row>' + xml[rm.end():]
        end = xml.index('</row>', rm.end())
        for m in re.finditer(r'<c r="([A-Z]+)[0-9]+"', xml[rm.end():end]):
            if _col_index(m.group(1)) > col:
                at = rm.end() + m.start()
                return xml[:at] + cell_xml + xml[at:]
        return xml[:end] + cell_xml + xml[end:]
    new_row = '<row r="%d">%s</row>' % (rown, cell_xml)
    for m in re.finditer(r'<row[^>]*\sr="([0-9]+)"', xml):
        if int(m.group(1)) > rown:
            return xml[:m.start()] + new_row + xml[m.start():]
    if '<sheetData/>' in xml:
        return xml.replace('<sheetData/>', '<sheetData>%s</sheetData>' % new_row, 1)
    at = xml.index('</sheetData>')
    return xml[:at] + new_row + xml[at:]


def _rewrite_cells(src, dst, sheet_name, edits, formulas):
    """Copy the xlsx, replacing only the given cells in the given sheet.

    edits: {cell: value}. formulas: {cell: formula text} of the sheet as openpyxl reads it, used to write
    out in full the other cells of a shared formula whose cell is overwritten."""
    target = _sheet_xml_name(src, sheet_name)
    with zipfile.ZipFile(src) as z:
        names = z.namelist()
        blobs = {n: z.read(n) for n in names}

    xml = blobs[target].decode('utf-8')
    overwrote_formula = False
    shared = set()
    for cell, value in edits.items():
        attrs, inner = _fmt(value)
        m = _cell_re(cell).search(xml)
        if m:
            old = m.group(3) or ''
            overwrote_formula |= '<f' in old
            sm = re.search(r'<f[^>]*\st="shared"[^>]*\ssi="([0-9]+)"', old)
            if sm:
                shared.add(sm.group(1))
            # keep the style index so the cell still looks like its neighbours
            st = re.search(r'\ss="(\d+)"', m.group(1))
            style = ' s="%s"' % st.group(1) if st else ''
            xml = xml[:m.start()] + '<c r="%s"%s%s>%s</c>' % (cell, style, attrs, inner) + xml[m.end():]
        else:
            xml = _put_new_cell(xml, cell, '<c r="%s"%s>%s</c>' % (cell, attrs, inner))
    # A shared formula whose cell was overwritten (perhaps its first cell, which holds the text the others
    # share): write each remaining cell of it as a formula of its own.
    if shared:
        members = [m.group(1) for m in re.finditer(r'<c r="([A-Z]+[0-9]+)"[^>]*?(?:/>|>(.*?)</c>)', xml, flags=re.S)
                   if m.group(2) and any(re.search(r'<f[^>]*\st="shared"[^>]*\ssi="%s"' % si, m.group(2))
                                         for si in shared)]
        for cell in members:
            cm = _cell_re(cell).search(xml)
            f_new = '<f>%s</f>' % _escape(formulas[cell])
            body = re.sub(r'<f[^>]*?(?:/>|>.*?</f>)', lambda _: f_new, cm.group(3), count=1, flags=re.S)
            xml = xml[:cm.start(3)] + body + xml[cm.end(3):]
    blobs[target] = xml.encode('utf-8')
    if overwrote_formula and 'xl/calcChain.xml' in blobs:
        # The calculation chain lists formula cells; one that is now a value would make Excel repair the
        # file. Without the chain Excel rebuilds it.
        names.remove('xl/calcChain.xml')
        del blobs['xl/calcChain.xml']
        ct = blobs['[Content_Types].xml'].decode('utf-8')
        blobs['[Content_Types].xml'] = re.sub(r'<Override[^>]*PartName="/xl/calcChain.xml"[^>]*/>', '', ct).encode('utf-8')
        rels = blobs['xl/_rels/workbook.xml.rels'].decode('utf-8')
        blobs['xl/_rels/workbook.xml.rels'] = re.sub(r'<Relationship[^>]*Target="[^"]*calcChain.xml"[^>]*/>', '',
                                                     rels).encode('utf-8')

    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in names:
            z.writestr(n, blobs[n])
